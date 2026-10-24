# -*- coding: utf-8 -*-
"""Excel çalışma kitabındaki seçili sayfaları web için JSON'a çevirir.

Her sayfa için: sütun genişlikleri (px), satır yükseklikleri (px), görünen hücre metni
(Türkçe sayı biçimi), ham sayı, stil (dolgu, yazı rengi, kalın, boyut, hiza, kenarlık),
birleşik hücreler, dondurulmuş bölme, koşullu biçimlerin sonucu (hücreye işlenmiş renk)
ve grafikler (tür, başlık, kategoriler, seriler, konum).

Rehberli derslerin (Datakit) kaynağıdır: çıktı data/<urun>/sayfalar/<slug>.json olur, scripts/derle.py oradan yayınlar.
Hücrelerin önbellekteki değerleri okunur: dosya Excel'de hesaplanıp kaydedilmiş olmalı. Derlemenin parçası değildir;
yalnızca openpyxl gerekir (pip install openpyxl). Excel dosyası repoda durmaz.

Kullanım: python3 scripts/rehber-disa-aktar.py KITAP.xlsx data/datakit/sayfalar "Sayfa Adı=slug" ...
"""
import sys, os, re, json, math, zipfile, colorsys, datetime, warnings
from decimal import Decimal, ROUND_HALF_UP
import xml.etree.ElementTree as ET
import openpyxl
from openpyxl.utils import get_column_letter, column_index_from_string

warnings.filterwarnings('ignore')

AYLAR = ['Oca', 'Şub', 'Mar', 'Nis', 'May', 'Haz', 'Tem', 'Ağu', 'Eyl', 'Eki', 'Kas', 'Ara']
NS = {
    'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
    'c': 'http://schemas.openxmlformats.org/drawingml/2006/chart',
    'xdr': 'http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing',
    'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
}

# ---------------------------------------------------------------- sayı biçimi

def yuvarla(x, basamak):
    q = Decimal(1).scaleb(-basamak)
    return Decimal(repr(abs(x))).quantize(q, rounding=ROUND_HALF_UP)


def tr_sayi(x, basamak, binlik=True):
    d = yuvarla(x, basamak)
    tam, _, kesir = format(d, 'f').partition('.')
    if binlik:
        tam = '{:,}'.format(int(tam)).replace(',', '.')
    s = tam + (',' + kesir if basamak else '')
    return s


def isaret(x, s):
    return ('-' + s) if x < 0 and s.strip('0,.') else s


def bicimle(v, fmt):
    """Excel sayı biçimini Türkçe görünüme çevirir. Döner: (metin, sayı_mı)."""
    if isinstance(v, (datetime.datetime, datetime.date)):
        if 'mmm' in fmt:
            return f"{AYLAR[v.month - 1]} {v.year % 100:02d}", False
        if 'yyyy' in fmt:
            return f"{v.day:02d}.{v.month:02d}.{v.year}", False
        if 'yy' in fmt:
            return f"{v.day:02d}.{v.month:02d}.{v.year % 100:02d}", False
        return f"{v.day:02d}.{v.month:02d}.{v.year}", False
    if isinstance(v, bool):
        return ('DOĞRU' if v else 'YANLIŞ'), False
    if not isinstance(v, (int, float)):
        return str(v), False
    x = float(v)
    f = fmt.replace('\\', '')
    artili = f.startswith('+')
    yuzde = '%' in f
    m = re.search(r'0\.(0+)', f)
    basamak = len(m.group(1)) if m else 0
    binlik = '#,##0' in f
    if f == 'General':
        if float(x).is_integer():
            return tr_sayi(x, 0, binlik=False) if abs(x) < 1e15 else str(x), True
        s = ('%.10g' % x)
        if 'e' in s:
            return s, True
        tam, _, kesir = s.partition('.')
        return tam.replace('-', '-') + (',' + kesir if kesir else ''), True
    if yuzde:
        s = '%' + tr_sayi(x * 100, basamak, binlik=False)
    else:
        s = tr_sayi(x, basamak, binlik=binlik)
    if '" puan"' in fmt or ' puan' in f:
        s += ' puan'
    if artili:
        if x > 0:
            s = '+' + s
        elif x < 0:
            s = '-' + s
        return s, True
    return isaret(x, s), True


def metni_turkcele(s):
    s = re.sub(r'\d{1,3}(?:,\d{3})+', lambda m: m.group(0).replace(',', '.'), s)
    s = re.sub(r'%(\d+)\.(\d+)', r'%\1,\2', s)
    s = re.sub(r'(?<![\d.])(\d+)\.(\d+)%', r'%\1,\2', s)
    return s

# ---------------------------------------------------------------- renkler

def tema_renkleri(z):
    try:
        x = z.read('xl/theme/theme1.xml')
    except KeyError:
        return {}
    t = ET.fromstring(x)
    cs = t.find('.//a:clrScheme', NS)
    sira = ['dk1', 'lt1', 'dk2', 'lt2', 'accent1', 'accent2', 'accent3', 'accent4', 'accent5', 'accent6', 'hlink', 'folHlink']
    renk = {}
    for ad in sira:
        e = cs.find('a:' + ad, NS)
        if e is None:
            continue
        c = e[0]
        val = c.get('lastClr') or c.get('val')
        renk[ad] = val
    # Excel tema dizini: 0=lt1 1=dk1 2=lt2 3=dk2 4..9 accent, 10 hlink, 11 folHlink
    dizin = ['lt1', 'dk1', 'lt2', 'dk2', 'accent1', 'accent2', 'accent3', 'accent4', 'accent5', 'accent6', 'hlink', 'folHlink']
    return {i: renk.get(ad) for i, ad in enumerate(dizin)}


def ton(hexrgb, tint):
    r, g, b = (int(hexrgb[i:i + 2], 16) / 255 for i in (0, 2, 4))
    h, l, s = colorsys.rgb_to_hls(r, g, b)
    l = l * (1 + tint) if tint < 0 else l * (1 - tint) + tint
    r, g, b = colorsys.hls_to_rgb(h, max(0, min(1, l)), s)
    return '%02X%02X%02X' % (round(r * 255), round(g * 255), round(b * 255))


def renk(col, tema):
    if col is None:
        return None
    try:
        if col.type == 'rgb' and isinstance(col.rgb, str):
            v = col.rgb[-6:]
            if col.rgb in ('00000000',):
                return None
            return '#' + v.upper()
        if col.type == 'theme' and col.theme is not None:
            base = tema.get(col.theme)
            if not base:
                return None
            if col.tint:
                base = ton(base, col.tint)
            return '#' + base.upper()
        if col.type == 'indexed':
            return None
    except Exception:
        return None
    return None

# ---------------------------------------------------------------- koşullu biçim

def kosullu_bicimler(ws, wsv, adlar):
    """Kitaptaki koşullu biçim kurallarını hesaplayıp {hücre: {'yazi':, 'kalin':, 'dolgu':}} döner."""
    sonuc = {}

    def ekle(ref, d):
        sonuc.setdefault(ref, {}).update({k: v for k, v in d.items() if v is not None})

    for rng in ws.conditional_formatting:
        hucreler = []
        for cr in str(rng.sqref).split():
            minc, minr, maxc, maxr = openpyxl.utils.cell.range_boundaries(cr)
            for r in range(minr, maxr + 1):
                for c in range(minc, maxc + 1):
                    hucreler.append((r, c))
        for kural in rng.rules:
            if kural.type == 'colorScale':
                renkler = ['#' + c.rgb[-6:] for c in kural.colorScale.color]
                degerler = [wsv.cell(r, c).value for r, c in hucreler]
                sayilar = [v for v in degerler if isinstance(v, (int, float)) and not isinstance(v, bool)]
                if not sayilar:
                    continue
                lo, hi = min(sayilar), max(sayilar)
                srt = sorted(sayilar)
                orta = srt[len(srt) // 2] if len(srt) % 2 else (srt[len(srt) // 2 - 1] + srt[len(srt) // 2]) / 2
                for (r, c), v in zip(hucreler, degerler):
                    if not isinstance(v, (int, float)) or isinstance(v, bool):
                        continue
                    if len(renkler) == 2:
                        t = 0 if hi == lo else (v - lo) / (hi - lo)
                        col = karistir(renkler[0], renkler[1], t)
                    else:
                        if v <= orta:
                            t = 0 if orta == lo else (v - lo) / (orta - lo)
                            col = karistir(renkler[0], renkler[1], t)
                        else:
                            t = 0 if hi == orta else (v - orta) / (hi - orta)
                            col = karistir(renkler[1], renkler[2], t)
                    ekle(f'{get_column_letter(c)}{r}', {'dolgu': col})
                continue
            if kural.type != 'expression' or not kural.formula:
                continue
            f = kural.formula[0]
            dx = kural.dxf
            stil = {}
            if dx is not None:
                if dx.font is not None:
                    if dx.font.color is not None and dx.font.color.rgb:
                        stil['yazi'] = '#' + dx.font.color.rgb[-6:]
                    if dx.font.b:
                        stil['kalin'] = True
                if dx.fill is not None:
                    bg = dx.fill.bgColor if dx.fill.bgColor is not None else None
                    if bg is not None and bg.rgb and bg.rgb not in ('00000000',):
                        stil['dolgu'] = '#' + bg.rgb[-6:]
            ilk_r, ilk_c = hucreler[0]
            for r, c in hucreler:
                if ifade_dogru(f, r - ilk_r, c - ilk_c, wsv, adlar):
                    ekle(f'{get_column_letter(c)}{r}', stil)
    return sonuc


def karistir(a, b, t):
    a = [int(a[i:i + 2], 16) for i in (1, 3, 5)]
    b = [int(b[i:i + 2], 16) for i in (1, 3, 5)]
    return '#%02X%02X%02X' % tuple(round(x + (y - x) * max(0, min(1, t))) for x, y in zip(a, b))


def hucre_kaydir(ref, dr, dc):
    m = re.fullmatch(r'(\$?)([A-Z]+)(\$?)(\d+)', ref)
    ca, col, ra, row = m.groups()
    c = column_index_from_string(col) + (0 if ca else dc)
    r = int(row) + (0 if ra else dr)
    return r, c


def ifade_dogru(f, dr, dc, wsv, adlar):
    def deger(ref):
        r, c = hucre_kaydir(ref, dr, dc)
        return wsv.cell(r, c).value
    m = re.fullmatch(r'NOT\(ISERROR\(SEARCH\("(.+?)",(\$?[A-Z]+\$?\d+)\)\)\)', f)
    if m:
        v = deger(m.group(2))
        return isinstance(v, str) and m.group(1).lower() in v.lower()
    m = re.fullmatch(r'AND\(ISNUMBER\((\$?[A-Z]+\$?\d+)\),(\$?[A-Z]+\$?\d+)([<>]=?)(.+)\)', f)
    if m:
        v = deger(m.group(2))
        if not isinstance(v, (int, float)) or isinstance(v, bool):
            return False
        sag = m.group(4)
        esik = adlar.get(sag)
        if esik is None:
            try:
                esik = float(sag)
            except ValueError:
                return False
        op = m.group(3)
        return {'<': v < esik, '>': v > esik, '<=': v <= esik, '>=': v >= esik}[op]
    m = re.fullmatch(r'(\$?[A-Z]+\$?\d+)="(.*)"', f)
    if m:
        return deger(m.group(1)) == m.group(2)
    raise ValueError('Tanınmayan koşullu biçim: ' + f)

# ---------------------------------------------------------------- grafikler

def grafikler(z, sayfa_yolu):
    rels = sayfa_yolu.replace('worksheets/', 'worksheets/_rels/') + '.rels'
    if rels not in z.namelist():
        return []
    sr = ET.fromstring(z.read(rels))
    cikis = []
    for rel in sr:
        if not rel.get('Type', '').endswith('/drawing'):
            continue
        dyol = os.path.normpath(os.path.join(os.path.dirname(sayfa_yolu), rel.get('Target'))).replace('\\', '/')
        dx = ET.fromstring(z.read(dyol))
        drels_yol = dyol.replace('drawings/', 'drawings/_rels/') + '.rels'
        drels = {r.get('Id'): r.get('Target') for r in ET.fromstring(z.read(drels_yol))} if drels_yol in z.namelist() else {}
        for anc in list(dx):
            fr = anc.find('xdr:from', NS); to = anc.find('xdr:to', NS)
            ch = anc.find('.//c:chart', NS)
            if ch is None or fr is None:
                continue
            rid = ch.get('{%s}id' % NS['r'])
            cyol = os.path.normpath(os.path.join(os.path.dirname(dyol), drels[rid])).replace('\\', '/')
            konum = {
                'bas': f"{get_column_letter(int(fr.find('xdr:col', NS).text) + 1)}{int(fr.find('xdr:row', NS).text) + 1}",
                'son': f"{get_column_letter(int(to.find('xdr:col', NS).text) + 1)}{int(to.find('xdr:row', NS).text) + 1}" if to is not None else None,
            }
            cikis.append(dict(grafik_oku(z.read(cyol)), konum=konum))
    return cikis


def onbellek(e):
    if e is None:
        return []
    noktalar = {}
    for pt in e.iter('{%s}pt' % NS['c']):
        v = pt.find('c:v', NS)
        noktalar[int(pt.get('idx'))] = v.text if v is not None else None
    n = e.find('.//c:ptCount', NS)
    say = int(n.get('val')) if n is not None else (max(noktalar) + 1 if noktalar else 0)
    return [noktalar.get(i) for i in range(say)]


def grafik_oku(x):
    t = ET.fromstring(x)
    baslik = ''.join(e.text or '' for e in t.findall('.//c:title//a:t', NS))
    plot = t.find('.//c:plotArea', NS)
    seriler = []; tur = None
    for grp in plot:
        etiket = grp.tag.split('}')[1]
        if not etiket.endswith('Chart'):
            continue
        g_tur = {'lineChart': 'cizgi', 'barChart': 'cubuk', 'pieChart': 'pasta', 'areaChart': 'alan', 'scatterChart': 'dagilim'}.get(etiket, etiket)
        bd = grp.find('c:barDir', NS)
        if bd is not None and bd.get('val') == 'bar':
            g_tur = 'yatay-cubuk'
        tur = tur or g_tur
        for ser in grp.findall('c:ser', NS):
            ad = ''.join(v.text or '' for v in ser.findall('c:tx//c:v', NS)) or ''.join(v.text or '' for v in ser.findall('c:tx//a:t', NS))
            cat = ser.find('c:cat', NS)
            kats_ham = onbellek(cat.find('.//c:strCache', NS) if cat is not None and cat.find('.//c:strCache', NS) is not None else (cat.find('.//c:numCache', NS) if cat is not None else None))
            ING = ['Jan','Feb','Mar','Apr','May','Jun','Jul','Aug','Sep','Oct','Nov','Dec']
            kats = [(AYLAR[ING.index(k[:3])] + k[3:]) if isinstance(k, str) and k[:3] in ING and re.fullmatch(r'[A-Z][a-z]{2} \d{2}', k) else k for k in kats_ham]
            val = ser.find('c:val', NS)
            degs = [float(v) if v not in (None, '') else None for v in onbellek(val.find('.//c:numCache', NS) if val is not None else None)]
            fmt_e = val.find('.//c:formatCode', NS) if val is not None else None
            renk_e = ser.find('.//c:spPr//a:srgbClr', NS)
            seriler.append({'ad': ad, 'tur': g_tur, 'kategoriler': kats, 'degerler': degs,
                            'bicim': fmt_e.text if fmt_e is not None else None,
                            'renk': ('#' + renk_e.get('val')) if renk_e is not None else None})
    return {'tur': tur, 'baslik': baslik, 'seriler': seriler}

# ---------------------------------------------------------------- ana akış

def px_genislik(w):
    return int(round(w * 7 + 5)) if w else 0


def disa_aktar(kitap, cikis, secim):
    z = zipfile.ZipFile(kitap)
    tema = tema_renkleri(z)
    wbv = openpyxl.load_workbook(kitap, data_only=True)
    wbf = openpyxl.load_workbook(kitap, data_only=False)
    adlar = {}
    for ad, dn in wbv.defined_names.items():
        try:
            dests = list(dn.destinations)
            if len(dests) == 1:
                s, ref = dests[0]
                v = wbv[s][ref.replace('$', '')].value
                if isinstance(v, (int, float)):
                    adlar[ad] = v
        except Exception:
            pass
    # sayfa adı -> xml yolu
    wbx = ET.fromstring(z.read('xl/workbook.xml'))
    rels = {r.get('Id'): r.get('Target') for r in ET.fromstring(z.read('xl/_rels/workbook.xml.rels'))}
    yol = {}
    for s in wbx.iter('{http://schemas.openxmlformats.org/spreadsheetml/2006/main}sheet'):
        rid = s.get('{%s}id' % NS['r'])
        t = rels[rid]
        yol[s.get('name')] = ('xl/' + t) if not t.startswith('/') else t[1:]
    os.makedirs(cikis, exist_ok=True)
    ozet = []
    for ad, slug in secim:
        wsv = wbv[ad]; wsf = wbf[ad]
        cf = kosullu_bicimler(wsf, wsv, adlar)
        # dolu alan
        son_r = son_c = 0
        for row in wsv.iter_rows():
            for c in row:
                if c.value not in (None, ''):
                    son_r = max(son_r, c.row); son_c = max(son_c, c.column)
        varsayilan = wsv.sheet_format.defaultColWidth or 8.43
        genislik = {}
        for key, dim in wsv.column_dimensions.items():
            lo = dim.min or column_index_from_string(key); hi = dim.max or lo
            for i in range(lo, hi + 1):
                genislik[i] = (dim.width if dim.width else varsayilan, bool(dim.hidden))
        sutunlar = []
        for i in range(1, son_c + 1):
            w, gizli = genislik.get(i, (varsayilan, False))
            sutunlar.append({'harf': get_column_letter(i), 'genislik': 0 if gizli else px_genislik(w)})
        stiller = []; stil_no = {}
        satirlar = []
        for r in range(1, son_r + 1):
            rd = wsv.row_dimensions.get(r)
            if rd is not None and rd.hidden:
                continue
            h = rd.height if rd is not None and rd.height else 15
            hucreler = []
            for c in range(1, son_c + 1):
                hv = wsv.cell(r, c); hf = wsf.cell(r, c)
                v = hv.value
                ref = f'{get_column_letter(c)}{r}'
                st = {}
                if hv.fill is not None and hv.fill.fill_type == 'solid':
                    d = renk(hv.fill.fgColor, tema)
                    if d and d != '#FFFFFF':
                        st['dolgu'] = d
                if hv.font is not None:
                    y = renk(hv.font.color, tema) if hv.font.color is not None else None
                    if y and y not in ('#000000',):
                        st['yazi'] = y
                    if hv.font.b:
                        st['kalin'] = True
                    if hv.font.i:
                        st['italik'] = True
                    if hv.font.sz and float(hv.font.sz) != 10:
                        st['boyut'] = float(hv.font.sz)
                al = hv.alignment
                if al is not None:
                    if al.horizontal in ('center', 'right', 'left', 'centerContinuous'):
                        st['hiza'] = 'center' if al.horizontal == 'centerContinuous' else al.horizontal
                    if al.wrap_text:
                        st['sar'] = True
                    if al.vertical in ('center', 'top'):
                        st['dikey'] = al.vertical
                b = hv.border
                if b is not None and (b.left.style or b.right.style or b.top.style or b.bottom.style):
                    st['kenar'] = True
                if ref in cf:
                    st.update(cf[ref])
                if v in (None, '') and not st:
                    continue
                kayit = {'c': get_column_letter(c)}
                if v not in (None, ''):
                    metin, sayi_mi = bicimle(v, hv.number_format)
                    if isinstance(v, str):
                        metin = metni_turkcele(metin)
                    kayit['v'] = metin
                    if sayi_mi and isinstance(v, (int, float)) and not isinstance(v, bool):
                        kayit['n'] = v if not (isinstance(v, float) and (math.isnan(v) or math.isinf(v))) else None
                        if 'hiza' not in st:
                            st['hiza'] = 'right'
                    if isinstance(hf.value, str) and hf.value.startswith('='):
                        kayit['f'] = True
                if st:
                    anahtar = json.dumps(st, sort_keys=True)
                    if anahtar not in stil_no:
                        stil_no[anahtar] = len(stiller); stiller.append(st)
                    kayit['s'] = stil_no[anahtar]
                hucreler.append(kayit)
            satirlar.append({'r': r, 'yukseklik': int(round(h * 4 / 3)), 'hucreler': hucreler})
        birlesik = [str(m) for m in wsv.merged_cells.ranges]
        # Excel'deki gibi metin taşması: sola dayalı, sarılmayan metin sağdaki boş hücrelere taşar.
        # Değeri ya da formülü olan hücre (formül boş metin döndürse bile) ve birleşik hücre taşmayı durdurur.
        birlesik_hucre = set()
        for m in wsv.merged_cells.ranges:
            for r in range(m.min_row, m.max_row + 1):
                for c in range(m.min_col, m.max_col + 1):
                    birlesik_hucre.add((r, c))

        def bos(r, c):
            if (r, c) in birlesik_hucre:
                return False
            v = wsv.cell(r, c).value; f = wsf.cell(r, c).value
            return v in (None, '') and f in (None, '')

        for sat in satirlar:
            r = sat['r']
            for h in sat['hucreler']:
                if 'v' not in h or 'n' in h:
                    continue
                st = stiller[h['s']] if 's' in h else {}
                c = column_index_from_string(h['c'])
                if st.get('sar') or st.get('hiza') in ('right', 'center') or (r, c) in birlesik_hucre:
                    continue
                k = 0
                while c + k + 1 <= son_c and bos(r, c + k + 1):
                    k += 1
                if k:
                    h['tasma'] = k
        # Grafikler: dolu alanın hemen yanında başlayıp taşan grafik için sütun listesi uzatılır (Dashboard N sütunu);
        # dolu alanın tamamen dışında duran grafik (Tahmin, S3) "disarida" işaretlenir, web'de tablonun üstünde çizilir.
        grafik_listesi = grafikler(z, yol[ad])
        for g in grafik_listesi:
            bas_c = column_index_from_string(re.match(r'[A-Z]+', g['konum']['bas']).group(0))
            son_gc = column_index_from_string(re.match(r'[A-Z]+', g['konum']['son']).group(0))
            if bas_c > son_c + 1:
                g['disarida'] = True
                continue
            for i in range(len(sutunlar) + 1, son_gc + 1):
                w, gizli = genislik.get(i, (varsayilan, False))
                sutunlar.append({'harf': get_column_letter(i), 'genislik': 0 if gizli else px_genislik(w)})
            son_gr = int(re.search(r'\d+', g['konum']['son']).group(0))
            for r in range(satirlar[-1]['r'] + 1, son_gr + 1):
                rd = wsv.row_dimensions.get(r)
                h = rd.height if rd is not None and rd.height else 15
                satirlar.append({'r': r, 'yukseklik': int(round(h * 4 / 3)), 'hucreler': []})
        veri = {
            'id': slug, 'ad': ad, 'kaynak': os.path.basename(kitap),
            'boyut': {'satir': son_r, 'sutun': son_c},
            'dondur': wsf.freeze_panes,
            'sutunlar': sutunlar, 'satirlar': satirlar, 'birlesik': birlesik,
            'stiller': stiller, 'grafikler': grafik_listesi,
        }
        with open(os.path.join(cikis, slug + '.json'), 'w', encoding='utf-8') as f:
            json.dump(veri, f, ensure_ascii=False, separators=(',', ':'))
        ozet.append((ad, slug, son_r, son_c, len(birlesik), len(veri['grafikler']), os.path.getsize(os.path.join(cikis, slug + '.json'))))
    return ozet


if __name__ == '__main__':
    if len(sys.argv) < 4:
        sys.exit('Kullanım: python3 scripts/rehber-disa-aktar.py KITAP.xlsx CIKIS_KLASORU "Sayfa Adı=slug" ...')
    kitap, cikis = sys.argv[1], sys.argv[2]
    secim = [a.split('=', 1) for a in sys.argv[3:]]
    for satir in disa_aktar(kitap, cikis, secim):
        print('%-18s %-18s satır %3d sütun %2d birleşik %2d grafik %d  %6d bayt' % satir)
