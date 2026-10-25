# RÖNTGEN (stok) ve RANDEVU
from lib import *
from datetime import date, time
from openpyxl import load_workbook

def rontgen():
    B=Book("RÖNTGEN","Stok takibi","Rafın içini gösterir: ne girdi, ne çıktı, ne kaldı, ne bitmek üzere.")
    NU,NH=300,1500; U="Ürünler"; H="Hareketler"
    HA,HB,HC,HD=[rg(H,c,NH) for c in "ABCD"]
    wu=B.sheet(U,"Ürünler","Sarı sütunları doldur. Giren, çıkan, mevcut ve durum Hareketler sayfasından hesaplanır.",{"A":28,"B":9,"C":15,"D":12,"E":12,"F":10,"G":10,"H":11,"I":15,"J":11,"K":14,"L":13,"M":4,"N":4})
    B.head(wu,4,["Ürün","Birim","Birim maliyet (₺)","Açılış stoku","Kritik seviye","Giren","Çıkan","Mevcut","Stok değeri (₺)","Durum","Son hareket","Hareketsiz gün","y1","y2"],calc=range(6,15))
    B.body(wu,R0,NU,14,calc=range(6,15),fmt={3:TLI,4:NUM,5:NUM,6:NUM,7:NUM,8:NUM,9:TL,11:DT,12:NUM},center=(2,10,11),bold=(8,))
    urun=[("Çay 1 kg","Paket",250,40,10),("Toz şeker 5 kg","Paket",219,30,8),("Ayçiçek yağı 5 L","Şişe",450,24,6),("Makarna 500 g","Paket",17,200,50),("Pirinç 1 kg","Paket",50,80,20),("Bulaşık deterjanı","Şişe",95,36,10)]
    B.put(wu,R0,urun)
    for i in range(NU):
        r=R0+i
        wu.cell(r,6,f'=IF(A{r}="","",SUMIFS({HD},{HC},A{r},{HB},"Giriş"))'); wu.cell(r,7,f'=IF(A{r}="","",SUMIFS({HD},{HC},A{r},{HB},"Çıkış"))')
        wu.cell(r,8,f'=IF(A{r}="","",D{r}+F{r}-G{r})'); wu.cell(r,9,f'=IF(A{r}="","",H{r}*C{r})')
        wu.cell(r,10,f'=IF(A{r}="","",IF(H{r}<=0,"Tükendi",IF(H{r}<=E{r},"Kritik","Yeterli")))')
        mx=f"SUMPRODUCT(MAX(({HC}=A{r})*{HA}))"
        wu.cell(r,11,f'=IF(A{r}="","",IF({mx}=0,"",{mx}))'); wu.cell(r,12,f'=IF(OR(A{r}="",K{r}=""),"",TODAY()-K{r})')
        wu.cell(r,13,f'=IF(OR(J{r}="Tükendi",J{r}="Kritik"),ROW(),"")'); wu.cell(r,14,f'=IF(A{r}="","",I{r}+ROW()/1000000)')
    wu.column_dimensions["M"].hidden=True; wu.column_dimensions["N"].hidden=True; wu.freeze_panes="B5"
    B.dv(wu,"uniq",f"A{R0}:A{R0+NU-1}",f"=COUNTIF($A${R0}:$A${R0+NU-1},A{R0})=1"); B.dv(wu,"nonneg",f"C{R0}:E{R0+NU-1}")
    for t,fc,tc in (("Tükendi",PINK,RED),("Kritik",AMB,AMBT)): B.cf_eq(wu,f"J{R0}:J{R0+NU-1}",t,fc,tc)
    wh=B.sheet(H,"Hareketler","Her giriş ya da çıkış bir satır. Sayım farkını da Giriş ya da Çıkış olarak yaz.",{"A":12,"B":10,"C":28,"D":10,"E":34,"F":22})
    B.head(wh,4,["Tarih","Tür","Ürün","Adet","Açıklama","Uyarı"],calc=(6,))
    B.body(wh,R0,NH,6,calc=(6,),fmt={1:DT,4:NUM},center=(1,2))
    har=[(date(2026,9,1),"Çıkış","Çay 1 kg",12,"Örnek"),(date(2026,9,3),"Çıkış","Makarna 500 g",90,"Örnek"),(date(2026,9,8),"Giriş","Makarna 500 g",100,"Örnek: toptancı"),(date(2026,9,12),"Çıkış","Toz şeker 5 kg",25,"Örnek"),
         (date(2026,9,20),"Çıkış","Ayçiçek yağı 5 L",24,"Örnek"),(date(2026,9,25),"Çıkış","Çay 1 kg",20,"Örnek"),(date(2026,10,1),"Giriş","Çay 1 kg",30,"Örnek: toptancı"),(date(2026,10,2),"Çıkış","Pirinç 1 kg",35,"Örnek"),
         (date(2026,10,3),"Çıkış","Makarna 500 g",120,"Örnek"),(date(2026,10,4),"Çıkış","Çay 1 kg",25,"Örnek")]
    B.put(wh,R0,har)
    for i in range(NH):
        r=R0+i; wh.cell(r,6,f'=IF(OR(C{r}="",D{r}=""),"",IF(OR(A{r}="",B{r}=""),"Tarih ya da tür eksik",IF(AND(B{r}="Çıkış",IFERROR(INDEX({rg(U,"H",NU)},MATCH(C{r},{rg(U,"A",NU)},0)),0)<0),"Stok eksiye düştü","")))')
    wh.freeze_panes="A5"; B.dv(wh,"date",f"A{R0}:A{R0+NH-1}"); B.dv(wh,"list",f"B{R0}:B{R0+NH-1}","Giriş,Çıkış"); B.dv(wh,"list",f"C{R0}:C{R0+NH-1}","="+rg(U,"A",NU)); B.dv(wh,"pos",f"D{R0}:D{R0+NH-1}")
    B.cf_f(wh,f"F{R0}:F{R0+NH-1}",f'F{R0}<>""',PINK,RED); B.cf_eq(wh,f"B{R0}:B{R0+NH-1}","Giriş",WHITE,GREEN); B.cf_eq(wh,f"B{R0}:B{R0+NH-1}","Çıkış",WHITE,RED)
    wo=B.sheet("Özet","Özet","Bu sayfada yalnızca yılı ve hareketsizlik eşiğini değiştir.",{"A":3,"B":16,"C":16,"D":16,"E":16,"F":18,"G":3,"H":28,"I":14,"J":14},grid=False)
    wo["B3"]="Yıl"; B.inp(wo,"C3",2026,"0"); wo["D3"]="Hareketsiz sayılma eşiği (gün)"; wo["D3"].font=font(9); B.inp(wo,"F3",60,"0")
    UA,UH,UI,UJ,UL,UM,UN=[rg(U,c,NU) for c in "AHIJLMN"]
    B.cards(wo,5,[("Ürün çeşidi",f"=COUNTA({UA})",NUM),("Stok değeri",f"=SUM({UI})",TL),("Kritik ürün",f'=COUNTIF({UJ},"Kritik")',NUM,1),("Tükenen ürün",f'=COUNTIF({UJ},"Tükendi")',NUM,1),("Hareketsiz ürün",f'=COUNTIF({UL},">="&F3)',NUM)])
    B.label(wo,"B8","Aylık hareket (adet)"); B.table(wo,9,2,["Ay","Giren","Çıkan"],12,fmt={3:NUM,4:NUM})
    for m,ay in enumerate(months(),1):
        r=9+m; wo.cell(r,2,ay); wo.cell(r,3,"="+mrange(HD,HA,"$C$3",m,f',{HB},"Giriş"')); wo.cell(r,4,"="+mrange(HD,HA,"$C$3",m,f',{HB},"Çıkış"'))
    B.label(wo,"H8","Stok uyarısı (ilk 10)"); B.table(wo,9,8,["Ürün","Mevcut","Durum"],10,fmt={9:NUM},center=(10,))
    for k in range(1,11):
        r=9+k; wo.cell(r,8,f"=IFERROR(INDEX('{U}'!$A$1:$A${R0+NU-1},SMALL({UM},{k})),\"\")")
        wo.cell(r,9,f'=IF(H{r}="","",INDEX({UH},MATCH(H{r},{UA},0)))'); wo.cell(r,10,f'=IF(H{r}="","",INDEX({UJ},MATCH(H{r},{UA},0)))')
    B.label(wo,"H21","Parası en çok bağlı 5 ürün"); B.table(wo,22,8,["Ürün","Stok değeri (₺)"],5,fmt={9:TL})
    for k in range(1,6):
        r=22+k; wo.cell(r,8,f'=IFERROR(IF(LARGE({UN},{k})<1,"",INDEX({UA},MATCH(LARGE({UN},{k}),{UN},0))),"")'); wo.cell(r,9,f'=IF(H{r}="","",INDEX({UI},MATCH(H{r},{UA},0)))')
    chart(wo,"B23","Aylık giren ve çıkan (adet)",Reference(wo,min_col=2,min_row=10,max_row=21),[Reference(wo,min_col=3,min_row=9,max_row=21),Reference(wo,min_col=4,min_row=9,max_row=21)],[GREEN,"D92D20"],w=15)
    B.basla([("NASIL KULLANILIR",None),("1. Ürünler","Ürünlerini, birim maliyetini, elindeki açılış stokunu ve kritik seviyeyi yaz."),("2. Hareketler","Her mal girişini ve çıkışını bir satır olarak gir. Sayımda fark çıkarsa farkı Giriş ya da Çıkış satırıyla düzelt."),("3. Özet","Stok değeri, kritik ve tükenen ürünler, hareketsiz ürünler ve aylık hareket kendiliğinden hesaplanır."),
      (None,None),("BİLMEN GEREKENLER",None),("Stok değeri","Mevcut adet × Ürünler sayfasındaki birim maliyet. Maliyeti değiştirirsen stok değeri de değişir."),("Hareketsiz ürün","Son hareketinden bu yana eşikten (varsayılan 60 gün) uzun süre geçen ürün. Hiç hareketi olmayan ürün sayılmaz."),("Kapasite",f"{NU} ürün ve {NH} hareket. Ürün adını sonradan değiştirirsen eski hareketlerdeki adı da değiştir.")])
    B.ai("Tedarikçi ve müşteri adlarını takma adla değiştir.",[("Sipariş listesi","[Ürünler sayfası] Durumu Kritik ya da Tükendi olan ürünler için sipariş listesi hazırla. Her üründen kritik seviyenin üç katına çıkacak kadar adet öner."),("Ölü stok","[Ürünler sayfası] Hareketsiz gün sayısı en yüksek ürünleri ve bunlara bağlı parayı listele. Eritmek için birer öneri ver."),("Hızlı dönenler","[Hareketler sayfası] Son 30 günde en çok çıkan 5 ürünü bul. Bu hızla her birinin mevcut stoku kaç gün yeter, hesapla."),("Kritik seviye önerisi","[Ürünler ve Hareketler sayfaları] Her ürünün haftalık ortalama çıkışına bakarak iki haftalık satışı karşılayacak kritik seviye öner."),("Sayım planı","[Ürünler sayfası] Stok değeri en yüksek ürünlerden başlayarak haftalık bir sayım planı çıkar."),("Hareket kontrolü","[Hareketler sayfası] Olağan dışı yüksek adetli, tarihi eksik ya da iki kez girilmiş gibi duran satırları göster.")])
    B.save(["Başla","Hareketler","Özet","Ürünler","AI Komutları"],OUT+"ExcelKlinigi_RONTGEN_Stok_Takibi.xlsx")
    w=load_workbook(OUT+"ExcelKlinigi_RONTGEN_Stok_Takibi.xlsx",data_only=True); u=w[U]; o=w["Özet"]
    mev=[u.cell(R0+i,8).value for i in range(6)]; exp=[40+30-57,30-25,0,200+100-210,80-35,36]
    assert mev==exp,(mev,exp); assert [u.cell(R0+i,10).value for i in range(6)]==["Yeterli","Kritik","Tükendi","Yeterli","Yeterli","Yeterli"]
    assert o["C6"].value==sum(m*c[2] for m,c in zip(exp,urun)); assert o["D6"].value==1 and o["E6"].value==1
    assert (o["C18"].value,o["D18"].value)==(100,171) and (o["C19"].value,o["D19"].value)==(30,180)
    assert o["H10"].value=="Toz şeker 5 kg" and o["H11"].value=="Ayçiçek yağı 5 L"; print("  RÖNTGEN doğrulandı")

def randevu():
    B=Book("RANDEVU","Randevu defteri","Kuaför, usta, danışman ve klinik için: kim, ne zaman, hangi hizmet, geldi mi, ne kazandırdı.")
    NR,NM,NHz=1000,300,50; R="Randevular"; M="Müşteriler"; Hz="Hizmetler"
    RA,RB,RC,RD,RE,RG,RH,RJ=[rg(R,c,NR) for c in "ABCDEGHJ"]
    wr=B.sheet(R,"Randevular","Her randevu bir satır. Müşteri gelince Durum sütununu Geldi yap; ciro yalnızca Geldi olanlardan hesaplanır.",{"A":12,"B":8,"C":24,"D":22,"E":11,"F":13,"G":13,"H":9,"I":12,"J":4})
    B.head(wr,4,["Tarih","Saat","Müşteri","Hizmet","Durum","Özel ücret (₺) isteğe bağlı","Tutar (₺)","Bitiş","Çakışma","y"],calc=(7,8,9,10))
    B.body(wr,R0,NR,10,calc=(7,8,9,10),fmt={1:DT,2:"HH:MM",6:TLI,7:TL,8:"HH:MM"},center=(1,2,5,8,9))
    rnd=[(date(2026,10,5),time(10,0),"Müşteri Alfa","Saç kesimi","Geldi",None),(date(2026,10,5),time(10,30),"Müşteri Delta","Fön","Geldi",None),(date(2026,10,5),time(11,0),"Müşteri Beta","Boya","Geldi",2200),
         (date(2026,10,5),time(12,30),"Müşteri Gama","Saç kesimi","Gelmedi",None),(date(2026,10,6),time(14,0),"Müşteri Alfa","Boya","Bekliyor",None),(date(2026,10,6),time(15,30),"Müşteri Delta","Saç kesimi","Bekliyor",None),
         (date(2026,10,6),time(15,30),"Müşteri Beta","Fön","İptal",None),(date(2026,9,20),time(13,0),"Müşteri Gama","Boya","Geldi",None),(date(2026,9,22),time(16,0),"Müşteri Delta","Saç kesimi","Geldi",None)]
    B.put(wr,R0,rnd)
    HzA,HzB,HzC=[rg(Hz,c,NHz) for c in "ABC"]
    for i in range(NR):
        r=R0+i
        wr.cell(r,7,f'=IF(D{r}="","",IF(E{r}="Geldi",IF(F{r}="",IFERROR(INDEX({HzC},MATCH(D{r},{HzA},0)),0),F{r}),0))')
        wr.cell(r,8,f'=IF(OR(B{r}="",D{r}=""),"",ROUND((B{r}+IFERROR(INDEX({HzB},MATCH(D{r},{HzA},0)),0)/1440)*1440,0)/1440)')
        wr.cell(r,9,f'=IF(OR(A{r}="",H{r}="",E{r}="İptal"),"",IF(COUNTIFS($A${R0}:$A${R0+NR-1},A{r},$B${R0}:$B${R0+NR-1},"<"&(H{r}-0.0003),$H${R0}:$H${R0+NR-1},">"&(B{r}+0.0003),$E${R0}:$E${R0+NR-1},"<>İptal")>1,"Çakışıyor",""))')
        wr.cell(r,10,f"=IF(AND(A{r}='Gün'!$C$3,B{r}<>\"\",E{r}<>\"İptal\"),B{r}+ROW()/100000000,\"\")")
    wr.column_dimensions["J"].hidden=True; wr.freeze_panes="A5"
    B.dv(wr,"date",f"A{R0}:A{R0+NR-1}"); B.dv(wr,"list",f"C{R0}:C{R0+NR-1}","="+rg(M,"A",NM)); B.dv(wr,"list",f"D{R0}:D{R0+NR-1}","="+HzA); B.dv(wr,"list",f"E{R0}:E{R0+NR-1}","Bekliyor,Geldi,Gelmedi,İptal"); B.dv(wr,"pos",f"F{R0}:F{R0+NR-1}")
    B.dv(wr,"custom",f"B{R0}:B{R0+NR-1}",f"=AND(ISNUMBER(B{R0}),B{R0}>=0,B{R0}<1)","Saati 14:30 biçiminde gir.")
    for t,fc,tc in (("Geldi",MINT,GREEN),("Gelmedi",PINK,RED),("İptal",GREY,"667085")): B.cf_eq(wr,f"E{R0}:E{R0+NR-1}",t,fc,tc)
    B.cf_eq(wr,f"I{R0}:I{R0+NR-1}","Çakışıyor",PINK,RED)
    wg=B.sheet("Gün","Günlük Program","Tarihi yaz; o günün randevuları saat sırasıyla listelenir.",{"A":3,"B":9,"C":14,"D":26,"E":24,"F":12,"G":16},grid=False)
    wg["B3"]="Tarih"; wg["B3"].font=font(10,True); B.inp(wg,"C3",date(2026,10,5),DT)
    B.cards(wg,5,[("Randevu",f'=COUNT({RJ})',NUM),("Beklenen ciro","=SUM(H10:H39)",TL)],col0=5)
    B.head(wg,9,["Saat","Bitiş","Müşteri","Hizmet","Durum","Telefon","ücret"],start=2,h=22); B.body(wg,10,30,7,calc=range(2,9),fmt={2:"HH:MM",3:"HH:MM",8:TL},center=(2,3,6),start=2,stripe=True)
    MA,MB=rg(M,"A",NM),rg(M,"B",NM)
    for k in range(1,31):
        r=9+k; ix=f"MATCH(SMALL({RJ},{k}),{RJ},0)"
        wg.cell(r,2,f'=IFERROR(INDEX({RB},{ix}),"")')
        for col,src in ((3,RH),(4,RC),(5,RD),(6,RE)): wg.cell(r,col,f'=IF(B{r}="","",INDEX({src},{ix}))')
        wg.cell(r,7,f'=IF(D{r}="","",IFERROR(INDEX({MB},MATCH(D{r},{MA},0))&"",""))')
        wg.cell(r,8,f'=IF(B{r}="","",IF(F{r}="Gelmedi",0,IF(INDEX({rg(R,"F",NR)},{ix})="",IFERROR(INDEX({HzC},MATCH(E{r},{HzA},0)),0),INDEX({rg(R,"F",NR)},{ix}))))')
    wg.column_dimensions["H"].hidden=True
    wh=B.sheet(Hz,"Hizmetler","Verdiğin hizmetleri, süresini ve ücretini yaz.",{"A":26,"B":11,"C":13,"D":13,"E":15,"F":4})
    B.head(wh,4,["Hizmet","Süre (dakika)","Ücret (₺)","Yapılan","Ciro (₺)","y"],calc=(4,5,6)); B.body(wh,R0,NHz,6,calc=(4,5,6),fmt={2:NUM,3:TLI,4:NUM,5:TL},center=(2,))
    B.put(wh,R0,[("Saç kesimi",30,600),("Fön",30,400),("Boya",90,2500)])
    for i in range(NHz):
        r=R0+i; wh.cell(r,4,f'=IF(A{r}="","",COUNTIFS({RD},A{r},{RE},"Geldi"))'); wh.cell(r,5,f'=IF(A{r}="","",SUMIFS({RG},{RD},A{r}))'); wh.cell(r,6,f'=IF(A{r}="","",D{r}+ROW()/1000000)')
    wh.column_dimensions["F"].hidden=True; B.dv(wh,"uniq",f"A{R0}:A{R0+NHz-1}",f"=COUNTIF($A${R0}:$A${R0+NHz-1},A{R0})=1"); B.dv(wh,"pos",f"B{R0}:C{R0+NHz-1}")
    wm=B.sheet(M,"Müşteriler","Müşterilerini yaz. Randevu sayısı, harcama ve son ziyaret kendiliğinden hesaplanır.",{"A":26,"B":16,"C":30,"D":11,"E":11,"F":15,"G":13,"H":4})
    B.head(wm,4,["Ad","Telefon","Not","Randevu","Gelmedi","Toplam harcama (₺)","Son ziyaret","y"],calc=(4,5,6,7,8)); B.body(wm,R0,NM,8,calc=(4,5,6,7,8),fmt={2:"@",4:NUM,5:NUM,6:TL,7:DT},center=(7,))
    B.put(wm,R0,[("Müşteri Alfa",None,"Örnek"),("Müşteri Delta",None,"Örnek"),("Müşteri Beta",None,"Örnek"),("Müşteri Gama",None,"Örnek")])
    for i in range(NM):
        r=R0+i; mx=f'SUMPRODUCT(MAX(({RC}=A{r})*({RE}="Geldi")*{RA}))'
        wm.cell(r,4,f'=IF(A{r}="","",COUNTIFS({RC},A{r},{RE},"<>İptal"))'); wm.cell(r,5,f'=IF(A{r}="","",COUNTIFS({RC},A{r},{RE},"Gelmedi"))'); wm.cell(r,6,f'=IF(A{r}="","",SUMIFS({RG},{RC},A{r}))')
        wm.cell(r,7,f'=IF(A{r}="","",IF({mx}=0,"",{mx}))'); wm.cell(r,8,f'=IF(A{r}="","",F{r}+ROW()/1000000)')
    wm.column_dimensions["H"].hidden=True; wm.freeze_panes="B5"; B.dv(wm,"uniq",f"A{R0}:A{R0+NM-1}",f"=COUNTIF($A${R0}:$A${R0+NM-1},A{R0})=1")
    wo=B.sheet("Özet","Özet","Bu sayfada yalnızca yılı değiştir.",{"A":3,"B":14,"C":14,"D":14,"E":14,"F":16,"G":16,"H":3,"I":26,"J":16},grid=False)
    wo["B3"]="Yıl"; B.inp(wo,"C3",2026,"0")
    B.cards(wo,5,[("Yıl cirosu","=F23",TL),("Randevu","=C23",NUM),("Gelen","=D23",NUM),("Gelmeyen","=E23",NUM,1),("Gelmeme oranı","=IF(C23=0,0,E23/C23)",PCT1)])
    B.label(wo,"B9","Aylık tablo"); B.table(wo,10,2,["Ay","Randevu","Geldi","Gelmedi","Ciro (₺)"],13,fmt={3:NUM,4:NUM,5:NUM,6:TL})
    for m,ay in enumerate(months(),1):
        r=10+m; wo.cell(r,2,ay); wo.cell(r,3,"="+crange(RA,"$C$3",m,f',{RE},"<>İptal",{RE},"<>"')); wo.cell(r,4,"="+crange(RA,"$C$3",m,f',{RE},"Geldi"')); wo.cell(r,5,"="+crange(RA,"$C$3",m,f',{RE},"Gelmedi"')); wo.cell(r,6,"="+mrange(RG,RA,"$C$3",m))
    wo.cell(23,2,"TOPLAM")
    for col,L in zip(range(3,7),"CDEF"): wo.cell(23,col,f"=SUM({L}11:{L}22)"); wo.cell(23,col).font=font(10,True)
    B.label(wo,"I9","En çok yapılan 5 hizmet"); B.table(wo,10,9,["Hizmet","Yapılan"],5,fmt={10:NUM})
    HF,HDc=rg(Hz,"F",NHz),rg(Hz,"D",NHz)
    for k in range(1,6):
        r=10+k; wo.cell(r,9,f'=IFERROR(IF(LARGE({HF},{k})<1,"",INDEX({HzA},MATCH(LARGE({HF},{k}),{HF},0))),"")'); wo.cell(r,10,f'=IF(I{r}="","",INDEX({HDc},MATCH(I{r},{HzA},0)))')
    B.label(wo,"I17","En çok harcayan 5 müşteri"); B.table(wo,18,9,["Müşteri","Harcama (₺)"],5,fmt={10:TL})
    MH,MF=rg(M,"H",NM),rg(M,"F",NM)
    for k in range(1,6):
        r=18+k; wo.cell(r,9,f'=IFERROR(IF(LARGE({MH},{k})<1,"",INDEX({MA},MATCH(LARGE({MH},{k}),{MH},0))),"")'); wo.cell(r,10,f'=IF(I{r}="","",INDEX({MF},MATCH(I{r},{MA},0)))')
    chart(wo,"B25","Aylık ciro (₺)",Reference(wo,min_col=2,min_row=11,max_row=22),[Reference(wo,min_col=6,min_row=10,max_row=22)],[GREEN],w=16)
    B.basla([("NASIL KULLANILIR",None),("1. Hizmetler","Hizmetlerini, kaç dakika sürdüğünü ve ücretini yaz."),("2. Müşteriler","Müşterilerini yaz. Randevu girerken açılır listeden seçersin."),("3. Randevular","Tarih, saat, müşteri ve hizmeti gir. Bitiş saati hizmet süresinden hesaplanır. Aynı saate iki randevu yazarsan Çakışıyor uyarısı çıkar."),("4. Gün","Tarihi yaz, o günün programını saat sırasıyla gör ya da yazdır."),("5. Özet","Aylık randevu, gelen, gelmeyen ve ciro; en çok yapılan hizmetler ve en çok harcayan müşteriler."),
      (None,None),("BİLMEN GEREKENLER",None),("Ciro","Yalnızca durumu Geldi olan randevular ciroya girer. Özel ücret yazmazsan hizmetin liste ücreti kullanılır."),("Çakışma","Tek kişinin çalıştığı varsayılır. Birden fazla çalışan varsa her çalışan için ayrı dosya kullan."),("Kapasite",f"{NR} randevu, {NM} müşteri, {NHz} hizmet. Gün sayfası bir günde en çok 30 randevu gösterir.")])
    B.ai("Müşteri adlarını ve telefonlarını yapıştırmadan önce sil ya da takma adla değiştir.",[("Gelmeyenler","[Randevular sayfası] Gelmedi durumundaki randevuları say. Hangi gün ve saatlerde daha çok gelmeme oluyor? Azaltmak için iki öneri ver."),("Hatırlatma mesajı","Yarınki randevular için müşteriye gönderilecek kısa ve kibar bir hatırlatma mesajı yaz. Tarih ve saat için boşluk bırak."),("Dolu ve boş saatler","[Randevular sayfası] Haftanın günlerine ve saat dilimlerine göre randevu yoğunluğunu çıkar. En boş üç dilim için kampanya öner."),("Geri dönmeyen müşteri","[Müşteriler sayfası] Son ziyareti 60 günden eski olan müşterileri listele ve onlara gönderilecek kısa bir davet mesajı yaz."),("Fiyat kontrolü","[Hizmetler sayfası] Dakika başına ücreti hesapla. En düşük kalan hizmetler için fiyat ya da süre önerisi ver."),("Aylık özet","[Özet sayfasındaki aylık tablo] Aylara göre ciro ve gelmeme oranını üç cümleyle yorumla.")])
    fn=OUT+"ExcelKlinigi_RANDEVU_Randevu_Defteri.xlsx"; B.save(["Başla","Randevular","Gün","Özet","Hizmetler","Müşteriler","AI Komutları"],fn)
    w=load_workbook(fn,data_only=True); r_=w[R]; o=w["Özet"]; g=w["Gün"]
    tut=[r_.cell(R0+i,7).value for i in range(9)]; assert tut==[600,400,2200,0,0,0,0,2500,600],tut
    assert [r_.cell(R0+i,9).value for i in range(9)]==[None]*9,[r_.cell(R0+i,9).value for i in range(9)]
    assert o["F20"].value==3200 and o["F19"].value==3100 and o["C20"].value==6 and o["E20"].value==1, (o["F20"].value,o["F19"].value,o["C20"].value)
    assert [g.cell(10+k,4).value for k in range(5)]==["Müşteri Alfa","Müşteri Delta","Müşteri Beta","Müşteri Gama",None]
    assert g["E6"].value==4 and g["F6"].value==3200,(g["E6"].value,g["F6"].value); print("  RANDEVU doğrulandı")
if __name__=="__main__": rontgen(); randevu()
