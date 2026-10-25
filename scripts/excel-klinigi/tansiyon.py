from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.chart import BarChart, Reference
from datetime import date

GREEN="1F6F43"; MINT="E8F5EE"; INPUT="FFF9DB"; GREY="F3F4F6"; RED="B42318"; DARK="1B1F23"
F="Arial"
def font(sz=10,b=False,c=DARK,i=False): return Font(name=F,size=sz,bold=b,color=c,italic=i)
def fill(c): return PatternFill("solid",start_color=c,end_color=c)
thin=Side(style="thin",color="D0D5DD"); box=Border(left=thin,right=thin,top=thin,bottom=thin)
TL='#,##0.00 "₺";[Red]-#,##0.00 "₺";"-"'
N=300; R0=5; R1=R0+N-1
K="'Kayıtlar'!"
def rng(c): return f"{K}${c}${R0}:${c}${R1}"
KA,KB,KE,KH=rng("A"),rng("B"),rng("E"),rng("H")
BUGUN="'Özet'!$C$3"
NK=40  # kişi satırı

wb=Workbook()
# ---------- Başla ----------
ws=wb.active; ws.title="Başla"; ws.sheet_view.showGridLines=False
ws.column_dimensions["A"].width=3; ws.column_dimensions["B"].width=24; ws.column_dimensions["C"].width=80
ws.merge_cells("B2:C2"); ws["B2"]="EXCEL KLİNİĞİ"; ws["B2"].font=font(11,True,"FFFFFF"); ws["B2"].fill=fill(GREEN); ws["B2"].alignment=Alignment(vertical="center",indent=1); ws.row_dimensions[2].height=24
ws.merge_cells("B3:C3"); ws["B3"]="TANSİYON  |  Alacak Borç Takibi"; ws["B3"].font=font(20,True,GREEN); ws.row_dimensions[3].height=36
ws.merge_cells("B4:C4"); ws["B4"]="Kim sana borçlu, sen kime borçlusun, hangisi gecikti. Kaydı gir, durumu dosya söylesin."; ws["B4"].font=font(10,c="475467")
rows=[
 ("NASIL KULLANILIR",None),
 ("1. Kişiler","Müşterilerini ve borçlu olduğun kişi ya da firmaları yaz. Buraya eklediğin her ad, kayıt girerken açılır listede görünür."),
 ("2. Kayıtlar","Her alacak ya da borç için bir satır gir: tür, kişi, açıklama, belge tarihi, vade tarihi, tutar. Para geldikçe ya da ödedikçe Ödenen sütununu güncelle."),
 ("3. Özet","Hiçbir şey yazma, yalnızca bak. Toplam alacak ve borç, gecikenler, önümüzdeki 7 gün, vade tablosu ve kişi bazında durum kendiliğinden hesaplanır."),
 ("4. AI Komutları","Kayıtlarını bir yapay zeka asistanına yapıştırıp kullanabileceğin hazır komutlar; hatırlatma mesajı yazdırma dahil."),
 (None,None),
 ("DURUM NE DEMEK",None),
 ("Ödendi","Kalan tutar sıfır."),
 ("Gecikti","Vade tarihi geçmiş ve hâlâ kalan var. Gecikme sütunu kaç gün geçtiğini gösterir."),
 ("Bu hafta","Vadesi bugün ya da önümüzdeki 7 gün içinde."),
 ("Bekliyor","Vadesine 7 günden fazla var."),
 ("Vadesiz","Vade tarihi girilmemiş. Gecikme hesaplanamaz, mümkünse vade yaz."),
 (None,None),
 ("RENKLER",None),
 ("Sarı hücre","Senin dolduracağın hücre."),
 ("Beyaz / yeşil hücre","Formül ya da başlık. Değiştirme, kendiliğinden hesaplanır."),
 (None,None),
 ("ÖNEMLİ",None),
 ("Örnek kayıtlar","Kayıtlar sayfasındaki 8 satır ve Kişiler sayfasındaki 5 ad örnektir. Başlamadan önce yalnızca sarı hücrelerdeki içeriği sil; Kalan, Durum ve Gecikme sütunlarına dokunma."),
 ("Tarih",f"Durumlar bilgisayarının bugünkü tarihine göre hesaplanır ve dosyayı her açtığında güncellenir."),
 ("Kapasite",f"Kayıtlar sayfası {N} satıra, Kişiler sayfası {NK} ada kadar hesaplanır."),
 ("Kişi adı değişirse","Bir kişinin adını sonradan değiştirirsen eski kayıtlardaki ad aynı kalır. Eski kayıtları da güncelle, yoksa kişi tablosunda görünmez."),
]
r=6
for a,b in rows:
    if a and b is None:
        ws.merge_cells(start_row=r,start_column=2,end_row=r,end_column=3)
        c=ws.cell(r,2,a); c.font=font(10,True,GREEN); c.fill=fill(MINT); c.alignment=Alignment(indent=1,vertical="center"); ws.row_dimensions[r].height=20
    elif a:
        ws.cell(r,2,a).font=font(10,True); ws.cell(r,3,b).font=font(10)
        ws.cell(r,2).alignment=Alignment(vertical="top",indent=1); ws.cell(r,3).alignment=Alignment(wrap_text=True,vertical="top")
        ws.row_dimensions[r].height=30 if len(b)>82 else 18
        if a=="Sarı hücre": ws.cell(r,2).fill=fill(INPUT)
    r+=1

# ---------- Kişiler ----------
wp=wb.create_sheet("Kişiler"); wp.sheet_view.showGridLines=False
for col,w in zip("ABC",[32,20,44]): wp.column_dimensions[col].width=w
wp["A1"]="Kişiler"; wp["A1"].font=font(16,True,GREEN)
wp["A2"]="Müşteri, tedarikçi, kiracı, kim varsa. Ad sütunu zorunlu, diğerleri isteğe bağlı."; wp["A2"].font=font(9,c="475467",i=True)
for j,h in enumerate(["Kişi / Firma","Telefon","Not"],1):
    c=wp.cell(3,j,h); c.font=font(10,True,"FFFFFF"); c.fill=fill(GREEN); c.border=box; c.alignment=Alignment(indent=1)
kisiler=["Müşteri Alfa","Müşteri Beta","Müşteri Gama","Tedarikçi Alfa","Ofis Kirası"]
for i in range(NK):
    for j in range(1,4):
        c=wp.cell(4+i,j); c.font=font(10); c.fill=fill(INPUT); c.border=box; c.alignment=Alignment(indent=1)
    if i<len(kisiler): wp.cell(4+i,1,kisiler[i]); wp.cell(4+i,3,"Örnek kişi, silebilirsin")
    wp.cell(4+i,2).number_format="@"
wp.freeze_panes="A4"
PK=f"Kişiler!$A$4:$A${3+NK}"

# ---------- Kayıtlar ----------
wr=wb.create_sheet("Kayıtlar")
for col,w in zip("ABCDEFGHIJK",[10,24,30,13,13,15,15,15,12,11,28]): wr.column_dimensions[col].width=w
wr["A1"]="Kayıtlar"; wr["A1"].font=font(16,True,GREEN)
wr["A2"]="Her satıra bir alacak ya da borç. Sarı sütunları doldur; Kalan, Durum ve Gecikme kendiliğinden hesaplanır."; wr["A2"].font=font(9,c="475467",i=True)
heads=["Tür","Kişi / Firma","Açıklama","Belge Tarihi","Vade Tarihi","Tutar (₺)","Ödenen (₺)","Kalan (₺)","Durum","Gecikme (gün)","Not"]
for j,h in enumerate(heads,1):
    c=wr.cell(4,j,h); c.font=font(10,True,"FFFFFF"); c.fill=fill(GREEN); c.border=box; c.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
wr.row_dimensions[4].height=24
samples=[
 ("Alacak","Müşteri Alfa","Örnek: Eylül mal teslimi",date(2026,8,20),date(2026,9,20),18500,5000),
 ("Alacak","Müşteri Beta","Örnek: Kurulum hizmeti",date(2026,9,10),date(2026,10,8),12000,None),
 ("Alacak","Müşteri Gama","Örnek: Toptan satış",date(2026,9,25),date(2026,11,15),7400,None),
 ("Alacak","Müşteri Alfa","Örnek: Temmuz faturası",date(2026,6,15),date(2026,7,15),9000,9000),
 ("Alacak","Müşteri Beta","Örnek: Haziran bakım işi",date(2026,6,1),date(2026,6,30),4200,None),
 ("Borç","Tedarikçi Alfa","Örnek: Ürün alımı",date(2026,9,15),date(2026,10,10),22000,10000),
 ("Borç","Tedarikçi Alfa","Örnek: Ağustos alımı",date(2026,8,10),date(2026,9,28),6500,None),
 ("Borç","Ofis Kirası","Örnek: Kasım kirası",date(2026,10,1),date(2026,11,1),12000,None),
]
for i in range(N):
    r=R0+i
    vals=samples[i] if i<len(samples) else (None,)*7
    for j,v in enumerate(vals,1): wr.cell(r,j,v)
    wr.cell(r,8,f'=IF(F{r}="","",F{r}-G{r})')
    wr.cell(r,9,f'=IF(F{r}="","",IF(H{r}<=0,"Ödendi",IF(E{r}="","Vadesiz",IF(E{r}<{BUGUN},"Gecikti",IF(E{r}<={BUGUN}+7,"Bu hafta","Bekliyor")))))')
    wr.cell(r,10,f'=IF(I{r}="Gecikti",{BUGUN}-E{r},"")')
    for j in range(1,12):
        c=wr.cell(r,j); c.font=font(10); c.border=box
        if j in (8,9,10): c.fill=fill("FFFFFF")
        else: c.fill=fill(INPUT)
    for j in (4,5): wr.cell(r,j).number_format="DD.MM.YYYY"; wr.cell(r,j).alignment=Alignment(horizontal="center")
    for j in (6,7): wr.cell(r,j).number_format='#,##0.00 "₺"'
    wr.cell(r,8).number_format=TL; wr.cell(r,8).font=font(10,True)
    wr.cell(r,1).alignment=Alignment(horizontal="center"); wr.cell(r,9).alignment=Alignment(horizontal="center"); wr.cell(r,10).alignment=Alignment(horizontal="center"); wr.cell(r,10).number_format="0"
wr.freeze_panes="C5"
def dv_add(dv,ref,title,msg):
    dv.errorTitle=title; dv.error=msg; dv.showErrorMessage=True; wr.add_data_validation(dv); dv.add(ref)
dv_add(DataValidation(type="list",formula1='"Alacak,Borç"',allow_blank=True),f"A{R0}:A{R1}","Tür","Listeden Alacak ya da Borç seç.")
dv_add(DataValidation(type="list",formula1=f"={PK}",allow_blank=True),f"B{R0}:B{R1}","Kişi","Kişiyi listeden seç. Yeni kişi için önce Kişiler sayfasına ekle.")
dv_add(DataValidation(type="date",operator="greaterThan",formula1="36526",allow_blank=True),f"D{R0}:E{R1}","Tarih","Geçerli bir tarih gir (örnek: 04.10.2026).")
dv_add(DataValidation(type="decimal",operator="greaterThan",formula1="0",allow_blank=True),f"F{R0}:F{R1}","Tutar","Tutarı sıfırdan büyük bir sayı olarak gir.")
dv_add(DataValidation(type="custom",formula1=f"=AND(ISNUMBER(G{R0}),G{R0}>=0,G{R0}<=F{R0})",allow_blank=True),f"G{R0}:G{R1}","Ödenen","Ödenen tutar sıfırdan küçük ya da toplam tutardan büyük olamaz.")
cf=wr.conditional_formatting
cf.add(f"I{R0}:I{R1}",CellIsRule(operator="equal",formula=['"Gecikti"'],fill=fill("FDE2E0"),font=Font(name=F,bold=True,color=RED)))
cf.add(f"I{R0}:I{R1}",CellIsRule(operator="equal",formula=['"Bu hafta"'],fill=fill("FEF0C7"),font=Font(name=F,bold=True,color="93370D")))
cf.add(f"I{R0}:I{R1}",CellIsRule(operator="equal",formula=['"Ödendi"'],fill=fill(MINT),font=Font(name=F,bold=True,color=GREEN)))
cf.add(f"A{R0}:A{R1}",CellIsRule(operator="equal",formula=['"Alacak"'],font=Font(name=F,bold=True,color=GREEN)))
cf.add(f"A{R0}:A{R1}",CellIsRule(operator="equal",formula=['"Borç"'],font=Font(name=F,bold=True,color=RED)))
cf.add(f"A{R0}:B{R1}",FormulaRule(formula=[f'AND($F{R0}<>"",OR($A{R0}="",$B{R0}=""))'],fill=fill("FDE2E0")))

# ---------- Özet ----------
wo=wb.create_sheet("Özet"); wo.sheet_view.showGridLines=False
for k,v in {"A":3,"B":22,"C":18,"D":18,"E":18,"F":18,"G":3,"H":28,"I":16,"J":17,"K":16,"L":16}.items(): wo.column_dimensions[k].width=v
wo["B1"]="Özet"; wo["B1"].font=font(16,True,GREEN)
wo["B3"]="Bugün"; wo["B3"].font=font(10,True); wo["B3"].alignment=Alignment(horizontal="right",indent=1)
wo["C3"]="=TODAY()"; wo["C3"].number_format="DD.MM.YYYY"; wo["C3"].font=font(12,True); wo["C3"].fill=fill(MINT); wo["C3"].border=box; wo["C3"].alignment=Alignment(horizontal="center")
wo["D3"]="← Kendiliğinden güncellenir, dokunma"; wo["D3"].font=font(9,c="475467",i=True)
cards=[("B","Toplam Alacak",f'=SUMIFS({KH},{KA},"Alacak")'),
       ("C","Geciken Alacak",f'=SUMIFS({KH},{KA},"Alacak",{KE},"<"&$C$3)'),
       ("D","Toplam Borç",f'=SUMIFS({KH},{KA},"Borç")'),
       ("E","Geciken Borç",f'=SUMIFS({KH},{KA},"Borç",{KE},"<"&$C$3)'),
       ("F","Net (Alacak − Borç)","=B6-D6")]
for col,lab,f in cards:
    a=wo[f"{col}5"]; a.value=lab; a.font=font(9,True,"FFFFFF"); a.fill=fill(GREEN); a.alignment=Alignment(horizontal="center"); a.border=box
    b=wo[f"{col}6"]; b.value=f; b.font=font(12,True,RED if "Geciken" in lab else DARK); b.fill=fill(MINT); b.alignment=Alignment(horizontal="center",vertical="center"); b.number_format=TL; b.border=box
wo.row_dimensions[6].height=28
wo["B8"]="Önümüzdeki 7 gün (bugün dahil)"; wo["B8"].font=font(11,True,GREEN)
for col,lab,tur in (("B","Tahsil edilecek","Alacak"),("D","Ödenecek","Borç")):
    a=wo[f"{col}9"]; a.value=lab; a.font=font(10,True); a.fill=fill(GREY); a.border=box; a.alignment=Alignment(indent=1)
    nxt=chr(ord(col)+1); b=wo[f"{nxt}9"]; b.value=f'=SUMIFS({KH},{KA},"{tur}",{KE},">="&$C$3,{KE},"<="&$C$3+7)'; b.font=font(10,True); b.number_format=TL; b.border=box
wo["B11"]="Vade Tablosu (kalan tutarlar)"; wo["B11"].font=font(11,True,GREEN)
for j,h in enumerate(["Vade Durumu","Alacak","Borç"]):
    c=wo.cell(12,2+j,h); c.font=font(10,True,"FFFFFF"); c.fill=fill(GREEN); c.border=box; c.alignment=Alignment(horizontal="center")
def bucket(tur,lo,hi):
    # lo/hi: gecikme günü sınırları (dahil); None = sınırsız
    parts=[KH,KA,f'"{tur}"']
    if lo is not None: parts+= [KE,f'"<="&$C$3-{lo}']
    if hi is not None: parts+= [KE,f'">="&$C$3-{hi}']
    return "=SUMIFS("+",".join(parts)+")"
buckets=[("Vadesi gelmemiş",None,0),("1-30 gün gecikmiş",1,30),("31-60 gün gecikmiş",31,60),("61-90 gün gecikmiş",61,90),("90 günden fazla",91,None)]
for i,(lab,lo,hi) in enumerate(buckets):
    r=13+i; wo.cell(r,2,lab); wo.cell(r,3,bucket("Alacak",lo,hi)); wo.cell(r,4,bucket("Borç",lo,hi))
wo.cell(18,2,"Vadesiz"); wo.cell(18,3,"=C19-SUM(C13:C17)"); wo.cell(18,4,"=D19-SUM(D13:D17)")
wo.cell(19,2,"TOPLAM"); wo.cell(19,3,"=B6"); wo.cell(19,4,"=D6")
for r in range(13,20):
    for col in (2,3,4):
        c=wo.cell(r,col); c.border=box; c.font=font(10,r==19)
        if col>2: c.number_format=TL
        if r==19: c.fill=fill(MINT)
        elif r%2==0: c.fill=fill(GREY)
    wo.cell(r,2).alignment=Alignment(indent=1)
wo["H11"]="Kişi Bazında Durum"; wo["H11"].font=font(11,True,GREEN)
for j,h in enumerate(["Kişi / Firma","Alacak","Geciken Alacak","Borç","Net"]):
    c=wo.cell(12,8+j,h); c.font=font(10,True,"FFFFFF"); c.fill=fill(GREEN); c.border=box; c.alignment=Alignment(horizontal="center")
for i in range(NK):
    r=13+i
    wo.cell(r,8,f'=IF(Kişiler!A{4+i}="","",Kişiler!A{4+i})')
    wo.cell(r,9,f'=IF(H{r}="","",SUMIFS({KH},{KA},"Alacak",{KB},H{r}))')
    wo.cell(r,10,f'=IF(H{r}="","",SUMIFS({KH},{KA},"Alacak",{KB},H{r},{KE},"<"&$C$3))')
    wo.cell(r,11,f'=IF(H{r}="","",SUMIFS({KH},{KA},"Borç",{KB},H{r}))')
    wo.cell(r,12,f'=IF(H{r}="","",I{r}-K{r})')
    for col in range(8,13):
        c=wo.cell(r,col); c.font=font(10); c.border=box
        if col>8: c.number_format=TL
        if i%2==1: c.fill=fill(GREY)
    wo.cell(r,8).alignment=Alignment(indent=1)
wo.conditional_formatting.add(f"J13:J{12+NK}",CellIsRule(operator="greaterThan",formula=["0"],font=Font(name=F,bold=True,color=RED)))
wo.conditional_formatting.add(f"L13:L{12+NK}",CellIsRule(operator="lessThan",formula=["0"],font=Font(name=F,bold=True,color=RED)))
wo.conditional_formatting.add("F6",CellIsRule(operator="lessThan",formula=["0"],font=Font(name=F,size=12,bold=True,color=RED)))
ch=BarChart(); ch.type="col"; ch.title="Vadeye Göre Alacak ve Borç"; ch.height=7.5; ch.width=17
ch.add_data(Reference(wo,min_col=3,max_col=4,min_row=12,max_row=18),titles_from_data=True)
ch.set_categories(Reference(wo,min_col=2,min_row=13,max_row=18))
ch.series[0].graphicalProperties.solidFill=GREEN; ch.series[1].graphicalProperties.solidFill="D92D20"
ch.legend.position="b"; ch.y_axis.numFmt='#,##0'; ch.y_axis.majorGridlines=None; ch.x_axis.delete=False; ch.y_axis.delete=False
wo.add_chart(ch,"B21")

# ---------- AI Komutları ----------
wa=wb.create_sheet("AI Komutları"); wa.sheet_view.showGridLines=False
wa.column_dimensions["A"].width=3; wa.column_dimensions["B"].width=26; wa.column_dimensions["C"].width=94
wa["B1"]="Yapay Zeka Komutları"; wa["B1"].font=font(16,True,GREEN)
wa.merge_cells("B2:C2"); wa["B2"]="Kayıtlar sayfasında başlık satırıyla birlikte kayıtlarını seç, kopyala, kullandığın yapay zeka asistanına yapıştır ve altına aşağıdaki komutlardan birini ekle."
wa["B2"].font=font(10,c="475467"); wa["B2"].alignment=Alignment(wrap_text=True,vertical="top"); wa.row_dimensions[2].height=30
wa.merge_cells("B3:C3"); wa["B3"]="Dikkat: Bu sayfada gerçek kişi ve firma adları var. Yapıştırmadan önce adları Müşteri Alfa, Müşteri Beta gibi takma adlarla değiştir; telefon ve IBAN paylaşma. Yapay zekanın verdiği rakamları Özet sayfasıyla karşılaştır."
wa["B3"].font=font(10,True,RED); wa["B3"].alignment=Alignment(wrap_text=True,vertical="top"); wa.row_dimensions[3].height=44
for j,h in enumerate(["Ne için","Komut (kopyala, yapıştır)"]):
    c=wa.cell(5,2+j,h); c.font=font(10,True,"FFFFFF"); c.fill=fill(GREEN); c.border=box; c.alignment=Alignment(indent=1)
prompts=[
 ("Genel durum","Yukarıdaki tablo benim alacak ve borç kayıtlarım. Toplam alacağımı, toplam borcumu ve gecikenleri hesapla, durumumu üç cümleyle özetle."),
 ("Kibar hatırlatma","Durumu Gecikti olan her alacak için müşteriye gönderebileceğim kısa ve kibar bir hatırlatma mesajı yaz. Mesajda tutar ve vade tarihi geçsin, suçlayıcı olmasın."),
 ("Sert hatırlatma","60 günden fazla gecikmiş alacaklar için daha net ama saygılı ikinci bir hatırlatma mesajı yaz. Son ödeme için bir tarih iste."),
 ("Kimi önce aramalıyım","Geciken alacakları tutar ve gecikme gününe göre sırala. Bugün ilk hangi üç kişiyle konuşmam gerektiğini nedenleriyle söyle."),
 ("Ödeme sırası","Borçlarımı vade tarihine göre sırala ve önümüzdeki 30 gün için hafta hafta bir ödeme planı çıkar."),
 ("Nakit dengesi","Önümüzdeki 30 günde tahsil etmeyi beklediğim tutarla ödemem gereken tutarı hafta hafta karşılaştır. Açık veren hafta varsa belirt."),
 ("Riskli müşteri","Hangi müşteriler sürekli geç ödüyor ya da birden fazla gecikmiş kaydı var? Yalnızca tablodaki verilere dayanarak listele."),
 ("Taksit önerisi","En büyük gecikmiş alacak için müşteriye sunabileceğim iki farklı taksit planı öner ve bunu anlatan kısa bir mesaj yaz."),
 ("Haftalık rapor","Bu tablodan bana tek sayfalık haftalık özet hazırla: tahsil edilenler, gecikenler, bu hafta vadesi gelenler."),
 ("Kayıt kontrolü","Bu kayıtlarda eksik ya da tutarsız görünen satır var mı? Vadesi belge tarihinden önce olan, tutarı boş kalan ya da iki kez girilmiş gibi duranları göster."),
]
for i,(a,b) in enumerate(prompts):
    r=6+i
    wa.cell(r,2,a).font=font(10,True); wa.cell(r,3,b).font=font(10)
    for col in (2,3):
        c=wa.cell(r,col); c.border=box; c.alignment=Alignment(wrap_text=True,vertical="center",indent=1)
        if i%2==1: c.fill=fill(GREY)
    wa.row_dimensions[r].height=32

wb._sheets=[wb["Başla"],wb["Kayıtlar"],wb["Özet"],wb["Kişiler"],wb["AI Komutları"]]
for s in wb.worksheets:
    s.sheet_properties.tabColor=GREEN
    s.page_setup.orientation="landscape"; s.page_setup.fitToWidth=1; s.page_setup.fitToHeight=0
    s.sheet_properties.pageSetUpPr.fitToPage=True
wb.active=0
wb.save("ExcelKlinigi_TANSIYON_Alacak_Borc_Takibi.xlsx")
