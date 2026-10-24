# RİTİM (alışkanlık), ACİL (borç kapatma), AMELİYAT (proje planı)
from lib import *
from datetime import date, timedelta
from openpyxl import load_workbook

def ritim():
    B=Book("RİTİM","Alışkanlık takibi","Her gün bir işaret koy; serini, haftalık hedefini ve aylık gidişini dosya saysın.")
    NA=15; T="Takip"; S="Seri"; A="Alışkanlıklar"; D0=6; ND=366; DL=D0+ND-1
    wa=B.sheet(A,"Alışkanlıklar","Yılı ve takip edeceğin alışkanlıkları yaz. Gerisi Takip sayfasındaki işaretlerinden hesaplanır.",{"A":28,"B":13,"C":11,"D":10,"E":12,"F":12,"G":11,"H":11})
    wa["J3"]="Yıl"; wa["J3"].font=font(10,True); B.inp(wa,"K3",2026,"0"); Y=f"'{A}'!$K$3"
    B.head(wa,4,["Alışkanlık","Haftalık hedef (gün)","Toplam gün","Bu ay","Bu hafta","Son 30 gün","Şimdiki seri","En uzun seri"],calc=range(3,9))
    B.body(wa,R0,NA,8,calc=range(3,9),fmt={2:NUM,3:NUM,4:NUM,6:"0%",7:NUM,8:NUM},center=range(2,9),bold=(7,))
    al=[("Yürüyüş (30 dk)",5),("Kitap okuma (20 sayfa)",7),("İngilizce çalışma",4),("Su (2 litre)",7)]
    B.put(wa,R0,al); B.dv(wa,"int",f"B{R0}:B{R0+NA-1}",(1,7))
    wt=B.sheet(T,"Takip","Yaptığın günün hücresine 1 yaz. Bugünün satırı yeşil çerçevelidir.",{"A":12,"B":11,**{CL(c):13 for c in range(3,3+NA)},"R":4,"S":4})
    wt.cell(4,1,"Tarih"); wt.cell(4,2,"Gün"); B.head(wt,4,["Tarih","Gün"]+[""]*NA+["y","y"],calc=range(1,3+NA+2),h=34)
    for j in range(NA): wt.cell(4,3+j,f"=IF('{A}'!A{R0+j}=\"\",\"\",'{A}'!A{R0+j})")
    marks={0:lambda k:k%7 not in (5,),1:lambda k:True,2:lambda k:k%7 in (0,2,4),3:lambda k:k%3!=0}
    base=date(2026,1,1)
    for i in range(ND):
        r=D0+i
        wt.cell(r,1,f"=DATE({Y},1,1)" if i==0 else (f"=A{r-1}+1" if i<365 else f'=IF(YEAR(A{r-1}+1)={Y},A{r-1}+1,"")'))
        wt.cell(r,2,f'=IF(A{r}="","",CHOOSE(WEEKDAY(A{r},2),"Pzt","Sal","Çar","Per","Cum","Cmt","Paz"))'); wt.cell(r,18,f'=IF(A{r}="",0,A{r})'); wt.cell(r,19,f'=IF(A{r}="",0,MONTH(A{r}))')
        for j in range(NA):
            c=wt.cell(r,3+j); c.fill=fill(INPUT); c.border=box; c.alignment=Alignment(horizontal="center"); c.font=font(10,True)
        wt.cell(r,1).number_format=DT; wt.cell(r,1).font=font(10); wt.cell(r,2).font=font(9,c="475467")
        d=base+timedelta(i)
        if date(2026,9,7)<=d<=date(2026,10,4):
            k=(d-date(2026,9,7)).days
            for j,f in marks.items():
                if f(k): wt.cell(r,3+j,1)
    wt.row_dimensions[5].height=4; wt.column_dimensions["R"].hidden=True; wt.column_dimensions["S"].hidden=True; wt.freeze_panes=f"C{D0}"
    B.dv(wt,"int",f"C{D0}:{CL(2+NA)}{DL}",(1,1),"Yaptıysan 1 yaz, yapmadıysan boş bırak.")
    B.cf_f(wt,f"A{D0}:{CL(2+NA)}{DL}",f"$R{D0}=TODAY()",MINT,GREEN); B.cf_f(wt,f"C{D0}:{CL(2+NA)}{DL}",f'C{D0}=1',"B7E4C7",GREEN)
    ws=B.sheet(S,"Seri (dokunma)","Art arda gün sayısının ara hesabı. Bu sayfayı değiştirme.",{"A":12})
    for i in range(ND):
        r=D0+i; ws.cell(r,1,f"='{T}'!R{r}")
        for j in range(NA):
            L=CL(3+j); ws.cell(r,3+j,f"=IF('{T}'!{L}{r}=\"\",0,1)" if i==0 else f"=IF('{T}'!{L}{r}=\"\",0,{L}{r-1}+1)")
    DN=f"'{T}'!$R${D0}:$R${DL}"; MN=f"'{T}'!$S${D0}:$S${DL}"
    for j in range(NA):
        r=R0+j; L=CL(3+j); col=f"'{T}'!${L}${D0}:${L}${DL}"; sc=f"'{S}'!${L}${D0}:${L}${DL}"
        wa.cell(r,3,f'=IF(A{r}="","",COUNT({col}))'); wa.cell(r,4,f'=IF(A{r}="","",SUMPRODUCT(({MN}=MONTH(TODAY()))*({DN}<=TODAY())*({col}<>"")))')
        wa.cell(r,5,f'=IF(A{r}="","",SUMPRODUCT(({DN}>=TODAY()-WEEKDAY(TODAY(),2)+1)*({DN}<=TODAY())*({col}<>""))&" / "&B{r})')
        wa.cell(r,6,f'=IF(A{r}="","",SUMPRODUCT(({DN}>TODAY()-30)*({DN}<=TODAY())*({col}<>""))/30)')
        ti=f"MATCH(TODAY(),{DN},0)"
        wa.cell(r,7,f'=IF(A{r}="","",IFERROR(IF(INDEX({col},{ti})<>"",INDEX({sc},{ti}),IF({ti}=1,0,INDEX({sc},{ti}-1))),0))'); wa.cell(r,8,f'=IF(A{r}="","",MAX({sc}))')
    wo=B.sheet("Aylar","Aylara Göre","Her alışkanlığı her ay kaç gün yaptığın.",{"A":28,**{CL(c):9 for c in range(2,14)},"N":10})
    B.head(wo,4,["Alışkanlık"]+[m[:3] for m in months()]+["Toplam"],calc=range(1,15),h=22); B.body(wo,R0,NA,14,calc=range(1,15),fmt={c:NUM for c in range(2,15)},center=range(2,15),stripe=True)
    for j in range(NA):
        r=R0+j; L=CL(3+j); col=f"'{T}'!${L}${D0}:${L}${DL}"; wo.cell(r,1,f"=IF('{A}'!A{r}=\"\",\"\",'{A}'!A{r})")
        for m in range(1,13): wo.cell(r,1+m,f'=IF($A{r}="","",SUMPRODUCT(({MN}={m})*({col}<>"")))')
        wo.cell(r,14,f'=IF(A{r}="","",SUM(B{r}:M{r}))')
    B.basla([("NASIL KULLANILIR",None),("1. Alışkanlıklar","Yılı, takip edeceğin alışkanlıkları ve haftada kaç gün yapmayı hedeflediğini yaz."),("2. Takip","Her gün, yaptığın alışkanlığın hücresine 1 yaz. Bugünün satırı yeşil görünür."),("3. Alışkanlıklar sayfasına dön","Toplam gün, bu ay, bu hafta ile hedef, son 30 gün oranı, şimdiki seri ve en uzun seri hesaplanır."),("4. Aylar","Her alışkanlığı her ay kaç gün yaptığını gösterir."),
      (None,None),("BİLMEN GEREKENLER",None),("Seri","Art arda işaretlenen gün sayısıdır. Bugünü henüz işaretlemediysen dünkü seri gösterilir; dün de boşsa seri sıfırlanır."),("Hafta","Hafta pazartesi başlar."),("Tek yıl","Bir dosya bir yılı tutar. Yeni yıl için boş bir kopyayla başla."),("Kapasite",f"{NA} alışkanlık, bir yıl.")])
    B.ai("Sağlıkla ilgili özel bilgileri paylaşmak zorunda değilsin; alışkanlık adlarını genel tutabilirsin.",[("Gidiş yorumu","[Alışkanlıklar sayfası] Hedefime göre hangi alışkanlıkta iyiyim, hangisinde geride kalıyorum? Kısa ve dürüst yorumla."),("Kopma günleri","[Takip sayfası] Serilerimin en çok hangi günlerde bozulduğunu bul. O günler için küçük bir önlem öner."),("Hedef ayarı","[Alışkanlıklar sayfası] Son 30 gün oranıma bakarak haftalık hedeflerimi gerçekçi olacak şekilde yeniden öner."),("Küçük adım","[Alışkanlık adı] alışkanlığını sürdüremiyorum. İki dakikada yapılabilecek daha küçük bir sürümünü öner."),("Aylık özet","[Aylar sayfası] Aylara göre gidişimi üç cümleyle özetle ve gelecek ay için tek bir odak öner."),("Hatırlatma planı","Bu alışkanlıklar için günün hangi saatine hangi hatırlatmayı koymalıyım? Sabah, öğle ve akşam olarak planla.")])
    fn=OUT+"ExcelKlinigi_RITIM_Aliskanlik_Takibi.xlsx"; B.save(["Başla",T,A,"Aylar",S,"AI Komutları"],fn)
    w=load_workbook(fn,data_only=True); a=w[A]; o=w["Aylar"]; today=date(2026,10,4)
    assert w[T].cell(D0,1).value.date()==date(2026,1,1) and w[T].cell(DL,1).value is None and w[T].cell(DL-1,1).value.date()==date(2026,12,31)
    days=[date(2026,9,7)+timedelta(k) for k in range(28)]
    for j,f in marks.items():
        done=[d for k,d in enumerate(days) if f(k)]; r=R0+j
        cur=0; d=today
        if d not in done: d-=timedelta(1)
        while d in done: cur+=1; d-=timedelta(1)
        best=0; run=0
        for k,dd in enumerate(days):
            run=run+1 if f(k) else 0; best=max(best,run)
        wk=sum(1 for d in done if today-timedelta(today.isoweekday()-1)<=d<=today)
        exp=(len(done),sum(1 for d in done if d.month==10),f"{wk} / {al[j][1]}",cur,best)
        got=(a.cell(r,3).value,a.cell(r,4).value,a.cell(r,5).value,a.cell(r,7).value,a.cell(r,8).value); assert got==exp,(j,got,exp)
        assert abs(a.cell(r,6).value-len([d for d in done if today-timedelta(30)<d<=today])/30)<1e-9
        assert o.cell(r,10).value==sum(1 for d in done if d.month==9) and o.cell(r,11).value==sum(1 for d in done if d.month==10)
    print("  RİTİM doğrulandı")

def acil():
    B=Book("ACİL","Borç kapatma planı","Borçlarını yaz, her ay ayırabileceğin tutarı gir; hangi borcu önce kapatacağını, ne zaman borçsuz kalacağını ve ne kadar faiz ödeyeceğini gör.")
    NB,NT=10,120; Bs="Borçlar"; P0=8
    wb_=B.sheet(Bs,"Borçlar","Her borcun kalan tutarını, aylık faiz oranını ve asgari aylık ödemesini yaz.",{"A":26,"B":16,"C":13,"D":16,"E":10,"F":15,"G":17,"H":30,"I":3,"J":34,"K":16,"L":4})
    B.head(wb_,4,["Borç","Kalan borç (₺)","Aylık faiz oranı","Asgari aylık ödeme (₺)","Öncelik","Kapanış ayı","Ödenecek faiz (₺)","Kontrol","","","","y"],calc=(5,6,7,8,12))
    for c in (9,10,11): wb_.cell(4,c).fill=fill(WHITE); wb_.cell(4,c).border=Border()
    B.body(wb_,R0,NB,8,calc=(5,6,7,8),fmt={2:TLI,3:"0.00%",4:TLI,5:NUM,6:"MMMM YYYY",7:TL},center=(3,5,6))
    borc=[("Kredi kartı",60000,0.045,3000),("İhtiyaç kredisi",120000,0.035,7500),("Kredili hesap",25000,0.05,1500),("Aileden borç",40000,0.0,2000)]
    B.put(wb_,R0,borc)
    for r,lab,val,nf in ((5,"Asgarilerin üstüne ayırabileceğin aylık tutar (₺)",5000,TLI),(6,"Yöntem",None,None),(7,"Planın başladığı ay",date(2026,10,1),"MMMM YYYY")):
        a=wb_.cell(r,10,lab); a.font=font(10,True); a.fill=fill(GREY); a.border=box
        if val is not None: B.inp(wb_,f"K{r}",val,nf)
    B.inp(wb_,"K6","Çığ"); B.dv(wb_,"list","K6","Çığ,Kartopu"); B.dv(wb_,"nonneg","K5"); B.dv(wb_,"date","K7")
    a=wb_.cell(8,10,"Toplam aylık ödeme gücün (₺)"); a.font=font(10,True); a.fill=fill(GREY); a.border=box; c=wb_.cell(8,11,f"=SUM(D{R0}:D{R0+NB-1})+K5"); c.number_format=TL; c.font=font(11,True); c.border=box
    wb_["J10"]="Çığ: faizi en yüksek borca yüklenir, toplamda en az faizi ödetir."; wb_["J11"]="Kartopu: en küçük borca yüklenir, ilk borç daha çabuk kapanır."
    for rr in (10,11): wb_.cell(rr,10).font=font(9,c="475467",i=True)
    B.dv(wb_,"nonneg",f"B{R0}:B{R0+NB-1}"); B.dv(wb_,"pct",f"C{R0}:C{R0+NB-1}"); B.dv(wb_,"nonneg",f"D{R0}:D{R0+NB-1}"); B.dv(wb_,"uniq",f"A{R0}:A{R0+NB-1}",f"=COUNTIF($A${R0}:$A${R0+NB-1},A{R0})=1")
    KEY=f"$L${R0}:$L${R0+NB-1}"
    for i in range(NB):
        r=R0+i; wb_.cell(r,12,f'=IF(OR(A{r}="",B{r}="",B{r}=0),9E+15,IF($K$6="Kartopu",B{r},-C{r}*1000000)+ROW()/1000)')
        wb_.cell(r,5,f'=IF(L{r}>8E+15,"",SUMPRODUCT(--({KEY}<L{r}))+1)')
        wb_.cell(r,6,f"=IF(E{r}=\"\",\"\",INDEX('Plan'!$D$4:$M$4,E{r}))"); wb_.cell(r,7,f"=IF(E{r}=\"\",\"\",INDEX('Plan'!$D$5:$M$5,E{r}))")
        wb_.cell(r,8,f'=IF(E{r}="","",IF(D{r}<=B{r}*C{r},"Asgari ödeme faizi karşılamıyor",IF(F{r}="Bitmiyor","{NT} ayda kapanmıyor","OK")))')
    wb_.column_dimensions["L"].hidden=True; B.cf_f(wb_,f"H{R0}:H{R0+NB-1}",f'AND(H{R0}<>"",H{R0}<>"OK")',PINK,RED)
    def plan(name,title,budget):
        ws=B.sheet(name,title,"Ara hesap sayfası. Borçlar öncelik sırasına göre soldan sağa dizilir. Değiştirme.",{"A":6,"B":13,"C":14})
        E=f"'{Bs}'!$E${R0}:$E${R0+NB-1}"
        for lab,r in (("Borç",1),("Açılış bakiyesi",2),("Aylık faiz",3),("Kapanış ayı",4),("Toplam faiz",5),("Asgari ödeme",6)): pass
        ws["A1"]=title; ws["C3"]="Öncelik"; ws["C3"].font=font(9,True)
        labels={2:"Borç",3:None,4:"Kapanış ayı",5:"Toplam faiz (₺)",6:"Açılış bakiyesi",7:"Faiz / asgari"}
        ws["A2"]=""; 
        # özet satırları: 2 ad, 4 kapanış, 5 faiz ; yardımcı: bakiye, oran, asgari uzak sütunlarda
        HC=14  # N sütunundan itibaren hesap blokları
        for k in range(1,NB+1):
            col=3+k; L=CL(col); ix=f"MATCH({k},{E},0)"
            ws.cell(3,col,k).font=font(9,True)
            ws.cell(2,col,f"=IFERROR(INDEX('{Bs}'!$A${R0}:$A${R0+NB-1},{ix})&\"\",\"\")").font=font(9,True)
            ws.cell(6,col,f"=IFERROR(INDEX('{Bs}'!$B${R0}:$B${R0+NB-1},{ix}),0)"); ws.cell(7,col,f"=IFERROR(INDEX('{Bs}'!$C${R0}:$C${R0+NB-1},{ix}),0)"); ws.cell(8,col,f"=IFERROR(INDEX('{Bs}'!$D${R0}:$D${R0+NB-1},{ix}),0)")
            b=HC+(k-1)*4; N_,M_,X_,Bn=CL(b),CL(b+1),CL(b+2),CL(b+3)
            for t in range(1,NT+1):
                r=P0+2+t; prev=f"{L}$6" if t==1 else f"{Bn}{r-1}"
                ws.cell(r,b,f"={prev}*(1+{L}$7)"); ws.cell(r,b+1,f"=MIN({N_}{r},{L}$8)")
                prevx="0" if k==1 else "+".join(f"{CL(HC+(j-1)*4+2)}{r}" for j in range(1,k))
                ws.cell(r,b+2,f"=MAX(0,MIN({N_}{r}-{M_}{r},$C{r}-({prevx})))"); ws.cell(r,b+3,f"=ROUND({N_}{r}-{M_}{r}-{X_}{r},2)")
            ws.cell(P0+2,b,f"=\"Borç {k}: toplam\""); ws.cell(P0+2,b+1,"asgari"); ws.cell(P0+2,b+2,"ek"); ws.cell(P0+2,b+3,"bakiye")
            r1,r2=P0+3,P0+2+NT
            ws.cell(4,col,f'=IF({L}6=0,"",IFERROR(INDEX($B${r1}:$B${r2},MATCH(0,{Bn}{r1}:{Bn}{r2},0)),"Bitmiyor"))').number_format="MMMM YYYY"
            ws.cell(5,col,f"=IF({L}6=0,0,SUM({N_}{r1}:{N_}{r2})-{L}6-(SUM({Bn}{r1}:{Bn}{r2})-{Bn}{r2}))").number_format=TL
        for r,lab in ((2,"Borç"),(4,"Kapanış ayı"),(5,"Toplam faiz (₺)"),(6,"Açılış bakiyesi"),(7,"Aylık faiz"),(8,"Asgari ödeme")): ws.cell(r,2,lab).font=font(9,True)
        B.head(ws,P0+2,["Ay","Tarih","Dağıtılacak ek (₺)","Toplam ödeme (₺)","O ayın faizi (₺)","Kalan borç (₺)"],calc=range(1,7),h=30)
        Ms=[CL(HC+(j-1)*4+1) for j in range(1,NB+1)]; Xs=[CL(HC+(j-1)*4+2) for j in range(1,NB+1)]; Ns=[CL(HC+(j-1)*4) for j in range(1,NB+1)]; Bns=[CL(HC+(j-1)*4+3) for j in range(1,NB+1)]
        for t in range(1,NT+1):
            r=P0+2+t; ws.cell(r,1,t); ws.cell(r,2,f"=EDATE('{Bs}'!$K$7,{t-1})").number_format="MMM YYYY"
            ws.cell(r,3,f"=MAX(0,{budget}-({'+'.join(m+str(r) for m in Ms)}))")
            ws.cell(r,4,"="+"+".join(f"{m}{r}+{x}{r}" for m,x in zip(Ms,Xs)))
            prevsum="+".join(f"{CL(3+k)}$6" for k in range(1,NB+1)) if t==1 else f"F{r-1}"
            ws.cell(r,5,f"=({'+'.join(n+str(r) for n in Ns)})-({prevsum})"); ws.cell(r,6,"="+"+".join(b+str(r) for b in Bns))
            for c in (3,4,5,6): ws.cell(r,c).number_format=TL
        ws.freeze_panes=f"C{P0+3}"; return ws
    plan("Plan","Plan (ek ödemeyle)",f"'{Bs}'!$K$8"); plan("Yalnız Asgari","Karşılaştırma (yalnız asgari ödemeyle)",f"SUM('{Bs}'!$D${R0}:$D${R0+NB-1})")
    r1,r2=P0+3,P0+2+NT
    wo=B.sheet("Özet","Özet","Bu sayfaya yazma. Rakamları Borçlar sayfasından değiştir.",{"A":3,"B":30,"C":20,"D":20,"E":20},grid=False)
    def end(sh): return f"IF(SUM('{Bs}'!$B${R0}:$B${R0+NB-1})=0,\"\",IFERROR(INDEX('{sh}'!$B${r1}:$B${r2},MATCH(0,'{sh}'!$F${r1}:$F${r2},0)),\"{NT} ayda bitmiyor\"))"
    def ay(sh): return f"IFERROR(MATCH(0,'{sh}'!$F${r1}:$F${r2},0),\"\")"
    B.cards(wo,4,[("Toplam borç",f"=SUM('{Bs}'!$B${R0}:$B${R0+NB-1})",TL),("Aylık ödeme gücü",f"='{Bs}'!K8",TL),("Borçsuz kalacağın ay","="+end("Plan"),"MMMM YYYY"),("Ödeyeceğin toplam faiz",f"=SUM('Plan'!$E${r1}:$E${r2})",TL,1)])
    B.label(wo,"B8","Ek ödeme yapmasaydın"); B.table(wo,9,2,["","Ek ödemeyle","Yalnız asgari ödemeyle","Fark"],3,fmt={})
    rows=[("Kapanış süresi (ay)","="+ay("Plan"),"="+ay("Yalnız Asgari"),'=IF(OR(C10="",D10=""),"",D10-C10)',NUM),("Borçsuz kalınan ay","="+end("Plan"),"="+end("Yalnız Asgari"),'=""',"MMMM YYYY"),
          ("Toplam faiz (₺)",f"=SUM('Plan'!$E${r1}:$E${r2})",f"=SUM('Yalnız Asgari'!$E${r1}:$E${r2})","=D12-C12",TL)]
    for k,(lab,a,b,c,nf) in enumerate(rows):
        r=10+k; wo.cell(r,2,lab).font=font(10,True)
        for col,v in ((3,a),(4,b),(5,c)): cc=wo.cell(r,col,v); cc.number_format=nf; cc.alignment=Alignment(horizontal="center")
    wo["B14"]="Asgari ödemeyle bile kapanan borçlardan boşalan tutar sıradaki borca aktarılır; iki sütunda da bu varsayılır."; wo["B14"].font=font(9,c="475467",i=True)
    wo["B15"]=f"Yalnız asgari ödeme {NT} ayda borcu bitirmiyorsa o sütundaki faiz, {NT} ayın faizidir; gerçekte daha fazladır."; wo["B15"].font=font(9,c="475467",i=True)
    chart(wo,"B17","Kalan toplam borç (₺)",Reference(B.wb["Plan"],min_col=1,min_row=r1,max_row=r1+59),[Reference(B.wb["Plan"],min_col=6,min_row=P0+2,max_row=r1+59),Reference(B.wb["Yalnız Asgari"],min_col=6,min_row=P0+2,max_row=r1+59)],[GREEN,"D92D20"],w=18,kind="line")
    wo._charts[0].series[0].tx=None; wo._charts[0].series[1].tx=None
    from openpyxl.chart.series import SeriesLabel
    wo._charts[0].series[0].tx=SeriesLabel(v="Ek ödemeyle"); wo._charts[0].series[1].tx=SeriesLabel(v="Yalnız asgari ödemeyle")
    B.basla([("NASIL KULLANILIR",None),("1. Borçlar","Her borcun kalan tutarını, aylık faiz oranını ve asgari aylık ödemesini yaz. Faizsiz borçta oranı 0 gir."),("2. Ödeme gücü","Asgari ödemelerin üstüne her ay ne kadar ayırabileceğini ve yöntemi (Çığ ya da Kartopu) seç."),("3. Özet","Ne zaman borçsuz kalacağını, toplam faizi ve yalnız asgari ödeseydin ne olacağını gör."),("4. Borçlar sayfasına dön","Her borcun öncelik sırası, kapanış ayı ve ödenecek faizi yazılır."),
      (None,None),("NASIL HESAPLANIR",None),("Aylık akış","Her ay önce bütün borçların faizi işler ve asgari ödemeleri yapılır. Kalan tutar öncelik sırasındaki ilk borca gider; o kapanınca sıradakine geçer. Kapanan borcun asgari ödemesi de sıradaki borca eklenir."),("Faiz oranı","Aylık oranı vergiler dahil gir. Yıllık oran biliyorsan kabaca 12'ye böl; bankanın kendi hesabı küçük farklarla çıkabilir."),
      (None,None),("BİLMEN GEREKENLER",None),("Banka hesabıyla birebir tutmaz","Gecikme faizi, dönem içi harcama, vergi ve masraf ayrıntıları hesaba katılmaz. Plan bir yol haritasıdır, ödeme tablosu değildir."),("Yeni borç yok varsayımı","Plan, kartla yeni harcama yapmadığını ve her ay aynı tutarı ödediğini varsayar."),("Finansal tavsiye değildir","Yapılandırma, borç birleştirme gibi kararlar için bankanla ya da bir uzmanla görüş."),("Kapasite",f"{NB} borç, {NT} ay.")])
    B.ai("Banka adı, kart numarası ve hesap bilgisi paylaşma; borç adlarını genel tut.",[("Yöntem karşılaştırma","[Borçlar sayfası] Çığ ve kartopu yöntemini benim borçlarım için karşılaştır: hangisi daha az faiz ödetir, hangisi ilk borcu daha erken kapatır?"),("Ek ödeme etkisi","Aylık ek ödememi [tutar] artırırsam borçsuz kalacağım ay ne kadar öne gelir? Kabaca hesapla, ben dosyada doğrulayacağım."),("Bütçeden yer açma","Aylık gelirim [tutar], sabit giderlerim [tutar]. Borca daha fazla ayırmak için nerelere bakabilirim? Somut kalemler öner."),("Banka görüşmesi","Kredi kartı borcumu yapılandırmak için bankayı arayacağım. Sormam gereken beş soruyu yaz."),("Motivasyon planı","Borç kapatma planım [ay] ay sürüyor. Her borç kapandığında kendime koyabileceğim küçük ara hedefler öner."),("Risk kontrolü","[Borçlar sayfası] Asgari ödemesi faizini karşılamayan ya da çok uzun süren borç var mı? Varsa neden tehlikeli olduğunu açıkla.")])
    fn=OUT+"ExcelKlinigi_ACIL_Borc_Kapatma_Plani.xlsx"; B.save(["Başla",Bs,"Özet","Plan","Yalnız Asgari","AI Komutları"],fn)
    # bağımsız simülasyon
    def sim(method,extra):
        ds=[dict(ad=a,b=float(b),r=r,m=m) for a,b,r,m in borc]; order=sorted(range(len(ds)),key=(lambda i:(ds[i]["b"],i)) if method=="Kartopu" else (lambda i:(-ds[i]["r"],i)))
        budget=sum(d["m"] for d in ds)+extra; faiz=[0.0]*len(ds); bit=[None]*len(ds); tot=[]
        for t in range(1,NT+1):
            need=[d["b"]*(1+d["r"]) for d in ds]
            for i,d in enumerate(ds): faiz[i]+=need[i]-d["b"]
            mn=[min(n,d["m"]) for n,d in zip(need,ds)]; pool=max(0,budget-sum(mn)); used=0
            for i in order:
                x=max(0,min(need[i]-mn[i],pool-used)); used+=x; ds[i]["b"]=round(need[i]-mn[i]-x,2)
                if ds[i]["b"]==0 and bit[i] is None: bit[i]=t
            tot.append(sum(d["b"] for d in ds))
        return order,bit,faiz,tot
    w=load_workbook(fn,data_only=True); bs=w[Bs]; o=w["Özet"]
    order,bit,faiz,tot=sim("Çığ",5000); o2,b2,f2,t2=sim("Çığ",0)
    assert [bs.cell(R0+i,5).value for i in range(4)]==[order.index(i)+1 for i in range(4)],[bs.cell(R0+i,5).value for i in range(4)]
    for i in range(4):
        got=bs.cell(R0+i,6).value; expd=date(2026,10,1); m_=bit[i]-1; expd=date(2026+(9+m_)//12,(9+m_)%12+1,1); assert got.date()==expd,(i,got,expd)
        assert abs(bs.cell(R0+i,7).value-faiz[i])<0.5,(i,bs.cell(R0+i,7).value,faiz[i])
    assert o["C10"].value==tot.index(0)+1 and o["D10"].value==t2.index(0)+1,(o["C10"].value,tot.index(0)+1,o["D10"].value,t2.index(0)+1)
    assert abs(o["C12"].value-sum(faiz))<1 and abs(o["D12"].value-sum(f2))<1
    print(f"  ACİL doğrulandı: ek ödemeyle {tot.index(0)+1} ay, {sum(faiz):,.0f} ₺ faiz; yalnız asgariyle {t2.index(0)+1} ay, {sum(f2):,.0f} ₺ faiz")
    # kartopu yöntemi de tutuyor mu
    import shutil, subprocess
    shutil.copy(fn,"t.xlsx"); w2=load_workbook("t.xlsx"); w2[Bs]["K6"]="Kartopu"; w2.save("t.xlsx"); subprocess.run([sys.executable,RECALC,"t.xlsx","300"],capture_output=True)
    w3=load_workbook("t.xlsx",data_only=True); order,bit,faiz,tot=sim("Kartopu",5000)
    assert [w3[Bs].cell(R0+i,5).value for i in range(4)]==[order.index(i)+1 for i in range(4)] and abs(w3["Özet"]["C12"].value-sum(faiz))<1 and w3["Özet"]["C10"].value==tot.index(0)+1; print("  ACİL kartopu yöntemi de doğrulandı")

def ameliyat():
    B=Book("AMELİYAT","Proje ve görev planı","Görevleri, sorumluları ve tarihleri yaz; zaman çizelgesi, gecikenler ve genel ilerleme kendiliğinden çıksın.")
    NG,NC,NKi=100,60,15; G="Görevler"; Ay="Ayarlar"
    wa=B.sheet(Ay,"Ayarlar","Proje adını, çizelgenin başlangıcını ve ölçeğini belirle.",{"A":3,"B":40,"C":20,"D":3,"E":24},grid=False)
    for r,lab,val,nf in ((4,"Proje adı","Dükkân yenileme",None),(5,"Çizelgenin başladığı tarih",date(2026,9,28),DT),(6,"Ölçek: bir sütun kaç gün? (1 = gün, 7 = hafta)",7,"0")):
        a=wa.cell(r,2,lab); a.font=font(10,True); a.fill=fill(GREY); a.border=box; B.inp(wa,f"C{r}",val,nf)
    B.dv(wa,"date","C5"); B.dv(wa,"int","C6",(1,31)); B.head(wa,4,["Kişiler"],start=5,h=22); B.body(wa,R0,NKi,1,start=5); B.put(wa,R0,[["Kişi Alfa"],["Usta"],["Elektrikçi"],["Tasarımcı"]],start=5)
    T="TODAY()"
    wg=B.sheet(G,"Görevler","Her görev bir satır. İlerlemeyi 0 ile 100 arasında yüzde olarak güncelle.",{"A":34,"B":16,"C":13,"D":13,"E":11,"F":9,"G":13,"H":11,"I":4})
    B.head(wg,4,["Görev","Sorumlu","Başlangıç","Bitiş","İlerleme","Süre (gün)","Durum","Kalan gün","y"],calc=(6,7,8,9)); B.body(wg,R0,NG,9,calc=(6,7,8,9),fmt={3:DT,4:DT,5:"0%",6:NUM,8:NUM},center=(3,4,5,6,7,8))
    gor=[("Plan ve ölçü alma","Kişi Alfa",date(2026,9,28),date(2026,10,2),1),("Malzeme siparişi","Kişi Alfa",date(2026,9,30),date(2026,10,3),0.6),("Elektrik tesisatı","Elektrikçi",date(2026,10,5),date(2026,10,14),0),
         ("Boya ve zemin","Usta",date(2026,10,12),date(2026,10,25),0),("Tabela tasarımı","Tasarımcı",date(2026,10,1),date(2026,10,20),0.3),("Raf montajı","Usta",date(2026,10,26),date(2026,11,6),0),("Açılış hazırlığı","Kişi Alfa",date(2026,11,9),date(2026,11,15),0)]
    B.put(wg,R0,gor)
    for i in range(NG):
        r=R0+i
        wg.cell(r,6,f'=IF(OR(C{r}="",D{r}="",D{r}<C{r}),"",D{r}-C{r}+1)')
        wg.cell(r,7,f'=IF(A{r}="","",IF(F{r}="","Tarih eksik",IF(E{r}>=1,"Tamamlandı",IF(D{r}<{T},"Gecikti",IF(C{r}>{T},"Başlamadı","Sürüyor")))))')
        wg.cell(r,8,f'=IF(OR(F{r}="",G{r}="Tamamlandı"),"",D{r}-{T})'); wg.cell(r,9,f'=IF(G{r}="Gecikti",D{r}+ROW()/1000000,"")')
    wg.column_dimensions["I"].hidden=True; wg.freeze_panes="B5"
    B.dv(wg,"list",f"B{R0}:B{R0+NG-1}",f"='{Ay}'!$E${R0}:$E${R0+NKi-1}"); B.dv(wg,"date",f"C{R0}:D{R0+NG-1}"); B.dv(wg,"pct",f"E{R0}:E{R0+NG-1}")
    for t,fc,tc in (("Tamamlandı",MINT,GREEN),("Gecikti",PINK,RED),("Sürüyor",AMB,AMBT),("Tarih eksik",PINK,RED)): B.cf_eq(wg,f"G{R0}:G{R0+NG-1}",t,fc,tc)
    wz=B.sheet("Çizelge","Zaman Çizelgesi","Bu sayfaya yazma. Koyu yeşil tamamlanan kısım, açık yeşil kalan kısım, kırmızı çerçeveli sütun bugün.",{"A":30,"B":14,**{CL(c):5 for c in range(3,3+NC)}})
    wz["A3"]=f"='{Ay}'!C4"; wz["A3"].font=font(11,True)
    B.head(wz,4,["Görev","Sorumlu"]+[""]*NC,calc=range(1,3+NC),h=30)
    for j in range(NC):
        c=wz.cell(4,3+j,f"='{Ay}'!$C$5+{j}*'{Ay}'!$C$6"); c.number_format="D MMM"; c.font=font(7,True,WHITE); c.alignment=Alignment(horizontal="center",vertical="center",text_rotation=90)
    for i in range(NG):
        r=R0+i; s_=f"'{G}'!"; wz.cell(r,1,f'=IF({s_}A{r}="","",{s_}A{r})').font=font(9); wz.cell(r,2,f'=IF({s_}A{r}="","",{s_}B{r}&"")').font=font(9)
        for j in range(NC):
            L=CL(3+j); ps=f"{L}$4"; pe=f"({L}$4+'{Ay}'!$C$6-1)"
            wz.cell(r,3+j,f'=IF({s_}$F{r}="","",IF(AND({s_}$C{r}<={pe},{s_}$D{r}>={ps}),IF({ps}<={s_}$C{r}+ROUND({s_}$F{r}*{s_}$E{r},0)-1,2,1),""))').font=font(6,c=WHITE)
    rng=f"C{R0}:{CL(2+NC)}{R0+NG-1}"
    wz.conditional_formatting.add(rng,CellIsRule(operator="equal",formula=["2"],fill=fill(GREEN),font=Font(name=F,size=6,color=GREEN)))
    wz.conditional_formatting.add(rng,CellIsRule(operator="equal",formula=["1"],fill=fill("9FD8B5"),font=Font(name=F,size=6,color="9FD8B5")))
    red=Side(style="medium",color=RED); wz.conditional_formatting.add(f"C4:{CL(2+NC)}{R0+NG-1}",FormulaRule(formula=[f"AND(TODAY()>=C$4,TODAY()<C$4+'{Ay}'!$C$6)"],border=Border(left=red,right=red)))
    wz.freeze_panes="C5"
    GA,GB,GD,GE,GF,GG,GI=[rg(G,c,NG) for c in "ABDEFGI"]
    wo=B.sheet("Özet","Özet","Bu sayfaya yazma.",{"A":3,"B":22,"C":16,"D":16,"E":16,"F":18,"G":3,"H":32,"I":14,"J":14},grid=False)
    wo["B3"]=f"='{Ay}'!C4"; wo["B3"].font=font(12,True)
    B.cards(wo,5,[("Görev",f"=COUNTA({GA})",NUM),("Tamamlanan",f'=COUNTIF({GG},"Tamamlandı")',NUM),("Geciken",f'=COUNTIF({GG},"Gecikti")',NUM,1),("7 gün içinde bitecek",f'=COUNTIFS({GD},">="&TODAY(),{GD},"<="&TODAY()+7,{GG},"<>Tamamlandı",{GG},"<>")',NUM),("Genel ilerleme",f"=IF(SUM({GF})=0,0,SUMPRODUCT({GF},{GE})/SUM({GF}))","0%")])
    wo["B8"]="Genel ilerleme, görevlerin süresine göre ağırlıklı ortalamadır."; wo["B8"].font=font(9,c="475467",i=True)
    B.label(wo,"B10","Sorumluya göre"); B.table(wo,11,2,["Sorumlu","Görev","Tamamlanan","Geciken","Açık iş (gün)"],NKi,fmt={3:NUM,4:NUM,5:NUM,6:NUM})
    for k in range(NKi):
        r=12+k; wo.cell(r,2,f"=IF('{Ay}'!E{R0+k}=\"\",\"\",'{Ay}'!E{R0+k})"); wo.cell(r,3,f'=IF(B{r}="","",COUNTIF({GB},B{r}))'); wo.cell(r,4,f'=IF(B{r}="","",COUNTIFS({GB},B{r},{GG},"Tamamlandı"))')
        wo.cell(r,5,f'=IF(B{r}="","",COUNTIFS({GB},B{r},{GG},"Gecikti"))'); wo.cell(r,6,f'=IF(B{r}="","",SUMPRODUCT(({GB}=B{r})*1,{GF},1-{GE}))')
    B.label(wo,"H10","Geciken görevler (en eski 10)"); B.table(wo,11,8,["Görev","Bitiş","Gecikme (gün)"],10,fmt={9:DT,10:NUM},center=(9,10))
    for k in range(1,11):
        r=11+k; ix=f"MATCH(SMALL({GI},{k}),{GI},0)"; wo.cell(r,8,f'=IFERROR(INDEX({GA},{ix}),"")'); wo.cell(r,9,f'=IF(H{r}="","",INDEX({GD},{ix}))'); wo.cell(r,10,f'=IF(H{r}="","",TODAY()-I{r})')
    B.basla([("NASIL KULLANILIR",None),("1. Ayarlar","Proje adını, çizelgenin başlayacağı tarihi, ölçeği (gün ya da hafta) ve kişileri yaz."),("2. Görevler","Her görevi sorumlusu, başlangıç ve bitiş tarihiyle gir. İş ilerledikçe yüzdeyi güncelle; yüzde 100 olunca Tamamlandı yazar."),("3. Çizelge","Görevlerin zaman çizelgesi. Koyu yeşil tamamlanan, açık yeşil kalan kısım. Bugünün sütunu kırmızı çerçevelidir."),("4. Özet","Geciken görevler, 7 gün içinde bitecekler, genel ilerleme ve kişi başına yük."),
      (None,None),("BİLMEN GEREKENLER",None),("Durum","Bitiş tarihi geçmiş ve yüzde 100 olmamışsa Gecikti; başlangıcı gelmemişse Başlamadı; arada ise Sürüyor yazar. Bugünün tarihine göre her açışta güncellenir."),("Çizelge sınırı",f"Çizelge {NC} sütun gösterir: ölçek 7 ise yaklaşık 14 ay, ölçek 1 ise 2 ay. Daha uzun proje için ölçeği büyüt."),("Bağımlılık yok","Dosya görevler arası bağımlılığı (biri bitmeden diğeri başlamaz) tutmaz. Bir görev kayarsa sonrakilerin tarihini elle güncellemen gerekir."),("Kapasite",f"{NG} görev, {NKi} kişi.")])
    B.ai("Kişi adlarını takma adla değiştir; müşteriye özel bilgileri sil.",[("Risk taraması","[Görevler sayfası] Geciken ve yavaş ilerleyen görevlere bak. Projenin bitiş tarihini en çok hangileri tehdit ediyor?"),("Haftalık plan","[Görevler sayfası] Önümüzdeki 7 gün için kişi kişi yapılacaklar listesi çıkar."),("Durum raporu","[Özet ve Görevler sayfaları] Müşteriye ya da yöneticiye gönderilecek beş cümlelik bir durum raporu yaz: nerede olduğumuz, ne gecikti, sıradaki adım."),("Yük dengesi","[Özet sayfasındaki sorumlu tablosu] İş yükü kimde birikmiş? Hangi görevleri kime kaydırabilirim?"),("Görev kırılımı","[Görev adı] görevini bir ile üç günlük alt görevlere böl ve sırala."),("Gecikme telafisi","[Geciken görevler] Bu gecikmeleri kapatmak için hangi görevler paralel yürütülebilir? Varsayımlarını yaz.")])
    fn=OUT+"ExcelKlinigi_AMELIYAT_Proje_Plani.xlsx"; B.save(["Başla",G,"Çizelge","Özet",Ay,"AI Komutları"],fn)
    w=load_workbook(fn,data_only=True); g=w[G]; o=w["Özet"]; z=w["Çizelge"]; today=date(2026,10,4)
    def st(c,d,e): return "Tamamlandı" if e>=1 else ("Gecikti" if d<today else ("Başlamadı" if c>today else "Sürüyor"))
    exp=[st(c,d,e) for _,_,c,d,e in gor]; assert [g.cell(R0+i,7).value for i in range(7)]==exp,exp
    dur=[(d-c).days+1 for _,_,c,d,_ in gor]; assert [g.cell(R0+i,6).value for i in range(7)]==dur
    prog=sum(x*e for x,(_,_,_,_,e) in zip(dur,gor))/sum(dur); assert abs(o["F6"].value-prog)<1e-9 and o["B6"].value==7 and o["C6"].value==1 and o["D6"].value==exp.count("Gecikti")
    assert o["H12"].value=="Malzeme siparişi" and o["J12"].value==1
    # çizelge: haftalık sütunlar 28 Eyl, 5 Eki, 12 Eki...; "Boya ve zemin" 12-25 Eki → 3. ve 4. sütun
    row=[z.cell(R0+3,3+j).value for j in range(8)]; assert row==[None,None,1,1,None,None,None,None],row
    row=[z.cell(R0,3+j).value for j in range(3)]; assert row==[2,None,None],row; print("  AMELİYAT doğrulandı; durumlar:",exp)
if __name__=="__main__": ritim(); acil(); ameliyat()
