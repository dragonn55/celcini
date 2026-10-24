# TAHLİL (satış analizi), VİZİTE (personel izin ve mesai)
from lib import *
from datetime import date, timedelta
from openpyxl import load_workbook
import random

def tahlil():
    B=Book("TAHLİL","Satış analizi","Satış kayıtlarını yapıştır; en çok kazandıran ürünler, en değerli müşteriler, aylık gidiş ve haftanın günleri kendiliğinden çıksın.")
    NV,NK=2000,200; V="Veri"
    wv=B.sheet(V,"Veri","Satışlarını buraya yaz ya da başka dosyadan yapıştır. Her satır bir satış kalemi.",{"A":12,"B":24,"C":26,"D":10,"E":15,"F":4,"G":4,"H":4,"I":4})
    B.head(wv,4,["Tarih","Müşteri","Ürün","Adet","Tutar (₺)","y1","y2","y3","y4"],calc=(6,7,8,9)); B.body(wv,R0,NV,9,calc=(6,7,8,9),fmt={1:DT,4:NUM,5:TLI},center=(1,))
    rnd=random.Random(11); urun=[("Defter A4",60),("Tükenmez kalem",15),("Fotokopi kâğıdı",190),("Klasör",75),("Boya kalemi seti",150),("Hesap makinesi",350),("Silgi",10),("Zımba",120)]
    mus=["Müşteri Alfa","Müşteri Beta","Müşteri Gama","Müşteri Epsilon","Müşteri Delta","Perakende"]
    data=[]
    for _ in range(80):
        d=date(2026,7,1)+timedelta(days=rnd.randint(0,95)); u,f=rnd.choice(urun[:6] if rnd.random()<0.85 else urun); a=rnd.choice([1,2,3,5,10,20,40]); data.append((d,rnd.choice(mus),u,a,a*f))
    data.sort(); B.put(wv,R0,data)
    for i in range(NV):
        r=R0+i
        wv.cell(r,6,f'=IF(C{r}="",0,IF(COUNTIF(C${R0}:C{r},C{r})=1,1,0))'); wv.cell(r,7,f"=SUM(F${R0}:F{r})")
        wv.cell(r,8,f'=IF(B{r}="",0,IF(COUNTIF(B${R0}:B{r},B{r})=1,1,0))'); wv.cell(r,9,f"=SUM(H${R0}:H{r})")
    for c in "FGHI": wv.column_dimensions[c].hidden=True
    wv.freeze_panes="A5"; B.dv(wv,"date",f"A{R0}:A{R0+NV-1}"); B.dv(wv,"nonneg",f"D{R0}:E{R0+NV-1}")
    VA,VB,VC,VD,VE,VG,VI=[rg(V,c,NV) for c in "ABCDEGI"]
    def analiz(name,title,sub,keycol,cumcol,heads,extra):
        ws=B.sheet(name,title,sub,{"A":6,"B":28,"C":12,"D":16,"E":9,"F":12,"G":8,"H":16,"I":14,"J":3,"K":4,"L":4,"M":4,"N":4})
        B.head(ws,4,heads,calc=range(1,15),h=30); B.body(ws,R0,NK,14,calc=range(1,15),fmt={3:NUM,4:TL,5:"0.0%",6:"0.0%",8:TL,9:DT},center=(1,7,9),stripe=True)
        for i in range(NK):
            r=R0+i; k=i+1
            ws.cell(r,11,f'=IF({k}>MAX({cumcol}),"",INDEX({keycol},MATCH({k},{cumcol},0)))')          # K: sırasız ad
            ws.cell(r,12,f'=IF(K{r}="","",SUMIFS({VE},{keycol},K{r}))'); ws.cell(r,13,f'=IF(K{r}="","",L{r}+ROW()/1000000000)')
            ix=f"MATCH(LARGE($M${R0}:$M${R0+NK-1},{k}),$M${R0}:$M${R0+NK-1},0)"
            ws.cell(r,2,f'=IFERROR(INDEX($K${R0}:$K${R0+NK-1},{ix}),"")'); ws.cell(r,1,f'=IF(B{r}="","",{k})')
            ws.cell(r,4,f'=IF(B{r}="","",SUMIFS({VE},{keycol},B{r}))'); ws.cell(r,5,f'=IF(B{r}="","",IF(SUM({VE})=0,0,D{r}/SUM({VE})))')
            ws.cell(r,6,f'=IF(B{r}="","",SUM(E${R0}:E{r}))')
            ws.cell(r,7,f'=IF(B{r}="","","A")' if i==0 else f'=IF(B{r}="","",IF(F{r-1}<0.8,"A",IF(F{r-1}<0.95,"B","C")))')
            extra(ws,r)
        for c in "KLMN": ws.column_dimensions[c].hidden=True
        ws.freeze_panes="C5"
        for t,fc,tc in (("A",MINT,GREEN),("B",AMB,AMBT),("C",GREY,"667085")): B.cf_eq(ws,f"G{R0}:G{R0+NK-1}",t,fc,tc)
        return ws
    def ex_u(ws,r):
        ws.cell(r,3,f'=IF(B{r}="","",SUMIFS({VD},{VC},B{r}))'); ws.cell(r,8,f'=IF(OR(B{r}="",C{r}=0),"",D{r}/C{r})')
        mx=f"SUMPRODUCT(MAX(({VC}=B{r})*{VA}))"; ws.cell(r,9,f'=IF(B{r}="","",{mx})')
    def ex_m(ws,r):
        ws.cell(r,3,f'=IF(B{r}="","",COUNTIF({VB},B{r}))'); ws.cell(r,8,f'=IF(OR(B{r}="",C{r}=0),"",D{r}/C{r})')
        mx=f"SUMPRODUCT(MAX(({VB}=B{r})*{VA}))"; ws.cell(r,9,f'=IF(B{r}="","",{mx})')
    wu=analiz("Ürünler","Ürün Analizi","Cirosu en yüksek üründen başlayarak sıralanır. A sınıfı: cironun ilk yüzde 80'ini getiren ürünler.",VC,VG,["Sıra","Ürün","Adet","Ciro (₺)","Pay","Birikimli pay","Sınıf","Ortalama birim fiyat (₺)","Son satış","","y","y","y","y"],ex_u)
    wm=analiz("Müşteriler","Müşteri Analizi","Cirosu en yüksek müşteriden başlayarak sıralanır.",VB,VI,["Sıra","Müşteri","Satış satırı","Ciro (₺)","Pay","Birikimli pay","Sınıf","Satır başına ortalama (₺)","Son alışveriş","","y","y","y","y"],ex_m)
    wo=B.sheet("Özet","Özet","Bu sayfada yalnızca yılı değiştir.",{"A":3,"B":15,"C":16,"D":14,"E":14,"F":16,"G":16,"H":3,"I":16,"J":16},grid=False)
    wo["B3"]="Yıl"; B.inp(wo,"C3",2026,"0")
    UG=rg("Ürünler","G",NK); UD=rg("Ürünler","D",NK)
    B.cards(wo,5,[("Toplam ciro",f"=SUM({VE})",TL),("Satış satırı",f"=COUNT({VE})",NUM),("Satır başına ortalama",f"=IF(COUNT({VE})=0,0,SUM({VE})/COUNT({VE}))",TL),("Ürün çeşidi",f"=MAX({VG})",NUM),("Müşteri",f"=MAX({VI})",NUM),("A sınıfı ürün",f'=COUNTIF({UG},"A")',NUM)])
    wo["B8"]=f'="A sınıfı ürünlerin cirodaki payı: "&TEXT(IF(SUM({VE})=0,0,SUMIFS({UD},{UG},"A")/SUM({VE})),"0%")'; wo["B8"].font=font(10,True,GREEN)
    B.label(wo,"B10","Aylık gidiş"); B.table(wo,11,2,["Ay","Ciro (₺)","Adet","Satış satırı","Önceki aya göre"],12,fmt={3:TL,4:NUM,5:NUM,6:PCT})
    for m,ay in enumerate(months(),1):
        r=11+m; wo.cell(r,2,ay); wo.cell(r,3,"="+mrange(VE,VA,"$C$3",m)); wo.cell(r,4,"="+mrange(VD,VA,"$C$3",m)); wo.cell(r,5,"="+crange(VA,"$C$3",m))
        wo.cell(r,6,'=""' if m==1 else f'=IF(C{r-1}=0,"",C{r}/C{r-1}-1)')
    B.label(wo,"I10","Haftanın günleri"); B.table(wo,11,9,["Gün","Ciro (₺)"],7,fmt={10:TL})
    for k,g in enumerate(["Pazartesi","Salı","Çarşamba","Perşembe","Cuma","Cumartesi","Pazar"],1):
        r=11+k; wo.cell(r,9,g); wo.cell(r,10,f"=SUMPRODUCT((WEEKDAY({VA},2)={k})*({VA}>0)*{VE})")
    chart(wo,"B25","Aylık ciro (₺)",Reference(wo,min_col=2,min_row=12,max_row=23),[Reference(wo,min_col=3,min_row=11,max_row=23)],[GREEN],w=16)
    B.basla([("NASIL KULLANILIR",None),("1. Veri","Satışlarını tarih, müşteri, ürün, adet ve tutar olarak yaz ya da başka bir dosyadan kopyalayıp yapıştır. Yapıştırırken yalnızca değerleri yapıştır."),("2. Ürünler","Ürünler cirosuna göre sıralanır; pay, birikimli pay ve A, B, C sınıfı hesaplanır."),("3. Müşteriler","Müşteriler cirosuna göre sıralanır; satış sayısı, ortalama ve son alışveriş tarihi görünür."),("4. Özet","Toplamlar, aylık gidiş, önceki aya göre değişim ve haftanın günlerine göre ciro."),
      (None,None),("BİLMEN GEREKENLER",None),("A, B, C sınıfı","Ürünler cirosuna göre dizilir. Cironun ilk yüzde 80'ini getirenler A, sonraki yüzde 15'i getirenler B, kalanlar C sınıfıdır. A sınıfı az sayıda üründür ama kazancın çoğunu onlar getirir."),("Yazım farkı","Aynı ürünü ya da müşteriyi iki farklı yazarsan iki ayrı satır olarak sayılır. Adları tek biçimde yaz."),("Kâr değil ciro","Dosya maliyet bilgisi istemez; bu yüzden kâr değil ciro analizi yapar."),("Kapasite",f"{NV} satış satırı; analizde {NK} ürün ve {NK} müşteri. Dosya dolu veriyle birkaç saniye geç hesaplayabilir.")])
    B.ai("Müşteri adlarını Müşteri Alfa, Müşteri Beta gibi takma adlarla değiştir.",[("Genel yorum","[Özet sayfası] Satışlarımı beş cümleyle yorumla: gidiş nasıl, hangi ay öne çıkıyor, dikkat etmem gereken ne?"),("A sınıfı strateji","[Ürünler sayfası] A sınıfı ürünlerim için stokta kalmama, fiyat ve tanıtım açısından birer öneri ver."),("C sınıfı karar","[Ürünler sayfası] C sınıfı ürünlerden hangilerini bırakmayı düşünmeliyim? Son satış tarihine ve ciroya göre gerekçelendir."),("Müşteri bağımlılığı","[Müşteriler sayfası] Cironun ne kadarı ilk üç müşteriden geliyor? Bu bir risk mi, nasıl azaltırım?"),("Geri kazanma","[Müşteriler sayfası] Son alışverişi 45 günden eski olan müşterileri listele ve onlara gönderilecek kısa bir mesaj yaz."),("Birlikte satılanlar","[Veri sayfası] Aynı müşterinin aynı gün aldığı ürünlere bakarak sık birlikte satılan ürün çiftlerini bul."),("Zayıf gün","[Özet sayfasındaki gün tablosu] En zayıf günüm hangisi? O gün için bir kampanya fikri öner.")])
    fn=OUT+"ExcelKlinigi_TAHLIL_Satis_Analizi.xlsx"; B.save(["Başla",V,"Özet","Ürünler","Müşteriler","AI Komutları"],fn)
    w=load_workbook(fn,data_only=True); u=w["Ürünler"]; m=w["Müşteriler"]; o=w["Özet"]
    from collections import defaultdict
    cu=defaultdict(float); au=defaultdict(float); cm=defaultdict(float); mon=defaultdict(float); wd=defaultdict(float)
    for d,ms,ur,a,t in data: cu[ur]+=t; au[ur]+=a; cm[ms]+=t; mon[d.month]+=t; wd[d.isoweekday()]+=t
    tot=sum(cu.values()); su=sorted(cu.items(),key=lambda x:-x[1]); sm=sorted(cm.items(),key=lambda x:-x[1])
    got=[(u.cell(R0+i,2).value,u.cell(R0+i,4).value) for i in range(len(su))]; assert got==[(a,b) for a,b in su],(got,su)
    cum=0; cls=[]
    for a,b in su: cls.append("A" if cum<0.8 else ("B" if cum<0.95 else "C")); cum+=b/tot
    assert [u.cell(R0+i,7).value for i in range(len(su))]==cls,cls; assert u.cell(R0+len(su),2).value is None
    assert [(m.cell(R0+i,2).value,m.cell(R0+i,4).value) for i in range(len(sm))]==[(a,b) for a,b in sm]
    assert o["B6"].value==tot and o["C6"].value==80 and o["E6"].value==len(cu) and o["F6"].value==len(cm) and o["G6"].value==cls.count("A")
    assert all(o.cell(11+k,3).value==mon.get(k,0) for k in range(1,13)); assert all(o.cell(11+k,10).value==wd.get(k,0) for k in range(1,8)); print("  TAHLİL doğrulandı; A sınıfı:",cls.count("A"),"ürün")

def vizite():
    B=Book("VİZİTE","Personel izin ve mesai takibi","Kim ne zaman izinli, kimin kaç gün izni kaldı, kim ne kadar mesai yaptı.")
    NP,NI,NMs=50,500,500; P="Personel"; I="İzinler"; Ms="Mesai"
    wa=B.sheet("Ayarlar","Ayarlar","Yılı ve izin sayma kuralını belirle.",{"A":3,"B":44,"C":12,"D":3,"E":16,"F":36},grid=False)
    for r,lab,val in ((4,"Yıl",2026),(5,"Yıllık izinde cumartesi iş günü sayılsın mı? (1 = evet, 0 = hayır)",1)):
        a=wa.cell(r,2,lab); a.font=font(10,True); a.fill=fill(GREY); a.border=box; a.alignment=Alignment(wrap_text=True); B.inp(wa,f"C{r}",val,"0")
    wa.row_dimensions[5].height=30
    wa["B7"]="Varsayılan 1'dir: yıllık izin hesabında pazar ve resmî tatiller izinden sayılmaz, cumartesi sayılır. Bu genel uygulamadır; kendi iş yerindeki kuralı ve sözleşmeni esas al."; wa["B7"].font=font(9,c="475467",i=True); wa["B7"].alignment=Alignment(wrap_text=True,vertical="top"); wa.row_dimensions[7].height=44
    B.head(wa,4,["Resmî tatil","Açıklama"],start=5,h=22); B.body(wa,R0,20,2,fmt={5:DT},center=(5,),start=5)
    for k,(mm,dd,ad) in enumerate(((1,1,"Yılbaşı"),(4,23,"Ulusal Egemenlik ve Çocuk Bayramı"),(5,1,"Emek ve Dayanışma Günü"),(5,19,"Gençlik ve Spor Bayramı"),(7,15,"Demokrasi ve Millî Birlik Günü"),(8,30,"Zafer Bayramı"),(10,29,"Cumhuriyet Bayramı"))):
        wa.cell(R0+k,5,f"=DATE($C$4,{mm},{dd})"); wa.cell(R0+k,6,ad)
    wa["E26"]="Ramazan ve Kurban Bayramı günleri her yıl değişir; boş satırlara kendin ekle."; wa["E26"].font=font(9,c="475467",i=True)
    HOL="Ayarlar!$E$5:$E$24"; Y="Ayarlar!$C$4"
    IA,IB,IC,ID,IF_=[rg(I,c,NI) for c in "ABCDF"]; MsA,MsB,MsC=[rg(Ms,c,NMs) for c in "ABC"]
    wp=B.sheet(P,"Personel","Sarı sütunları doldur. Kullanılan ve kalan izin, rapor ve mesai kendiliğinden hesaplanır.",{"A":26,"B":20,"C":13,"D":12,"E":13,"F":11,"G":11,"H":12,"I":12,"J":10})
    B.head(wp,4,["Ad","Görev","İşe giriş","Yıllık izin hakkı (gün)","Kullanılan yıllık izin","Kalan izin","Rapor (gün)","Ücretsiz izin (gün)","Mesai (saat)","Kıdem (yıl)"],calc=range(5,11))
    B.body(wp,R0,NP,10,calc=range(5,11),fmt={3:DT,4:NUM,5:NUM,6:NUM,7:NUM,8:NUM,9:'#,##0.0;-#,##0.0;"-"',10:"0.0"},center=(3,),bold=(6,))
    per=[("Personel Alfa","Usta",date(2019,3,4),20),("Personel Beta","Kasiyer",date(2023,6,12),14),("Personel Gama","Depo görevlisi",date(2025,1,20),14)]
    B.put(wp,R0,per)
    yr=f'{IC},">="&DATE({Y},1,1),{IC},"<"&DATE({Y}+1,1,1)'
    for i in range(NP):
        r=R0+i
        wp.cell(r,5,f'=IF(A{r}="","",SUMIFS({IF_},{IA},A{r},{IB},"Yıllık izin",{yr}))'); wp.cell(r,6,f'=IF(A{r}="","",D{r}-E{r})')
        wp.cell(r,7,f'=IF(A{r}="","",SUMIFS({IF_},{IA},A{r},{IB},"Rapor",{yr}))'); wp.cell(r,8,f'=IF(A{r}="","",SUMIFS({IF_},{IA},A{r},{IB},"Ücretsiz izin",{yr}))')
        wp.cell(r,9,f'=IF(A{r}="","",SUMIFS({MsC},{MsA},A{r},{MsB},">="&DATE({Y},1,1),{MsB},"<"&DATE({Y}+1,1,1)))'); wp.cell(r,10,f'=IF(OR(A{r}="",C{r}=""),"",(TODAY()-C{r})/365.25)')
    wp.freeze_panes="B5"; B.dv(wp,"uniq",f"A{R0}:A{R0+NP-1}",f"=COUNTIF($A${R0}:$A${R0+NP-1},A{R0})=1"); B.dv(wp,"date",f"C{R0}:C{R0+NP-1}"); B.dv(wp,"nonneg",f"D{R0}:D{R0+NP-1}")
    B.cf_f(wp,f"F{R0}:F{R0+NP-1}",f'AND(ISNUMBER(F{R0}),F{R0}<0)',PINK,RED)
    PA=rg(P,"A",NP)
    wi=B.sheet(I,"İzinler","Her izin ya da rapor bir satır. Gün sayısı kendiliğinden hesaplanır.",{"A":26,"B":16,"C":13,"D":13,"E":30,"F":9,"G":22})
    B.head(wi,4,["Personel","Tür","Başlangıç","Bitiş (dahil)","Not","Gün","Kontrol"],calc=(6,7)); B.body(wi,R0,NI,7,calc=(6,7),fmt={3:DT,4:DT,6:NUM},center=(2,3,4,6))
    izin=[("Personel Alfa","Yıllık izin",date(2026,7,13),date(2026,7,25)),("Personel Beta","Yıllık izin",date(2026,8,24),date(2026,9,2)),("Personel Gama","Rapor",date(2026,9,14),date(2026,9,16)),("Personel Alfa","Yıllık izin",date(2026,10,26),date(2026,10,31)),("Personel Beta","Ücretsiz izin",date(2026,10,5),date(2026,10,6)),("Personel Gama","Mazeret izni",date(2026,10,4),date(2026,10,4))]
    B.put(wi,R0,izin)
    for i in range(NI):
        r=R0+i; c,d=f"C{r}",f"D{r}"
        paz=f"INT((WEEKDAY({c}-1)+{d}-{c})/7)"; cmt=f"INT((WEEKDAY({c}-7)+{d}-{c})/7)"
        tat=f"SUMPRODUCT(({HOL}>={c})*({HOL}<={d})*(WEEKDAY({HOL},2)<IF(Ayarlar!$C$5=1,7,6)))"
        wi.cell(r,6,f'=IF(OR({c}="",{d}="",{d}<{c}),"",IF(B{r}="Yıllık izin",{d}-{c}+1-{paz}-IF(Ayarlar!$C$5=1,0,{cmt})-{tat},{d}-{c}+1))')
        wi.cell(r,7,f'=IF(A{r}="","",IF(OR(B{r}="",{c}="",{d}=""),"Eksik bilgi",IF({d}<{c},"Bitiş başlangıçtan önce","OK")))')
    wi.freeze_panes="A5"; B.dv(wi,"list",f"A{R0}:A{R0+NI-1}","="+PA); B.dv(wi,"list",f"B{R0}:B{R0+NI-1}","Yıllık izin,Rapor,Ücretsiz izin,Mazeret izni"); B.dv(wi,"date",f"C{R0}:D{R0+NI-1}")
    B.cf_f(wi,f"G{R0}:G{R0+NI-1}",f'AND(G{R0}<>"",G{R0}<>"OK")',PINK,RED)
    wm=B.sheet(Ms,"Mesai","Fazla çalışmayı kişi, tarih ve saat olarak yaz.",{"A":26,"B":13,"C":10,"D":34})
    B.head(wm,4,["Personel","Tarih","Saat","Açıklama"]); B.body(wm,R0,NMs,4,fmt={2:DT,3:"0.0"},center=(2,3))
    B.put(wm,R0,[("Personel Alfa",date(2026,9,12),3,"Örnek"),("Personel Gama",date(2026,9,19),4.5,"Örnek"),("Personel Alfa",date(2026,10,3),2,"Örnek")])
    B.dv(wm,"list",f"A{R0}:A{R0+NMs-1}","="+PA); B.dv(wm,"date",f"B{R0}:B{R0+NMs-1}"); B.dv(wm,"pos",f"C{R0}:C{R0+NMs-1}")
    wy=B.sheet("Ay","Ay Çizelgesi","Ayı yaz. İ yıllık izin, R rapor, Ü ücretsiz izin, M mazeret izni. Gri sütunlar pazar günleridir.",{"A":3,"B":24,**{CL(c):4.2 for c in range(3,34)},"AH":9},grid=False)
    wy["B3"]="Ay (1-12)"; B.inp(wy,"C3",10,"0"); wy.merge_cells("C3:D3"); B.dv(wy,"int","C3",(1,12))
    B.head(wy,5,["Personel"]+[""]*31+["İzinli gün"],start=2,h=20)
    for d in range(1,32):
        c=wy.cell(5,2+d,f"=IF(DAY(DATE({Y},$C$3,{d}))={d},DATE({Y},$C$3,{d}),\"\")"); c.number_format="D"; c.font=font(8,True,WHITE)
    B.body(wy,6,NP,33,calc=range(2,35),center=range(3,35),start=2)
    for i in range(NP):
        r=6+i; wy.cell(r,2,f"=IF('{P}'!A{R0+i}=\"\",\"\",'{P}'!A{R0+i})")
        for d in range(1,32):
            L=CL(2+d); ix=f"SUMPRODUCT(MAX(({IA}=$B{r})*({IC}<={L}$5)*({ID}>={L}$5)*ROW({IA})))"
            wy.cell(r,2+d,f'=IF(OR($B{r}="",{L}$5=""),"",IF({ix}=0,"",LEFT(SUBSTITUTE(SUBSTITUTE(INDEX(\'{I}\'!$B$1:$B${R0+NI-1},{ix}),"Yıllık izin","İ"),"Ücretsiz izin","Ü"),1)))').font=font(8,True)
        wy.cell(r,34,f'=IF(B{r}="","",COUNTIF(C{r}:AG{r},"?"))')
    rng=f"C6:AG{5+NP}"
    for t,fc,tc in (("İ",MINT,GREEN),("R",PINK,RED),("Ü",AMB,AMBT),("M","DBEAFE","1D4ED8")): B.cf_eq(wy,rng,t,fc,tc)
    B.cf_f(wy,rng,'AND(C$5<>"",WEEKDAY(C$5,2)=7)',"E5E7EB",None); wy.freeze_panes="C6"
    wo=B.sheet("Özet","Özet","Bu sayfaya yazma. Yılı Ayarlar sayfasından değiştir.",{"A":3,"B":20,"C":18,"D":18,"E":18,"F":18},grid=False)
    PE,PF,PG,PI=[rg(P,c,NP) for c in "EFGI"]
    B.cards(wo,4,[("Personel",f"=COUNTA({PA})",NUM),("Bugün izinde",f'=SUMPRODUCT(({IA}<>"")*({IC}<=TODAY())*({ID}>=TODAY()))',NUM,1),("Kullanılan yıllık izin (gün)",f"=SUM({PE})",NUM),("Kalan yıllık izin (gün)",f"=SUM({PF})",NUM),("Mesai (saat)",f"=SUM({PI})",'#,##0.0')])
    B.label(wo,"B8","Aylara göre izin ve mesai"); B.table(wo,9,2,["Ay","Yıllık izin (gün)","Rapor (gün)","Mesai (saat)"],12,fmt={3:NUM,4:NUM,5:'#,##0.0;-#,##0.0;"-"'})
    for m,ay in enumerate(months(),1):
        r=9+m; wo.cell(r,2,ay); wo.cell(r,3,"="+mrange(IF_,IC,Y,m,f',{IB},"Yıllık izin"')); wo.cell(r,4,"="+mrange(IF_,IC,Y,m,f',{IB},"Rapor"')); wo.cell(r,5,"="+mrange(MsC,MsB,Y,m))
    wo["B23"]="İzinler, başladığı aya yazılır."; wo["B23"].font=font(9,c="475467",i=True)
    B.basla([("NASIL KULLANILIR",None),("1. Ayarlar","Yılı yaz. Ramazan ve Kurban Bayramı günlerini resmî tatil listesine ekle."),("2. Personel","Çalışanları, işe giriş tarihini ve o yılki yıllık izin hakkını yaz."),("3. İzinler","Her izni ya da raporu başlangıç ve bitiş tarihiyle gir. Gün sayısı kendiliğinden hesaplanır."),("4. Mesai","Fazla çalışmayı kişi, tarih ve saat olarak yaz."),("5. Ay ve Özet","Ay çizelgesinde kimin hangi gün izinli olduğunu, Özet'te toplamları gör."),
      (None,None),("BİLMEN GEREKENLER",None),("Yıllık izin günü","Başlangıç ve bitiş arasındaki günlerden pazarlar ve listedeki resmî tatiller düşülür. Ayarlar'da cumartesiyi de düşecek şekilde değiştirebilirsin. Rapor ve diğer izinler takvim günü olarak sayılır."),("İzin hakkı","Yıllık izin hakkını sen yazarsın; dosya kıdeme göre hak hesaplamaz ve önceki yıldan devreden izni kendiliğinden eklemez."),("Yıl geçişi","Bir izin iki yıla yayılıyorsa başladığı yıla yazılır; iki ayrı satır olarak girmen daha doğru olur."),("Resmî belge değil","Bordro ve puantaj programı yerine geçmez; kendi takibin içindir."),("Kapasite",f"{NP} personel, {NI} izin kaydı, {NMs} mesai kaydı.")])
    B.ai("Çalışan adlarını Personel Alfa, Personel Beta gibi takma adlarla değiştir; rapor nedenlerini paylaşma.",[("İzin çakışması","[İzinler sayfası] Aynı görevdeki kişilerin aynı günlerde izinli olduğu durumları bul ve hangi günlerde eksik kalacağımı söyle."),("İzin planı","[Personel sayfası] Kalan izni en çok olan kişileri listele. Yıl sonuna kadar izinleri eritmek için kişi başı aylık bir plan öner."),("Mesai yükü","[Mesai sayfası] Kişi ve ay bazında mesai saatlerini topla. Yük kimde birikmiş, dengelemek için ne yapabilirim?"),("Rapor örüntüsü","[İzinler sayfası] Rapor kayıtlarının haftanın günlerine ve aylara dağılımını çıkar. Yorum yapmadan yalnızca dağılımı ver."),("Duyuru metni","Yıllık izin taleplerinin en geç iki hafta önce bildirilmesini isteyen kısa ve kibar bir duyuru yaz."),("Yoğun dönem","[Özet sayfasındaki aylık tablo] İzinlerin en çok yığıldığı ayları bul ve o aylar için vardiya önerisi ver.")])
    fn=OUT+"ExcelKlinigi_VIZITE_Personel_Izin_Takibi.xlsx"; B.save(["Başla",I,Ms,"Ay","Özet",P,"Ayarlar","AI Komutları"],fn)
    w=load_workbook(fn,data_only=True); iz=w[I]; p=w[P]; ay=w["Ay"]; o=w["Özet"]
    hol={date(2026,1,1),date(2026,4,23),date(2026,5,1),date(2026,5,19),date(2026,7,15),date(2026,8,30),date(2026,10,29)}
    def gun(t,a,b):
        if t!="Yıllık izin": return (b-a).days+1
        return sum(1 for k in range((b-a).days+1) if (a+timedelta(k)).isoweekday()!=7 and (a+timedelta(k)) not in hol)
    exp=[gun(t,a,b) for _,t,a,b in izin]; got=[iz.cell(R0+i,6).value for i in range(len(izin))]; assert got==exp,(got,exp)
    assert [p.cell(R0+i,5).value for i in range(3)]==[exp[0]+exp[3],exp[1],0]; assert [p.cell(R0+i,6).value for i in range(3)]==[20-exp[0]-exp[3],14-exp[1],14]
    assert p.cell(R0+2,7).value==3 and p.cell(R0+1,8).value==2 and p.cell(R0,9).value==5 and p.cell(R0+2,9).value==4.5
    row=lambda r:[ay.cell(r,2+d).value for d in range(1,32)]
    a=row(6); assert [d for d in range(1,32) if a[d-1]=="İ"]==[26,27,28,29,30,31],a
    f=row(7); assert [d for d in range(1,32) if f[d-1]=="Ü"]==[5,6]; m_=row(8); assert [d for d in range(1,32) if m_[d-1]=="M"]==[4]
    assert ay.cell(6,34).value==6 and o["B5"].value==3; print("  VİZİTE doğrulandı; yıllık izin günleri:",exp)
if __name__=="__main__": tahlil(); vizite()
