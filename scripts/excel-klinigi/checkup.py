from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.chart import BarChart, LineChart, Reference
from datetime import date
import json

GREEN="1F6F43"; MINT="E8F5EE"; INPUT="FFF9DB"; GREY="F3F4F6"; RED="B42318"; DARK="1B1F23"; WHITE="FFFFFF"
F="Arial"
def font(sz=10,b=False,c=DARK,i=False): return Font(name=F,size=sz,bold=b,color=c,italic=i)
def fill(c): return PatternFill("solid",start_color=c,end_color=c)
thin=Side(style="thin",color="D0D5DD"); box=Border(left=thin,right=thin,top=thin,bottom=thin)
TL='#,##0.00 "₺";[Red]-#,##0.00 "₺";"-"'; TLIN='#,##0.00 "₺"'; QTY='#,##0.##;[Red]-#,##0.##;"-"'
R0=5
NS,NA,NG,NH,NU,NC=1000,500,500,500,200,100
def last(n): return R0+n-1
def rg(sheet,col,n): return f"'{sheet}'!${col}${R0}:${col}${last(n)}"
SA,SB,SC,SD,SF,SG,SH=[rg("Satışlar",c,NS) for c in "ABCDFGH"]
AA,AB,AC,AD,AF,AG=[rg("Alışlar",c,NA) for c in "ABCDFG"]
GA,GD=rg("Giderler","A",NG),rg("Giderler","D",NG)
HB,HC,HD=[rg("Cari Hareket",c,NH) for c in "BCD"]
UA,UC,UD,UI,UJ,UK,UL,UM,UN=[rg("Ürünler",c,NU) for c in "ACDIJKLMN"]
CA,CG,CJ,CL=[rg("Cariler",c,NC) for c in "AGJL"]

wb=Workbook()
def title(ws,t,sub):
    ws["A1"]=t; ws["A1"].font=font(16,True,GREEN)
    ws["A2"]=sub; ws["A2"].font=font(9,c="475467",i=True)
def header(ws,row,heads,start=1,formula_cols=()):
    for j,h in enumerate(heads):
        c=ws.cell(row,start+j,h); c.font=font(10,True,WHITE); c.fill=fill("3B8A5E" if (start+j) in formula_cols else GREEN); c.border=box
        c.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
    ws.row_dimensions[row].height=28
def style_rows(ws,n,ncols,formula_cols=(),money=(),qty=(),dates=(),center=()):
    for i in range(n):
        r=R0+i
        for j in range(1,ncols+1):
            c=ws.cell(r,j); c.font=font(10); c.border=box
            c.fill=fill(WHITE if j in formula_cols else INPUT)
            if j in money: c.number_format=TL if j in formula_cols else TLIN
            if j in qty: c.number_format=QTY
            if j in dates: c.number_format="DD.MM.YYYY"; c.alignment=Alignment(horizontal="center")
            if j in center: c.alignment=Alignment(horizontal="center")
def dv(ws,v,ref,t,msg):
    v.errorTitle=t; v.error=msg; v.showErrorMessage=True; ws.add_data_validation(v); v.add(ref)
def widths(ws,ws_w):
    for k,v in ws_w.items(): ws.column_dimensions[k].width=v

# ================= Başla =================
ws=wb.active; ws.title="Başla"; ws.sheet_view.showGridLines=False
widths(ws,{"A":3,"B":24,"C":88})
ws.merge_cells("B2:C2"); ws["B2"]="EXCEL KLİNİĞİ"; ws["B2"].font=font(11,True,WHITE); ws["B2"].fill=fill(GREEN); ws["B2"].alignment=Alignment(vertical="center",indent=1); ws.row_dimensions[2].height=24
ws.merge_cells("B3:C3"); ws["B3"]="CHECK-UP  |  Esnaf İşletme Defteri"; ws["B3"].font=font(20,True,GREEN); ws.row_dimensions[3].height=36
ws.merge_cells("B4:C4"); ws["B4"]="Satış, alış, stok, veresiye, kasa ve kâr tek dosyada. Hareketi gir, işletmenin durumunu Panel söylesin."; ws["B4"].font=font(10,c="475467")
rows=[
 ("İLK KURULUM (bir kez)",None),
 ("1. Ayarlar","İşletme adını, açılış kasa bakiyesini ve gider kategorilerini yaz."),
 ("2. Ürünler","Sattığın ürünleri gir: ad, birim, alış maliyeti, satış fiyatı, eldeki açılış stoku ve kritik seviye."),
 ("3. Cariler","Veresiye çalıştığın müşterileri ve tedarikçileri gir. Önceden kalan alacak ya da borç varsa açılış sütunlarına yaz."),
 (None,None),
 ("GÜNLÜK KULLANIM",None),
 ("Satışlar","Her satış bir satır. Ürünü seç, adedi yaz. Birim fiyatı boş bırakırsan Ürünler sayfasındaki satış fiyatı kullanılır. Veresiye satışta müşteri seçmek zorunlu."),
 ("Alışlar","Her mal alımı bir satır. Birim maliyeti boş bırakırsan Ürünler sayfasındaki alış maliyeti kullanılır. Alış girince stok artar."),
 ("Giderler","Kira, fatura, personel gibi mal alımı dışındaki harcamalar."),
 ("Cari Hareket","Veresiye müşteriden para aldığında Tahsilat, tedarikçiye veresiye borcunu ödediğinde Ödeme gir."),
 ("Panel","Hiçbir şey yazma, yalnızca yılı seç. Ciro, kâr, kasa, alacak, borç, stok ve uyarılar kendiliğinden hesaplanır."),
 (None,None),
 ("NASIL HESAPLANIR",None),
 ("Stok","Mevcut = açılış stoku + alışlar − satışlar. Mevcut kritik seviyeye inince Kritik, sıfıra inince Tükendi yazar."),
 ("Satılan malın maliyeti","Satılan adet × Ürünler sayfasındaki güncel alış maliyeti. Alış maliyetini değiştirirsen geçmiş satışların maliyeti ve kârı da yeni rakama göre hesaplanır."),
 ("Kasa + Banka","Açılış bakiyesi + peşin satışlar + tahsilatlar − peşin alışlar − ödemeler − giderler. Nakit, kart ve havale ayrımı yapılmaz, hepsi tek bakiyedir."),
 ("Alacak ve borç","Alacak = açılış alacağı + veresiye satışlar − tahsilatlar. Borç = açılış borcu + veresiye alışlar − ödemeler."),
 ("KDV","Dosya KDV ayrımı yapmaz. Tutarları KDV dahil ya da hariç gir, ama hep aynı şekilde gir. Resmî defter ya da beyanname yerine geçmez."),
 (None,None),
 ("RENKLER",None),
 ("Sarı hücre","Senin dolduracağın hücre."),
 ("Beyaz hücre, açık yeşil başlık","Formül. Değiştirme, kendiliğinden hesaplanır."),
 (None,None),
 ("ÖNEMLİ",None),
 ("Örnek veriler","Dosya örnek bir kırtasiye dükkânının verileriyle dolu gelir. Başlamadan önce Satışlar, Alışlar, Giderler, Cari Hareket, Ürünler ve Cariler sayfalarında yalnızca sarı hücrelerin içeriğini sil."),
 ("Ad değiştirme","Ürün ve cari adları hareketleri birbirine bağlar. Bir adı sonradan değiştirirsen eski hareketlerdeki adı da değiştir, yoksa stok ve bakiye eksik çıkar. Aynı ad iki kez girilemez."),
 ("Kapasite",f"{NS} satış, {NA} alış, {NG} gider, {NH} cari hareket, {NU} ürün, {NC} cari. Dolunca boş bir kopyayla yeni döneme başla; stok ve bakiyeleri açılış olarak gir."),
]
r=6
for a,b in rows:
    if a and b is None:
        ws.merge_cells(start_row=r,start_column=2,end_row=r,end_column=3)
        c=ws.cell(r,2,a); c.font=font(10,True,GREEN); c.fill=fill(MINT); c.alignment=Alignment(indent=1,vertical="center"); ws.row_dimensions[r].height=20
    elif a:
        ws.cell(r,2,a).font=font(10,True); ws.cell(r,3,b).font=font(10)
        ws.cell(r,2).alignment=Alignment(vertical="top",indent=1,wrap_text=True); ws.cell(r,3).alignment=Alignment(wrap_text=True,vertical="top")
        ws.row_dimensions[r].height=30 if len(b)>92 else 18
        if a=="Sarı hücre": ws.cell(r,2).fill=fill(INPUT)
    r+=1

# ================= Ayarlar =================
wy=wb.create_sheet("Ayarlar"); wy.sheet_view.showGridLines=False
widths(wy,{"A":3,"B":28,"C":30,"D":3,"E":30})
wy["B1"]="Ayarlar"; wy["B1"].font=font(16,True,GREEN)
wy["B2"]="Sarı hücreleri bir kez doldur."; wy["B2"].font=font(9,c="475467",i=True)
for r_,lab,val,nf in ((4,"İşletme Adı","Örnek Kırtasiye",None),(5,"Açılış Kasa + Banka Bakiyesi",10000,TLIN)):
    a=wy.cell(r_,2,lab); a.font=font(10,True); a.fill=fill(GREY); a.border=box; a.alignment=Alignment(indent=1)
    b=wy.cell(r_,3,val); b.font=font(10,True); b.fill=fill(INPUT); b.border=box
    if nf: b.number_format=nf
wy["B6"]="Dosyayı kullanmaya başladığın gün kasada ve bankada olan toplam para (kullanıcı girişi)."; wy["B6"].font=font(9,c="475467",i=True)
c=wy["E3"]; c.value="Gider Kategorileri"; c.font=font(10,True,WHITE); c.fill=fill(GREEN); c.border=box; c.alignment=Alignment(indent=1)
gk=["Kira","Faturalar","İnternet / Telefon","Personel","Vergi / SGK","Kargo","Reklam / Tanıtım","Ulaşım / Yakıt","Bakım / Onarım","Banka / POS Masrafı","Diğer Gider"]
for i in range(20):
    c=wy.cell(4+i,5,gk[i] if i<len(gk) else None); c.font=font(10); c.fill=fill(INPUT); c.border=box; c.alignment=Alignment(indent=1)
dv(wy,DataValidation(type="decimal",operator="between",formula1="-1000000000",formula2="1000000000"),"C5","Bakiye","Bir sayı gir.")

# ================= Ürünler =================
wu=wb.create_sheet("Ürünler")
widths(wu,{"A":28,"B":9,"C":14,"D":14,"E":12,"F":12,"G":11,"H":11,"I":11,"J":15,"K":11,"L":16,"M":4,"N":4})
title(wu,"Ürünler ve Stok","Sarı sütunları doldur. Giren, çıkan, mevcut stok ve durum Satışlar ile Alışlar sayfalarından hesaplanır.")
header(wu,4,["Ürün","Birim","Alış Maliyeti (₺)","Satış Fiyatı (₺)","Açılış Stoku","Kritik Seviye","Giren","Çıkan","Mevcut","Stok Değeri (₺)","Durum","Toplam Satış (₺)","y1","y2"],formula_cols=range(7,15))
style_rows(wu,NU,14,formula_cols=range(7,15),money=(3,4,10,12),qty=(5,6,7,8,9),center=(2,11))
urunler=[("Defter A4","Adet",35,60,300,60),("Tükenmez Kalem","Adet",8,15,800,150),("Fotokopi Kâğıdı","Paket",140,190,90,25),
         ("Klasör","Adet",45,75,200,40),("Boya Kalemi Seti","Set",90,150,70,10),("Hesap Makinesi","Adet",220,350,8,3)]
for i in range(NU):
    r=R0+i
    if i<len(urunler):
        for j,v in enumerate(urunler[i],1): wu.cell(r,j,v)
    wu.cell(r,7,f'=IF(A{r}="","",SUMIFS({AD},{AC},A{r}))')
    wu.cell(r,8,f'=IF(A{r}="","",SUMIFS({SD},{SC},A{r}))')
    wu.cell(r,9,f'=IF(A{r}="","",E{r}+G{r}-H{r})')
    wu.cell(r,10,f'=IF(A{r}="","",I{r}*C{r})')
    wu.cell(r,11,f'=IF(A{r}="","",IF(I{r}<=0,"Tükendi",IF(I{r}<=F{r},"Kritik","Yeterli")))')
    wu.cell(r,12,f'=IF(A{r}="","",SUMIFS({SG},{SC},A{r}))')
    wu.cell(r,13,f'=IF(A{r}="","",L{r}+ROW()/1000000)')
    wu.cell(r,14,f'=IF(OR(K{r}="Tükendi",K{r}="Kritik"),ROW(),"")')
    wu.cell(r,9).font=font(10,True)
wu.column_dimensions["M"].hidden=True; wu.column_dimensions["N"].hidden=True
wu.freeze_panes="B5"
dv(wu,DataValidation(type="custom",formula1=f"=COUNTIF($A${R0}:$A${last(NU)},A{R0})=1",allow_blank=True),f"A{R0}:A{last(NU)}","Ürün","Bu ürün adı zaten var. Her ürün adı bir kez girilebilir.")
dv(wu,DataValidation(type="decimal",operator="greaterThanOrEqual",formula1="0",allow_blank=True),f"C{R0}:F{last(NU)}","Sayı","Sıfır ya da daha büyük bir sayı gir.")
wu.conditional_formatting.add(f"K{R0}:K{last(NU)}",CellIsRule(operator="equal",formula=['"Tükendi"'],fill=fill("FDE2E0"),font=Font(name=F,bold=True,color=RED)))
wu.conditional_formatting.add(f"K{R0}:K{last(NU)}",CellIsRule(operator="equal",formula=['"Kritik"'],fill=fill("FEF0C7"),font=Font(name=F,bold=True,color="93370D")))
wu.conditional_formatting.add(f"K{R0}:K{last(NU)}",CellIsRule(operator="equal",formula=['"Yeterli"'],font=Font(name=F,color=GREEN)))

# ================= Cariler =================
wc=wb.create_sheet("Cariler")
widths(wc,{"A":28,"B":16,"C":15,"D":15,"E":15,"F":15,"G":16,"H":15,"I":15,"J":16,"K":16,"L":4})
title(wc,"Cariler (Müşteri ve Tedarikçi)","Sarı sütunları doldur. Bakiyeler Satışlar, Alışlar ve Cari Hareket sayfalarından hesaplanır.")
header(wc,4,["Ad","Telefon","Açılış Alacağı (₺)","Açılış Borcu (₺)","Veresiye Satış","Tahsilat","ALACAK BAKİYE","Veresiye Alış","Ödeme","BORÇ BAKİYE","Net (Alacak − Borç)","y1"],formula_cols=range(5,13))
style_rows(wc,NC,12,formula_cols=range(5,13),money=(3,4,5,6,7,8,9,10,11))
cariler=[("Müşteri Alfa",None,2500,None),("Müşteri Beta",None,None,None),("Müşteri Gama",None,None,None),("Tedarikçi Alfa",None,None,1500),("Tedarikçi Beta",None,None,None)]
for i in range(NC):
    r=R0+i
    if i<len(cariler):
        for j,v in enumerate(cariler[i],1): wc.cell(r,j,v)
    wc.cell(r,2).number_format="@"
    wc.cell(r,5,f'=IF(A{r}="","",SUMIFS({SG},{SB},A{r},{SF},"Veresiye"))')
    wc.cell(r,6,f'=IF(A{r}="","",SUMIFS({HD},{HB},"Tahsilat",{HC},A{r}))')
    wc.cell(r,7,f'=IF(A{r}="","",C{r}+E{r}-F{r})')
    wc.cell(r,8,f'=IF(A{r}="","",SUMIFS({AG},{AB},A{r},{AF},"Veresiye"))')
    wc.cell(r,9,f'=IF(A{r}="","",SUMIFS({HD},{HB},"Ödeme",{HC},A{r}))')
    wc.cell(r,10,f'=IF(A{r}="","",D{r}+H{r}-I{r})')
    wc.cell(r,11,f'=IF(A{r}="","",G{r}-J{r})')
    wc.cell(r,12,f'=IF(A{r}="","",G{r}+ROW()/1000000)')
    for j in (7,10): wc.cell(r,j).font=font(10,True)
wc.column_dimensions["L"].hidden=True
wc.freeze_panes="B5"
dv(wc,DataValidation(type="custom",formula1=f"=COUNTIF($A${R0}:$A${last(NC)},A{R0})=1",allow_blank=True),f"A{R0}:A{last(NC)}","Cari","Bu ad zaten var. Her cari adı bir kez girilebilir.")
dv(wc,DataValidation(type="decimal",operator="greaterThanOrEqual",formula1="0",allow_blank=True),f"C{R0}:D{last(NC)}","Sayı","Sıfır ya da daha büyük bir sayı gir.")
wc.conditional_formatting.add(f"K{R0}:K{last(NC)}",CellIsRule(operator="lessThan",formula=["0"],font=Font(name=F,bold=True,color=RED)))

LIST_U=f"='Ürünler'!$A${R0}:$A${last(NU)}"; LIST_C=f"=Cariler!$A${R0}:$A${last(NC)}"
def price(col_letter_src,r): return f"IFERROR(INDEX({col_letter_src},MATCH(C{r},{UA},0)),0)"

# ================= Satışlar =================
wsat=wb.create_sheet("Satışlar")
widths(wsat,{"A":12,"B":24,"C":26,"D":9,"E":14,"F":11,"G":15,"H":15,"I":15,"J":26})
title(wsat,"Satışlar","Her satış bir satır. Birim fiyatı boş bırakırsan liste fiyatı kullanılır. Veresiye satışta müşteri seç.")
header(wsat,4,["Tarih","Müşteri","Ürün","Adet","Birim Fiyat (₺) isteğe bağlı","Ödeme","Tutar (₺)","Maliyet (₺)","Kâr (₺)","Not"],formula_cols=(7,8,9))
style_rows(wsat,NS,10,formula_cols=(7,8,9),money=(5,7,8,9),qty=(4,),dates=(1,),center=(6,))
satis=[(date(2026,9,2),None,"Defter A4",80,None,"Peşin"),(date(2026,9,5),"Müşteri Gama","Tükenmez Kalem",400,13,"Veresiye"),
 (date(2026,9,9),None,"Fotokopi Kâğıdı",20,None,"Peşin"),(date(2026,9,14),"Müşteri Beta","Klasör",120,None,"Veresiye"),
 (date(2026,9,20),None,"Boya Kalemi Seti",24,None,"Peşin"),(date(2026,9,27),"Müşteri Alfa","Fotokopi Kâğıdı",80,180,"Veresiye"),
 (date(2026,10,1),None,"Hesap Makinesi",6,None,"Peşin"),(date(2026,10,2),None,"Defter A4",140,None,"Peşin"),
 (date(2026,10,3),"Müşteri Gama","Boya Kalemi Seti",40,140,"Veresiye"),(date(2026,10,4),None,"Tükenmez Kalem",240,None,"Peşin")]
for i in range(NS):
    r=R0+i
    if i<len(satis):
        for j,v in enumerate(satis[i],1): wsat.cell(r,j,v)
        wsat.cell(r,10,"Örnek kayıt")
    wsat.cell(r,7,f'=IF(OR(C{r}="",D{r}=""),"",D{r}*IF(E{r}="",{price(UD,r)},E{r}))')
    wsat.cell(r,8,f'=IF(G{r}="","",D{r}*{price(UC,r)})')
    wsat.cell(r,9,f'=IF(G{r}="","",G{r}-H{r})')
wsat.freeze_panes="A5"
L=last(NS)
dv(wsat,DataValidation(type="date",operator="greaterThan",formula1="36526",allow_blank=True),f"A{R0}:A{L}","Tarih","Geçerli bir tarih gir (örnek: 04.10.2026).")
dv(wsat,DataValidation(type="list",formula1=LIST_C,allow_blank=True),f"B{R0}:B{L}","Müşteri","Listeden seç. Yeni müşteri için önce Cariler sayfasına ekle.")
dv(wsat,DataValidation(type="list",formula1=LIST_U,allow_blank=True),f"C{R0}:C{L}","Ürün","Listeden seç. Yeni ürün için önce Ürünler sayfasına ekle.")
dv(wsat,DataValidation(type="decimal",operator="greaterThan",formula1="0",allow_blank=True),f"D{R0}:E{L}","Sayı","Sıfırdan büyük bir sayı gir.")
dv(wsat,DataValidation(type="list",formula1='"Peşin,Veresiye"',allow_blank=True),f"F{R0}:F{L}","Ödeme","Peşin ya da Veresiye seç.")
wsat.conditional_formatting.add(f"A{R0}:F{L}",FormulaRule(formula=[f'OR(AND($F{R0}="Veresiye",$B{R0}=""),AND($C{R0}<>"",OR($A{R0}="",$D{R0}="",$F{R0}="")))'],fill=fill("FDE2E0")))
wsat.conditional_formatting.add(f"I{R0}:I{L}",CellIsRule(operator="lessThan",formula=["0"],font=Font(name=F,bold=True,color=RED)))

# ================= Alışlar =================
wal=wb.create_sheet("Alışlar")
widths(wal,{"A":12,"B":24,"C":26,"D":9,"E":14,"F":11,"G":15,"H":26})
title(wal,"Alışlar","Her mal alımı bir satır. Birim maliyeti boş bırakırsan Ürünler sayfasındaki alış maliyeti kullanılır.")
header(wal,4,["Tarih","Tedarikçi","Ürün","Adet","Birim Maliyet (₺) isteğe bağlı","Ödeme","Tutar (₺)","Not"],formula_cols=(7,))
style_rows(wal,NA,8,formula_cols=(7,),money=(5,7),qty=(4,),dates=(1,),center=(6,))
alis=[(date(2026,9,8),"Tedarikçi Alfa","Fotokopi Kâğıdı",30,None,"Veresiye"),(date(2026,9,22),"Tedarikçi Beta","Defter A4",100,34,"Peşin"),(date(2026,10,2),"Tedarikçi Alfa","Tükenmez Kalem",200,None,"Veresiye")]
for i in range(NA):
    r=R0+i
    if i<len(alis):
        for j,v in enumerate(alis[i],1): wal.cell(r,j,v)
        wal.cell(r,8,"Örnek kayıt")
    wal.cell(r,7,f'=IF(OR(C{r}="",D{r}=""),"",D{r}*IF(E{r}="",{price(UC,r)},E{r}))')
wal.freeze_panes="A5"; L=last(NA)
dv(wal,DataValidation(type="date",operator="greaterThan",formula1="36526",allow_blank=True),f"A{R0}:A{L}","Tarih","Geçerli bir tarih gir (örnek: 04.10.2026).")
dv(wal,DataValidation(type="list",formula1=LIST_C,allow_blank=True),f"B{R0}:B{L}","Tedarikçi","Listeden seç. Yeni tedarikçi için önce Cariler sayfasına ekle.")
dv(wal,DataValidation(type="list",formula1=LIST_U,allow_blank=True),f"C{R0}:C{L}","Ürün","Listeden seç. Yeni ürün için önce Ürünler sayfasına ekle.")
dv(wal,DataValidation(type="decimal",operator="greaterThan",formula1="0",allow_blank=True),f"D{R0}:E{L}","Sayı","Sıfırdan büyük bir sayı gir.")
dv(wal,DataValidation(type="list",formula1='"Peşin,Veresiye"',allow_blank=True),f"F{R0}:F{L}","Ödeme","Peşin ya da Veresiye seç.")
wal.conditional_formatting.add(f"A{R0}:F{L}",FormulaRule(formula=[f'OR(AND($F{R0}="Veresiye",$B{R0}=""),AND($C{R0}<>"",OR($A{R0}="",$D{R0}="",$F{R0}="")))'],fill=fill("FDE2E0")))

# ================= Giderler =================
wg=wb.create_sheet("Giderler")
widths(wg,{"A":12,"B":26,"C":40,"D":16})
title(wg,"Giderler","Mal alımı dışındaki harcamalar. Mal alımını Alışlar sayfasına gir.")
header(wg,4,["Tarih","Kategori","Açıklama","Tutar (₺)"])
style_rows(wg,NG,4,money=(4,),dates=(1,))
gider=[(date(2026,9,1),"Kira","Örnek: Eylül kirası",6000),(date(2026,9,12),"Faturalar","Örnek: Elektrik",1250),(date(2026,9,25),"Kargo","Örnek: Kargo ücreti",380),
       (date(2026,10,1),"Kira","Örnek: Ekim kirası",6000),(date(2026,10,3),"İnternet / Telefon","Örnek: İnternet faturası",420)]
for i,row in enumerate(gider):
    for j,v in enumerate(row,1): wg.cell(R0+i,j,v)
wg.freeze_panes="A5"; L=last(NG)
dv(wg,DataValidation(type="date",operator="greaterThan",formula1="36526",allow_blank=True),f"A{R0}:A{L}","Tarih","Geçerli bir tarih gir (örnek: 04.10.2026).")
dv(wg,DataValidation(type="list",formula1="=Ayarlar!$E$4:$E$23",allow_blank=True),f"B{R0}:B{L}","Kategori","Listeden seç. Yeni kategori için Ayarlar sayfasına ekle.")
dv(wg,DataValidation(type="decimal",operator="greaterThan",formula1="0",allow_blank=True),f"D{R0}:D{L}","Tutar","Sıfırdan büyük bir sayı gir.")
wg.conditional_formatting.add(f"A{R0}:D{L}",FormulaRule(formula=[f'AND($D{R0}<>"",$A{R0}="")'],fill=fill("FDE2E0")))

# ================= Cari Hareket =================
wh=wb.create_sheet("Cari Hareket")
widths(wh,{"A":12,"B":13,"C":28,"D":16,"E":40})
title(wh,"Cari Hareket","Tahsilat: veresiye müşteriden para aldım. Ödeme: tedarikçiye veresiye borcumu ödedim.")
header(wh,4,["Tarih","Tür","Kişi / Firma","Tutar (₺)","Açıklama"])
style_rows(wh,NH,5,money=(4,),dates=(1,),center=(2,))
har=[(date(2026,9,20),"Tahsilat","Müşteri Gama",3000,"Örnek kayıt"),(date(2026,10,1),"Ödeme","Tedarikçi Alfa",2000,"Örnek kayıt"),(date(2026,10,3),"Tahsilat","Müşteri Beta",4000,"Örnek kayıt")]
for i,row in enumerate(har):
    for j,v in enumerate(row,1): wh.cell(R0+i,j,v)
wh.freeze_panes="A5"; L=last(NH)
dv(wh,DataValidation(type="date",operator="greaterThan",formula1="36526",allow_blank=True),f"A{R0}:A{L}","Tarih","Geçerli bir tarih gir (örnek: 04.10.2026).")
dv(wh,DataValidation(type="list",formula1='"Tahsilat,Ödeme"',allow_blank=True),f"B{R0}:B{L}","Tür","Tahsilat ya da Ödeme seç.")
dv(wh,DataValidation(type="list",formula1=LIST_C,allow_blank=True),f"C{R0}:C{L}","Kişi","Listeden seç. Yeni kişi için önce Cariler sayfasına ekle.")
dv(wh,DataValidation(type="decimal",operator="greaterThan",formula1="0",allow_blank=True),f"D{R0}:D{L}","Tutar","Sıfırdan büyük bir sayı gir.")
wh.conditional_formatting.add(f"B{R0}:B{L}",CellIsRule(operator="equal",formula=['"Tahsilat"'],font=Font(name=F,bold=True,color=GREEN)))
wh.conditional_formatting.add(f"B{R0}:B{L}",CellIsRule(operator="equal",formula=['"Ödeme"'],font=Font(name=F,bold=True,color=RED)))
wh.conditional_formatting.add(f"A{R0}:D{L}",FormulaRule(formula=[f'AND($D{R0}<>"",OR($A{R0}="",$B{R0}="",$C{R0}=""))'],fill=fill("FDE2E0")))

# ================= Panel =================
wp=wb.create_sheet("Panel"); wp.sheet_view.showGridLines=False
widths(wp,{"A":3,"B":15,"C":17,"D":17,"E":17,"F":17,"G":17,"H":3,"I":28,"J":17,"K":13})
wp["B1"]='=Ayarlar!C4&"  |  CHECK-UP"'; wp["B1"].font=font(16,True,GREEN)
wp["B3"]="Yıl"; wp["B3"].font=font(10,True); wp["B3"].alignment=Alignment(horizontal="right",indent=1)
c=wp["C3"]; c.value=2026; c.font=font(12,True); c.fill=fill(INPUT); c.border=box; c.alignment=Alignment(horizontal="center"); c.number_format="0"
wp["D3"]="← Bu sayfada yalnızca yılı değiştir"; wp["D3"].font=font(9,c="475467",i=True)
dv(wp,DataValidation(type="whole",operator="between",formula1="2000",formula2="2100"),"C3","Yıl","2000 ile 2100 arasında bir yıl yaz.")
def cards(row,label,items):
    wp.cell(row,2,label).font=font(11,True,GREEN)
    for j,(lab,f,nf,red) in enumerate(items):
        a=wp.cell(row+1,2+j,lab); a.font=font(9,True,WHITE); a.fill=fill(GREEN); a.alignment=Alignment(horizontal="center"); a.border=box
        b=wp.cell(row+2,2+j,f); b.font=font(12,True,RED if red else DARK); b.fill=fill(MINT); b.alignment=Alignment(horizontal="center",vertical="center"); b.number_format=nf; b.border=box
    wp.row_dimensions[row+2].height=28
PCT='0.0%;[Red]-0.0%;"-"'
cards(5,"Seçilen yıl",[("Ciro","=C28",TL,0),("Satış Maliyeti","=D28",TL,0),("Brüt Kâr","=E28",TL,0),("Giderler","=F28",TL,0),("Net Kâr","=G28",TL,0),("Net Kâr / Ciro","=IF(C28=0,0,G28/C28)",PCT,0)])
kasa=(f"=Ayarlar!C5+SUMIFS({SG},{SF},\"Peşin\")+SUMIFS({HD},{HB},\"Tahsilat\")-SUMIFS({AG},{AF},\"Peşin\")-SUMIFS({HD},{HB},\"Ödeme\")-SUM({GD})")
cards(9,"Şu anki durum (tüm kayıtlar)",[("Kasa + Banka",kasa,TL,0),("Toplam Alacak",f"=SUM({CG})",TL,0),("Toplam Borç",f"=SUM({CJ})",TL,1),("Stok Değeri",f"=SUM({UJ})",TL,0),
      ("Kritik / Tükenen Ürün",f'=COUNTIF({UK},"Kritik")+COUNTIF({UK},"Tükendi")','0;-0;"-"',1),("Net Durum","=B11+C11-D11+E11",TL,0)])
wp["B12"]="Net Durum = Kasa + Banka + Alacak − Borç + Stok Değeri"; wp["B12"].font=font(8,c="475467",i=True)
wp["B14"]="Aylık Kâr Zarar"; wp["B14"].font=font(11,True,GREEN)
for j,h in enumerate(["Ay","Satış","Maliyet","Brüt Kâr","Gider","Net Kâr"]):
    c=wp.cell(15,2+j,h); c.font=font(10,True,WHITE); c.fill=fill(GREEN); c.border=box; c.alignment=Alignment(horizontal="center")
aylar=["Ocak","Şubat","Mart","Nisan","Mayıs","Haziran","Temmuz","Ağustos","Eylül","Ekim","Kasım","Aralık"]
def monthly(val,dt,m): return f'=SUMIFS({val},{dt},">="&DATE($C$3,{m},1),{dt},"<"&DATE($C$3,{m+1},1))'
for m,ay in enumerate(aylar,1):
    r=15+m
    wp.cell(r,2,ay); wp.cell(r,3,monthly(SG,SA,m)); wp.cell(r,4,monthly(SH,SA,m)); wp.cell(r,5,f"=C{r}-D{r}"); wp.cell(r,6,monthly(GD,GA,m)); wp.cell(r,7,f"=E{r}-F{r}")
    for col in range(2,8):
        c=wp.cell(r,col); c.font=font(10); c.border=box
        if col>2: c.number_format=TL
        if m%2==0: c.fill=fill(GREY)
    wp.cell(r,2).alignment=Alignment(indent=1)
wp.cell(28,2,"TOPLAM")
for col,L_ in zip(range(3,8),"CDEFG"): wp.cell(28,col,f"=SUM({L_}16:{L_}27)")
for col in range(2,8):
    c=wp.cell(28,col); c.font=font(10,True); c.fill=fill(MINT); c.border=box
    if col>2: c.number_format=TL
wp.cell(28,2).alignment=Alignment(indent=1)
wp.conditional_formatting.add("E16:E28",CellIsRule(operator="lessThan",formula=["0"],font=Font(name=F,bold=True,color=RED)))
wp.conditional_formatting.add("G16:G28",CellIsRule(operator="lessThan",formula=["0"],font=Font(name=F,bold=True,color=RED)))
wp.conditional_formatting.add("F7",CellIsRule(operator="lessThan",formula=["0"],font=Font(name=F,size=12,bold=True,color=RED)))
def side(row,ttl,heads):
    wp.cell(row,9,ttl).font=font(11,True,GREEN)
    for j,h in enumerate(heads):
        c=wp.cell(row+1,9+j,h); c.font=font(10,True,WHITE); c.fill=fill(GREEN); c.border=box; c.alignment=Alignment(horizontal="center")
def side_style(r,i,n,fmt2=TL,fmt3=None):
    for j in range(n):
        c=wp.cell(r,9+j); c.font=font(10); c.border=box
        if i%2==1: c.fill=fill(GREY)
    wp.cell(r,9).alignment=Alignment(indent=1); wp.cell(r,10).number_format=fmt2
    if n==3: wp.cell(r,11).alignment=Alignment(horizontal="center")
side(14,"En Çok Satan 5 Ürün (tüm kayıtlar)",["Ürün","Satış (₺)"])
for k in range(1,6):
    r=15+k
    wp.cell(r,9,f'=IFERROR(IF(LARGE({UM},{k})<1,"",INDEX({UA},MATCH(LARGE({UM},{k}),{UM},0))),"")')
    wp.cell(r,10,f'=IF(I{r}="","",INDEX({UL},MATCH(I{r},{UA},0)))'); side_style(r,k,2)
side(22,"Stok Uyarısı (ilk 8)",["Ürün","Mevcut","Durum"])
for k in range(1,9):
    r=23+k
    wp.cell(r,9,f"=IFERROR(INDEX('Ürünler'!$A$1:$A${last(NU)},SMALL({UN},{k})),\"\")")
    wp.cell(r,10,f'=IF(I{r}="","",INDEX({UI},MATCH(I{r},{UA},0)))')
    wp.cell(r,11,f'=IF(I{r}="","",INDEX({UK},MATCH(I{r},{UA},0)))'); side_style(r,k,3,QTY)
wp.conditional_formatting.add("K24:K31",CellIsRule(operator="equal",formula=['"Tükendi"'],font=Font(name=F,bold=True,color=RED)))
wp.conditional_formatting.add("K24:K31",CellIsRule(operator="equal",formula=['"Kritik"'],font=Font(name=F,bold=True,color="93370D")))
side(33,"En Çok Alacaklı Olduğun 5 Kişi",["Kişi / Firma","Alacak (₺)"])
for k in range(1,6):
    r=34+k
    wp.cell(r,9,f'=IFERROR(IF(LARGE({CL},{k})<1,"",INDEX({CA},MATCH(LARGE({CL},{k}),{CL},0))),"")')
    wp.cell(r,10,f'=IF(I{r}="","",INDEX({CG},MATCH(I{r},{CA},0)))'); side_style(r,k,2)
ch=BarChart(); ch.type="col"; ch.title="Aylık Satış, Gider ve Net Kâr"; ch.height=8; ch.width=19.5
ch.add_data(Reference(wp,min_col=3,min_row=15,max_row=27),titles_from_data=True)
ch.add_data(Reference(wp,min_col=6,min_row=15,max_row=27),titles_from_data=True)
ch.add_data(Reference(wp,min_col=7,min_row=15,max_row=27),titles_from_data=True)
ch.set_categories(Reference(wp,min_col=2,min_row=16,max_row=27))
for s,col in zip(ch.series,(GREEN,"D92D20","F79009")): s.graphicalProperties.solidFill=col
ch.legend.position="b"; ch.y_axis.numFmt='#,##0'; ch.y_axis.majorGridlines=None; ch.x_axis.delete=False; ch.y_axis.delete=False
wp.add_chart(ch,"B30")

# ================= AI Komutları =================
wa=wb.create_sheet("AI Komutları"); wa.sheet_view.showGridLines=False
widths(wa,{"A":3,"B":26,"C":96})
wa["B1"]="Yapay Zeka Komutları"; wa["B1"].font=font(16,True,GREEN)
wa.merge_cells("B2:C2"); wa["B2"]="İlgili sayfada başlık satırıyla birlikte kayıtlarını seç, kopyala, kullandığın yapay zeka asistanına yapıştır ve altına komutu ekle. Hangi sayfanın yapıştırılacağı her komutta yazıyor."
wa["B2"].font=font(10,c="475467"); wa["B2"].alignment=Alignment(wrap_text=True,vertical="top"); wa.row_dimensions[2].height=30
wa.merge_cells("B3:C3"); wa["B3"]="Dikkat: Müşteri ve tedarikçi adlarını yapıştırmadan önce Müşteri Alfa, Tedarikçi Alfa gibi takma adlarla değiştir; telefon paylaşma. Yapay zekanın verdiği rakamları Panel sayfasıyla karşılaştır."
wa["B3"].font=font(10,True,RED); wa["B3"].alignment=Alignment(wrap_text=True,vertical="top"); wa.row_dimensions[3].height=30
for j,h in enumerate(["Ne için","Komut (kopyala, yapıştır)"]):
    c=wa.cell(5,2+j,h); c.font=font(10,True,WHITE); c.fill=fill(GREEN); c.border=box; c.alignment=Alignment(indent=1)
prompts=[
 ("Aylık durum raporu","[Panel sayfasındaki Aylık Kâr Zarar tablosunu yapıştır] Bu tablo dükkânımın aylık satış, maliyet, gider ve net kârı. Durumu beş cümleyle özetle, en iyi ve en kötü ayı nedenleriyle belirt."),
 ("En kârlı ürünler","[Satışlar sayfasını yapıştır] Ürün bazında toplam ciroyu, toplam kârı ve kâr oranını hesapla. Sonucu kâra göre sıralı tablo olarak ver."),
 ("Zam önerisi","[Ürünler sayfasını yapıştır] Alış maliyeti ile satış fiyatı arasındaki farkı en düşük olan ürünleri bul. Her biri için yüzde 10 zam yaparsam birim kârım kaç lira artar, göster."),
 ("Sipariş listesi","[Ürünler sayfasını yapıştır] Durumu Kritik ya da Tükendi olan ürünler için tedarikçiye göndereceğim sipariş listesini hazırla. Her üründen kritik seviyenin üç katına çıkacak kadar adet öner."),
 ("Ölü stok","[Ürünler ve Satışlar sayfalarını yapıştır] Stokta olup son 60 günde hiç satılmamış ya da çok az satılmış ürünleri listele. Bu ürünlere bağlı parayı Stok Değeri sütunundan topla."),
 ("Veresiye hatırlatma","[Cariler sayfasını yapıştır] Alacak bakiyesi olan her müşteri için kısa ve kibar bir ödeme hatırlatma mesajı yaz. Mesajda bakiye tutarı geçsin."),
 ("Gider kontrolü","[Giderler sayfasını yapıştır] Giderlerimi kategoriye göre topla, ciroya oranla yüksek görünenleri belirt ve azaltmak için iki somut öneri ver."),
 ("En iyi müşteriler","[Satışlar sayfasını yapıştır] Müşteri bazında toplam alışveriş tutarını ve alışveriş sayısını çıkar. İlk beş müşteriyi listele."),
 ("Satış günleri","[Satışlar sayfasını yapıştır] Haftanın hangi günlerinde daha çok satış yapıyorum? Gün gün toplam ciroyu ver ve en zayıf gün için bir kampanya fikri öner."),
 ("Kayıt kontrolü","[Satışlar ve Alışlar sayfalarını yapıştır] Eksik ya da tutarsız görünen satır var mı? Tarihi boş, adedi olağan dışı yüksek ya da iki kez girilmiş gibi duranları göster."),
]
for i,(a,b) in enumerate(prompts):
    r=6+i
    wa.cell(r,2,a).font=font(10,True); wa.cell(r,3,b).font=font(10)
    for col in (2,3):
        c=wa.cell(r,col); c.border=box; c.alignment=Alignment(wrap_text=True,vertical="center",indent=1)
        if i%2==1: c.fill=fill(GREY)
    wa.row_dimensions[r].height=32

order=["Başla","Panel","Satışlar","Alışlar","Giderler","Cari Hareket","Ürünler","Cariler","Ayarlar","AI Komutları"]
wb._sheets=[wb[n] for n in order]
for s in wb.worksheets:
    s.sheet_properties.tabColor=GREEN if s.title in ("Başla","Panel","AI Komutları") else ("F79009" if s.title in ("Ürünler","Cariler","Ayarlar") else "2E90FA")
    s.page_setup.orientation="landscape"; s.page_setup.fitToWidth=1; s.page_setup.fitToHeight=0
    s.sheet_properties.pageSetUpPr.fitToPage=True
wb.active=0
wb.save("ExcelKlinigi_CHECKUP_Esnaf_Isletme_Defteri.xlsx")
json.dump({"urunler":urunler,"cariler":cariler,"satis":[(str(a),b,c,d,e,f) for a,b,c,d,e,f in satis],"alis":[(str(a),b,c,d,e,f) for a,b,c,d,e,f in alis],
           "gider":[(str(a),b,c,d) for a,b,c,d in gider],"har":[(str(a),b,c,d,e) for a,b,c,d,e in har]},open("samples.json","w"),ensure_ascii=False)
