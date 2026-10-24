import json
from openpyxl import load_workbook
S=json.load(open("samples.json"))
U={u[0]:dict(alis=u[2],satis=u[3],acilis=u[4],kritik=u[5]) for u in S["urunler"]}
ciro=cogs=0; pes_s=0; ver_s={}; cikan={}; ursat={}; mon={}
for t,mu,ur,ad,bf,od in S["satis"]:
    tut=ad*(bf if bf else U[ur]["satis"]); mal=ad*U[ur]["alis"]
    ciro+=tut; cogs+=mal; cikan[ur]=cikan.get(ur,0)+ad; ursat[ur]=ursat.get(ur,0)+tut
    m=int(t[5:7]); a=mon.setdefault(m,[0,0,0]); a[0]+=tut; a[1]+=mal
    if od=="Peşin": pes_s+=tut
    else: ver_s[mu]=ver_s.get(mu,0)+tut
pes_a=0; ver_a={}; giren={}
for t,td,ur,ad,bm,od in S["alis"]:
    tut=ad*(bm if bm else U[ur]["alis"]); giren[ur]=giren.get(ur,0)+ad
    if od=="Peşin": pes_a+=tut
    else: ver_a[td]=ver_a.get(td,0)+tut
gid=0
for t,k,a,tu in S["gider"]:
    gid+=tu; mon.setdefault(int(t[5:7]),[0,0,0])[2]+=tu
tah={}; ode={}
for t,tur,k,tu,_ in S["har"]:
    (tah if tur=="Tahsilat" else ode)[k]=(tah if tur=="Tahsilat" else ode).get(k,0)+tu
alacak={c[0]:(c[2] or 0)+ver_s.get(c[0],0)-tah.get(c[0],0) for c in S["cariler"]}
borc={c[0]:(c[3] or 0)+ver_a.get(c[0],0)-ode.get(c[0],0) for c in S["cariler"]}
stok={u:U[u]["acilis"]+giren.get(u,0)-cikan.get(u,0) for u in U}
stokdeg=sum(stok[u]*U[u]["alis"] for u in U)
kasa=10000+pes_s+sum(tah.values())-pes_a-sum(ode.values())-gid
kritik=[u for u in U if stok[u]<=U[u]["kritik"]]
exp=dict(ciro=ciro,cogs=cogs,brut=ciro-cogs,gider=gid,net=ciro-cogs-gid,kasa=kasa,alacak=sum(alacak.values()),borc=sum(borc.values()),stok=stokdeg,kritik=len(kritik))
exp["netdurum"]=kasa+exp["alacak"]-exp["borc"]+stokdeg
wb=load_workbook("ExcelKlinigi_CHECKUP_Esnaf_Isletme_Defteri.xlsx",data_only=True); p=wb["Panel"]
got=dict(ciro=p["B7"].value,cogs=p["C7"].value,brut=p["D7"].value,gider=p["E7"].value,net=p["F7"].value,kasa=p["B11"].value,alacak=p["C11"].value,borc=p["D11"].value,stok=p["E11"].value,kritik=p["F11"].value,netdurum=p["G11"].value)
ok=True
for k in exp:
    f=abs(exp[k]-got[k])<0.005; ok&=f; print(f"{k:9} beklenen={exp[k]:>10} dosya={got[k]:>10} {'OK' if f else 'HATA'}")
print("baslik:",p["B1"].value,"| marj:",round(p["G7"].value,4))
for m in (9,10):
    row=[p.cell(15+m,c).value for c in range(3,8)]; e=[mon[m][0],mon[m][1],mon[m][0]-mon[m][1],mon[m][2],mon[m][0]-mon[m][1]-mon[m][2]]
    print(m,row,e,"OK" if row==e else "HATA"); ok&=row==e
print("top5:",[(p.cell(r,9).value,p.cell(r,10).value) for r in range(16,21)])
print("beklenen top:",sorted(ursat.items(),key=lambda x:-x[1])[:5])
print("uyarı:",[(p.cell(r,9).value,p.cell(r,10).value,p.cell(r,11).value) for r in range(24,28)]," beklenen:",[(u,stok[u]) for u in kritik])
print("alacaklı:",[(p.cell(r,9).value,p.cell(r,10).value) for r in range(35,40)]," beklenen:",{k:v for k,v in alacak.items() if v>0})
print("borç:",borc)
print("TÜMÜ OK" if ok else "FARK VAR")
