# REÇETE (teklif), DOZ (düzenli ödemeler), DİYET (bütçe)
from lib import *
from datetime import date
from openpyxl import load_workbook

def recete():
    B=Book("REÇETE","Teklif hazırlama","Ürünü seç, adedi yaz; fiyat, iskonto, KDV ve toplam kendiliğinden çıksın. Yazdır ya da PDF olarak gönder.")
    NU,NM,ND,NL=200,100,500,20; U="Ürünler"; M="Müşteriler"; D="Teklif Defteri"
    wa=B.sheet("Ayarlar","Ayarlar","Firma bilgilerin teklifin üstünde görünür.",{"A":3,"B":30,"C":44},grid=False)
    for r,lab,val,nf in ((4,"Firma adı","Örnek Ticaret",None),(5,"Adres","Örnek Mah. Örnek Sok. No: 1, İstanbul",None),(6,"Telefon","0 (000) 000 00 00",None),(7,"E-posta","ornek@ornek.com",None),(8,"Varsayılan KDV oranı",0.20,"0%"),(9,"Teklif geçerlilik süresi (gün)",15,"0"),(10,"Teklif altı notu","Fiyatlara KDV dahil değildir; KDV ayrıca gösterilmiştir. Teslimat ve ödeme koşulları ayrıca kararlaştırılır.",None)):
        a=wa.cell(r,2,lab); a.font=font(10,True); a.fill=fill(GREY); a.border=box; b=wa.cell(r,3,val); b.font=font(10); b.fill=fill(INPUT); b.border=box; b.alignment=Alignment(wrap_text=True,vertical="top")
        if nf: b.number_format=nf
    wa.row_dimensions[10].height=44; wa["B12"]="KDV oranı kullanıcı girişidir. Sattığın ürün ya da hizmetin oranını mali müşavirinden doğrula."; wa["B12"].font=font(9,c="475467",i=True)
    wu=B.sheet(U,"Ürünler ve Hizmetler","Teklifte kullanacağın kalemler. KDV oranını boş bırakırsan Ayarlar'daki oran kullanılır.",{"A":34,"B":10,"C":18,"D":12})
    B.head(wu,4,["Ürün ya da hizmet","Birim","Birim fiyat (₺, KDV hariç)","KDV oranı"]); B.body(wu,R0,NU,4,fmt={3:TLI,4:"0%"},center=(2,4))
    urun=[("Web sitesi kurulumu","Adet",25000,None),("Aylık bakım","Ay",3500,None),("Logo tasarımı","Adet",8000,None),("Eğitim (yarım gün)","Seans",6000,0.10),("Basılı katalog","Adet",45,0.10)]
    B.put(wu,R0,urun); B.dv(wu,"uniq",f"A{R0}:A{R0+NU-1}",f"=COUNTIF($A${R0}:$A${R0+NU-1},A{R0})=1"); B.dv(wu,"nonneg",f"C{R0}:C{R0+NU-1}"); B.dv(wu,"pct",f"D{R0}:D{R0+NU-1}")
    wm=B.sheet(M,"Müşteriler","Teklif vereceğin kişi ve firmalar.",{"A":30,"B":44,"C":18,"D":24})
    B.head(wm,4,["Ad","Adres","Telefon","Yetkili"]); B.body(wm,R0,NM,4,fmt={3:"@"}); B.put(wm,R0,[("Örnek Müşteri A.Ş.","Örnek Cad. No: 10, İstanbul",None,"Satın alma sorumlusu"),("Örnek Dükkân",None,None,None)])
    B.dv(wm,"uniq",f"A{R0}:A{R0+NM-1}",f"=COUNTIF($A${R0}:$A${R0+NM-1},A{R0})=1")
    UA,UB,UC,UD=[rg(U,c,NU) for c in "ABCD"]; MA,MB,MD=rg(M,"A",NM),rg(M,"B",NM),rg(M,"D",NM)
    wt=B.wb.create_sheet("Teklif"); wt.sheet_view.showGridLines=False
    for k,v in {"A":5,"B":36,"C":9,"D":9,"E":15,"F":10,"G":9,"H":17,"I":15,"J":4}.items(): wt.column_dimensions[k].width=v
    wt["A1"]="=Ayarlar!C4"; wt["A1"].font=font(16,True,GREEN); wt["A2"]='=Ayarlar!C5&"   "&Ayarlar!C6&"   "&Ayarlar!C7'; wt["A2"].font=font(9,c="475467")
    wt["H1"]="TEKLİF"; wt["H1"].font=font(22,True,GREEN); wt["H1"].alignment=Alignment(horizontal="right"); wt.merge_cells("H1:I1")
    for r,lab in ((4,"Teklif no"),(5,"Tarih"),(6,"Geçerlilik tarihi")): c=wt.cell(r,8,lab); c.font=font(10,True); c.alignment=Alignment(horizontal="right")
    B.inp(wt,"I4","T-2026-001",None); B.inp(wt,"I5",date(2026,10,5),DT); wt["I6"]="=IF(I5=\"\",\"\",I5+Ayarlar!C9)"; wt["I6"].number_format=DT; wt["I6"].font=font(11,True); wt["I6"].alignment=Alignment(horizontal="center"); wt["I6"].border=box
    wt["A4"]="Sayın"; wt["A4"].font=font(10,True); wt.merge_cells("B4:E4"); B.inp(wt,"B4","Örnek Müşteri A.Ş."); wt["B4"].alignment=Alignment(horizontal="left")
    wt.merge_cells("B5:E5"); wt["B5"]=f'=IF(B4="","",IFERROR(INDEX({MB},MATCH(B4,{MA},0))&"",""))'; wt.merge_cells("B6:E6"); wt["B6"]=f'=IF(B4="","",IFERROR(INDEX({MD},MATCH(B4,{MA},0))&"",""))'
    B.head(wt,9,["#","Ürün ya da hizmet","Birim","Adet","Özel birim fiyat (₺)","İskonto","KDV","Tutar (₺, KDV hariç)","KDV tutarı (₺)","y"],calc=(1,3,7,8,9,10),h=30)
    B.body(wt,10,NL,10,calc=(1,3,7,8,9,10),fmt={4:NUM,5:TLI,6:"0%",7:"0%",8:TL,9:TL},center=(1,3,4,6,7))
    B.put(wt,10,[(None,"Web sitesi kurulumu",None,1,None,0.10),(None,"Aylık bakım",None,6,3000,None),(None,"Eğitim (yarım gün)",None,2,None,None)])
    for i in range(NL):
        r=10+i; lp=f"IFERROR(INDEX({UC},MATCH(B{r},{UA},0)),0)"
        wt.cell(r,1,f'=IF(B{r}="","",{i+1})'); wt.cell(r,3,f'=IF(B{r}="","",IFERROR(INDEX({UB},MATCH(B{r},{UA},0))&"",""))')
        wt.cell(r,7,f'=IF(B{r}="","",IF(IFERROR(INDEX({UD},MATCH(B{r},{UA},0)),"")="",Ayarlar!$C$8,INDEX({UD},MATCH(B{r},{UA},0))))')
        wt.cell(r,10,f'=IF(OR(B{r}="",D{r}=""),0,D{r}*IF(E{r}="",{lp},E{r}))')
        wt.cell(r,8,f'=IF(OR(B{r}="",D{r}=""),"",J{r}*(1-F{r}))'); wt.cell(r,9,f'=IF(H{r}="","",H{r}*G{r})')
    wt.column_dimensions["J"].hidden=True; e=10+NL
    for k,(lab,f_) in enumerate((("Ara toplam",f"=SUM(J10:J{e-1})"),("İskonto",f"=SUM(J10:J{e-1})-SUM(H10:H{e-1})"),("KDV hariç toplam",f"=SUM(H10:H{e-1})"),("KDV",f"=SUM(I10:I{e-1})"),("GENEL TOPLAM",f"=SUM(H10:H{e-1})+SUM(I10:I{e-1})"))):
        r=e+1+k; a=wt.cell(r,8,lab); a.font=font(10,True); a.alignment=Alignment(horizontal="right"); b=wt.cell(r,9,f_); b.number_format=TL; b.font=font(11 if k<4 else 13,True,DARK if k<4 else GREEN); b.border=box
        if k==4: b.fill=fill(MINT)
    wt.merge_cells(start_row=e+1,start_column=1,end_row=e+4,end_column=6); c=wt.cell(e+1,1,"=Ayarlar!C10"); c.font=font(9,c="475467"); c.alignment=Alignment(wrap_text=True,vertical="top")
    wt.cell(e+8,2,"Teklifi veren").font=font(10,True); wt.cell(e+8,8,"Teklifi kabul eden").font=font(10,True)
    B.dv(wt,"list","B4","="+MA); B.dv(wt,"list",f"B10:B{e-1}","="+UA); B.dv(wt,"pos",f"D10:E{e-1}"); B.dv(wt,"pct",f"F10:F{e-1}"); B.dv(wt,"date","I5")
    wt.print_area=f"A1:I{e+9}"; wt.page_setup.orientation="portrait"
    wd=B.sheet(D,"Teklif Defteri","Gönderdiğin her teklifi buraya bir satır olarak yaz. Dosya teklifleri kendiliğinden kaydetmez.",{"A":14,"B":12,"C":30,"D":17,"E":15,"F":34})
    B.head(wd,4,["Teklif no","Tarih","Müşteri","Tutar (₺, KDV dahil)","Durum","Not"]); B.body(wd,R0,ND,6,fmt={2:DT,4:TLI},center=(1,2,5))
    B.put(wd,R0,[("T-2026-001",date(2026,9,4),"Örnek Müşteri A.Ş.",54000,"Kabul edildi","Örnek"),("T-2026-002",date(2026,9,18),"Örnek Dükkân",9600,"Reddedildi","Örnek"),("T-2026-003",date(2026,10,1),"Örnek Müşteri A.Ş.",21600,"Gönderildi","Örnek"),("T-2026-004",date(2026,10,3),"Örnek Dükkân",12000,"Kabul edildi","Örnek")])
    B.dv(wd,"date",f"B{R0}:B{R0+ND-1}"); B.dv(wd,"list",f"C{R0}:C{R0+ND-1}","="+MA); B.dv(wd,"pos",f"D{R0}:D{R0+ND-1}"); B.dv(wd,"list",f"E{R0}:E{R0+ND-1}","Gönderildi,Kabul edildi,Reddedildi,Yanıt yok")
    for t,fc,tc in (("Kabul edildi",MINT,GREEN),("Reddedildi",PINK,RED),("Yanıt yok",GREY,"667085")): B.cf_eq(wd,f"E{R0}:E{R0+ND-1}",t,fc,tc)
    wd.freeze_panes="A5"
    DB,DD,DE=rg(D,"B",ND),rg(D,"D",ND),rg(D,"E",ND)
    wo=B.sheet("Özet","Özet","Bu sayfada yalnızca yılı değiştir. Rakamlar Teklif Defteri'nden gelir.",{"A":3,"B":14,"C":15,"D":17,"E":15,"F":17,"G":15},grid=False)
    wo["B3"]="Yıl"; B.inp(wo,"C3",2026,"0")
    B.cards(wo,5,[("Teklif","=C23",NUM),("Teklif tutarı","=D23",TL),("Kabul edilen","=E23",NUM),("Kabul tutarı","=F23",TL),("Kabul oranı","=IF(C23=0,0,E23/C23)",PCT1),("Yanıt bekleyen tutar",f'=SUMIFS({DD},{DE},"Gönderildi")',TL)])
    B.label(wo,"B9","Aylık tablo"); B.table(wo,10,2,["Ay","Teklif","Tutar (₺)","Kabul","Kabul tutarı (₺)"],13,fmt={3:NUM,4:TL,5:NUM,6:TL})
    for m,ay in enumerate(months(),1):
        r=10+m; wo.cell(r,2,ay); wo.cell(r,3,"="+crange(DB,"$C$3",m)); wo.cell(r,4,"="+mrange(DD,DB,"$C$3",m)); wo.cell(r,5,"="+crange(DB,"$C$3",m,f',{DE},"Kabul edildi"')); wo.cell(r,6,"="+mrange(DD,DB,"$C$3",m,f',{DE},"Kabul edildi"'))
    wo.cell(23,2,"TOPLAM")
    for col,L in zip(range(3,7),"CDEF"): wo.cell(23,col,f"=SUM({L}11:{L}22)"); wo.cell(23,col).font=font(10,True)
    B.basla([("NASIL KULLANILIR",None),("1. Ayarlar","Firma bilgilerini, varsayılan KDV oranını ve teklif geçerlilik süresini yaz."),("2. Ürünler ve Müşteriler","Teklifte kullanacağın kalemleri ve müşterilerini bir kez gir."),("3. Teklif","Müşteriyi ve kalemleri açılır listeden seç, adedi yaz. İstersen özel fiyat ve iskonto gir. Yazdır ya da Dosya > Dışa Aktar ile PDF yap."),("4. Teklif Defteri","Gönderdiğin teklifi bir satır olarak deftere yaz ve sonucunu güncelle. Özet buradan hesaplanır."),
      (None,None),("BİLMEN GEREKENLER",None),("Tek teklif","Teklif sayfası her seferinde bir teklif tutar. Yeni teklif için kalemleri silip yeniden doldur ya da dosyanın kopyasını al."),("Resmî belge değil","Teklif bir fiyat önerisidir; fatura ya da e-fatura yerine geçmez. KDV oranları kullanıcı girişidir."),("Kapasite",f"Bir teklifte {NL} kalem, {NU} ürün, {NM} müşteri, defterde {ND} teklif.")])
    B.ai("Müşteri adlarını takma adla değiştir; fiyatların gizliyse tutarları paylaşma.",[("Kapak yazısı","Aşağıdaki teklif için müşteriye gönderilecek kısa bir e-posta yaz. Teklifin ne içerdiğini iki cümleyle özetle, geçerlilik tarihini belirt."),("Takip mesajı","Bir hafta önce gönderdiğim teklife yanıt gelmedi. Israrcı olmayan, kısa bir takip mesajı yaz."),("İskonto etkisi","[Teklif kalemleri] Toplam iskontoyu yüzde 5 artırırsam genel toplam ve kâr payım ne olur? Birim maliyetleri ben söyleyeceğim, önce sor."),("Reddedilenler","[Teklif Defteri] Reddedilen tekliflerin tutar aralığına ve müşterilerine bak. Ortak bir neden görünüyor mu?"),("Kabul oranı","[Özet sayfasındaki aylık tablo] Kabul oranımı aylara göre yorumla ve artırmak için iki öneri ver."),("Kalem açıklaması","Şu kalem için teklifte kullanılacak iki cümlelik, anlaşılır bir açıklama yaz: [kalem adı].")])
    fn=OUT+"ExcelKlinigi_RECETE_Teklif_Hazirlama.xlsx"; B.save(["Başla","Teklif","Teklif Defteri" if False else D,"Özet",U,M,"Ayarlar","AI Komutları"],fn)
    w=load_workbook(fn,data_only=True); t=w["Teklif"]; o=w["Özet"]
    # 25000*1*0.9=22500 (KDV %20=4500); 6*3000=18000 (3600); 2*6000=12000 (%10=1200)
    assert [t.cell(10+i,8).value for i in range(3)]==[22500,18000,12000]; assert [round(t.cell(10+i,9).value,2) for i in range(3)]==[4500,3600,1200]
    assert [t.cell(31+k,9).value for k in range(5)]==[55000,2500,52500,9300,61800],[t.cell(31+k,9).value for k in range(5)]
    assert (o["B6"].value,o["C6"].value,o["D6"].value,o["E6"].value,o["F6"].value,o["G6"].value)==(4,97200,2,66000,0.5,21600); print("  REÇETE doğrulandı")

def doz():
    B=Book("DOZ","Düzenli ödeme ve abonelik takibi","Kira, fatura, abonelik, taksit: her ay ne kadar çıkıyor, sırada hangisi var, yıl boyunca hangi ay ağır.")
    N=100; O="Ödemeler"; T="'Özet'!$C$3"
    w=B.sheet(O,"Ödemeler","Düzenli ödediğin her şeyi bir satıra yaz. Bitiş tarihini yalnızca taksit gibi sonu belli ödemelerde doldur.",{"A":28,"B":18,"C":14,"D":11,"E":14,"F":14,"G":4,"H":15,"I":15,"J":15,"K":10,"L":12,"M":4,"N":4,"O":4})
    B.head(w,4,["Ödeme","Kategori","Tutar (₺)","Sıklık","İlk ödeme tarihi","Bitiş tarihi (varsa)","ay","Aylık karşılığı (₺)","Yıllık karşılığı (₺)","Sıradaki ödeme","Kalan gün","Durum","y1","y2","y3"],calc=range(7,16))
    B.body(w,R0,N,15,calc=range(7,16),fmt={3:TLI,5:DT,6:DT,8:TL,9:TL,10:DT,11:NUM},center=(4,5,6,10,11,12))
    od=[("Ev kirası","Konut",22000,"Aylık",date(2026,1,5),None),("Elektrik","Faturalar",1400,"Aylık",date(2026,1,20),None),("İnternet","Faturalar",650,"Aylık",date(2026,2,12),None),("Film ve dizi aboneliği","Abonelik",230,"Aylık",date(2025,11,8),None),
        ("Bulut depolama","Abonelik",1200,"Yıllık",date(2026,3,15),None),("Telefon taksidi","Taksit",2500,"Aylık",date(2026,4,10),date(2027,3,10)),("Kasko","Sigorta",9000,"6 Aylık",date(2026,6,1),None),("Spor salonu","Sağlık",4500,"3 Aylık",date(2026,8,1),None),("Eski taksit","Taksit",1000,"Aylık",date(2025,1,10),date(2025,12,10))]
    B.put(w,R0,od)
    for i in range(N):
        r=R0+i
        w.cell(r,7,f'=IF(D{r}="Aylık",1,IF(D{r}="3 Aylık",3,IF(D{r}="6 Aylık",6,IF(D{r}="Yıllık",12,""))))')
        w.cell(r,13,f'=IF(OR(A{r}="",C{r}="",G{r}="",E{r}=""),"",MAX(0,ROUNDUP(((YEAR({T})-YEAR(E{r}))*12+MONTH({T})-MONTH(E{r}))/G{r},0)))')
        w.cell(r,14,f'=IF(M{r}="","",IF(EDATE(E{r},G{r}*M{r})<{T},EDATE(E{r},G{r}*(M{r}+1)),EDATE(E{r},G{r}*M{r})))')
        w.cell(r,10,f'=IF(N{r}="","",IF(AND(F{r}<>"",N{r}>F{r}),"Bitti",N{r}))'); w.cell(r,11,f'=IF(ISNUMBER(J{r}),J{r}-{T},"")')
        w.cell(r,12,f'=IF(J{r}="","",IF(J{r}="Bitti","Bitti",IF(K{r}<=7,"Bu hafta",IF(K{r}<=30,"Bu ay","Bekliyor"))))')
        w.cell(r,8,f'=IF(OR(J{r}="",J{r}="Bitti"),"",C{r}/G{r})'); w.cell(r,9,f'=IF(H{r}="","",H{r}*12)'); w.cell(r,15,f'=IF(ISNUMBER(J{r}),J{r}+ROW()/1000000,"")')
    for c in "GMNO": w.column_dimensions[c].hidden=True
    w.freeze_panes="B5"; B.dv(w,"pos",f"C{R0}:C{R0+N-1}"); B.dv(w,"list",f"D{R0}:D{R0+N-1}","Aylık,3 Aylık,6 Aylık,Yıllık"); B.dv(w,"date",f"E{R0}:F{R0+N-1}"); B.dv(w,"list",f"B{R0}:B{R0+N-1}","='Kategoriler'!$A$5:$A$24")
    for t,fc,tc in (("Bu hafta",PINK,RED),("Bu ay",AMB,AMBT),("Bitti",GREY,"667085")): B.cf_eq(w,f"L{R0}:L{R0+N-1}",t,fc,tc)
    wk=B.sheet("Kategoriler","Kategoriler","Kendi kategorilerini yaz.",{"A":26}); B.head(wk,4,["Kategori"]); B.body(wk,R0,20,1); B.put(wk,R0,[[k] for k in ("Konut","Faturalar","Abonelik","Taksit","Sigorta","Sağlık","Eğitim","Diğer")])
    OA,OB,OC,OE,OF,OG,OH,OI,OJ,OK,OL,OO=[rg(O,c,N) for c in "ABCEFGHIJKLO"]
    wt=B.sheet("Takvim","Yıllık Takvim","Yılı Özet sayfasından değiştir. Her ödemenin hangi ay düştüğünü ve ayların toplam yükünü gösterir.",{"A":28,**{CL(c):12 for c in range(2,15)}})
    B.head(wt,4,["Ödeme"]+months()+["Yıl toplamı"],calc=range(1,15),h=22); B.body(wt,R0,N+1,14,calc=range(1,15),fmt={c:NUM for c in range(2,15)},stripe=True)
    wt.cell(R0,1,"AY TOPLAMI").font=font(10,True)
    for m in range(1,13):
        L=CL(m+1); wt.cell(R0,m+1,f"=SUM({L}{R0+1}:{L}{R0+N})").font=font(10,True)
        for i in range(N):
            r=R0+1+i; s_=R0+i; d=f"(('Özet'!$C$4*12+{m})-(YEAR('{O}'!$E{s_})*12+MONTH('{O}'!$E{s_})))"
            wt.cell(r,m+1,f"=IF(OR('{O}'!$A{s_}=\"\",'{O}'!$G{s_}=\"\",'{O}'!$E{s_}=\"\"),\"\",IF(AND({d}>=0,MOD({d},'{O}'!$G{s_})=0,OR('{O}'!$F{s_}=\"\",DATE('Özet'!$C$4,{m},1)<='{O}'!$F{s_})),'{O}'!$C{s_},0))")
    for i in range(N+1):
        r=R0+i; wt.cell(r,14,f'=IF(COUNT(B{r}:M{r})=0,"",SUM(B{r}:M{r}))')
        if i: wt.cell(r,1,f"=IF('{O}'!A{R0+i-1}=\"\",\"\",'{O}'!A{R0+i-1})")
    wt.freeze_panes="B6"
    wo=B.sheet("Özet","Özet","Bugünün tarihi kendiliğinden gelir. Takvim için yılı değiştirebilirsin.",{"A":3,"B":20,"C":16,"D":16,"E":16,"F":16,"G":3,"H":28,"I":14,"J":14},grid=False)
    wo["B3"]="Bugün"; wo["C3"]="=TODAY()"; wo["C3"].number_format=DT; wo["C3"].font=font(11,True); wo["C3"].border=box; wo["C3"].fill=fill(MINT); wo["C3"].alignment=Alignment(horizontal="center")
    wo["B4"]="Takvim yılı"; B.inp(wo,"C4",2026,"0")
    B.cards(wo,6,[("Aylık ortalama yük",f"=SUM({OH})",TL),("Yıllık karşılığı",f"=SUM({OI})",TL),("7 gün içinde ödenecek",f'=SUMIFS({OC},{OK},">=0",{OK},"<=7")',TL,1),("30 gün içinde ödenecek",f'=SUMIFS({OC},{OK},">=0",{OK},"<=30")',TL),("Süren ödeme sayısı",f"=COUNT({OH})",NUM)])
    B.label(wo,"B10","Kategoriye göre aylık yük"); B.table(wo,11,2,["Kategori","Aylık (₺)","Pay"],12,fmt={3:TL,4:"0%"})
    for k in range(12):
        r=12+k; wo.cell(r,2,f"=IF('Kategoriler'!A{R0+k}=\"\",\"\",'Kategoriler'!A{R0+k})"); wo.cell(r,3,f'=IF(B{r}="","",SUMIFS({OH},{OB},B{r}))'); wo.cell(r,4,f'=IF(OR(B{r}="",$B$7=0),"",C{r}/$B$7)')
    B.label(wo,"H10","Sıradaki 10 ödeme"); B.table(wo,11,8,["Ödeme","Tarih","Tutar (₺)"],10,fmt={9:DT,10:TL},center=(9,))
    for k in range(1,11):
        r=11+k; ix=f"MATCH(SMALL({OO},{k}),{OO},0)"; wo.cell(r,8,f'=IFERROR(INDEX({OA},{ix}),"")'); wo.cell(r,9,f'=IF(H{r}="","",INDEX({OJ},{ix}))'); wo.cell(r,10,f'=IF(H{r}="","",INDEX({OC},{ix}))')
    chart(wo,"B26","Aylara göre ödeme yükü (₺)",Reference(wt,min_col=2,max_col=13,min_row=4),[],[GREEN],w=16)
    ch=wo._charts[0]; ch.add_data(Reference(wt,min_col=2,max_col=13,min_row=R0),from_rows=True,titles_from_data=False); ch.series[0].graphicalProperties.solidFill=GREEN; ch.legend=None
    B.basla([("NASIL KULLANILIR",None),("1. Ödemeler","Kira, fatura, abonelik, sigorta, taksit: düzenli ödediğin her şeyi tutarı, sıklığı ve ilk ödeme tarihiyle yaz."),("2. Özet","Aylık ortalama yükün, önümüzdeki 7 ve 30 günde ödeyeceklerin, kategori dağılımı ve sıradaki 10 ödeme."),("3. Takvim","Seçtiğin yılda her ödemenin hangi aya düştüğü ve ayların toplamı. Yıllık ve altı aylık ödemelerin ağırlaştırdığı ayları buradan görürsün."),
      (None,None),("BİLMEN GEREKENLER",None),("Aylık karşılığı","Yıllık 1.200 liralık bir ödemenin aylık karşılığı 100 liradır. Bitmiş ödemeler toplama girmez."),("Sıradaki ödeme","İlk ödeme tarihinden sıklığa göre ileri sayılır. Ayın 31'i gibi günlerde kısa aylarda ayın son gününe kayar."),("Tutar değişirse","Fatura gibi değişen ödemelerde ortalama tutarı yaz. Dosya geçmiş tutarları ayrı ayrı tutmaz."),("Kapasite",f"{N} ödeme, 20 kategori.")])
    B.ai("Sözleşme ve abone numaralarını paylaşma.",[("Abonelik temizliği","[Ödemeler sayfası] Abonelik kategorisindeki ödemeleri yıllık karşılığına göre sırala. Hangilerini iptal edersem yılda ne kadar kazanırım?"),("Ağır aylar","[Takvim sayfası] En yüksek ödeme yükü olan üç ayı bul. O aylar için her ay ne kadar kenara koymam gerektiğini hesapla."),("Yıllığa çevir","[Ödemeler sayfası] Aylık ödediğim aboneliklerden hangileri genelde yıllık planda daha ucuzdur? Araştırmam için liste çıkar."),("Ödeme günü düzeni","[Ödemeler sayfası] Ödemelerimin ayın hangi günlerine yığıldığını çıkar. Maaş günüm ayın [gün]; hangi ödemelerin tarihini değiştirmeye çalışmalıyım?"),("Sabit gider oranı","Aylık gelirim [tutar]. Bu ödemelerin gelire oranını hesapla ve yorumla."),("Pazarlık mesajı","[Hizmet adı] aboneliğimin fiyatı arttı. Müşteri hizmetlerine indirim ya da daha uygun paket sormak için kısa bir mesaj yaz.")])
    fn=OUT+"ExcelKlinigi_DOZ_Duzenli_Odeme_Takibi.xlsx"; B.save(["Başla",O,"Özet","Takvim","Kategoriler","AI Komutları"],fn)
    wv=load_workbook(fn,data_only=True); o=wv["Özet"]; s=wv[O]; t=wv["Takvim"]; today=o["C3"].value.date()
    from dateutil.relativedelta import relativedelta
    per={"Aylık":1,"3 Aylık":3,"6 Aylık":6,"Yıllık":12}; ay=0
    for i,(ad,kat,tut,sk,ilk,bit) in enumerate(od):
        p=per[sk]; k=0
        while ilk+relativedelta(months=p*k)<today: k+=1
        nxt=ilk+relativedelta(months=p*k); exp="Bitti" if (bit and nxt>bit) else nxt
        got=s.cell(R0+i,10).value; got=got.date() if hasattr(got,"date") else got
        assert got==exp,(ad,got,exp)
        if exp!="Bitti": ay+=tut/p
    assert abs(o["B7"].value-ay)<0.01,(o["B7"].value,ay)
    # takvim 2026: kira 12 ay; kasko haziran ve aralık; spor ağustos, kasım; bulut mart; telefon nisan-aralık
    rowsum=lambda r:[t.cell(r,c).value for c in range(2,14)]
    assert rowsum(R0+1)==[22000]*12; assert rowsum(R0+7)==[0,0,0,0,0,9000,0,0,0,0,0,9000]; assert rowsum(R0+8)==[0]*7+[4500,0,0,4500,0]; assert rowsum(R0+5)==[0,0,1200]+[0]*9; assert rowsum(R0+6)==[0,0,0]+[2500]*9; assert rowsum(R0+9)==[0]*12
    assert t.cell(R0,7).value==22000+1400+650+230+2500+9000; print("  DOZ doğrulandı (bugün:",today,")")

def diyet():
    B=Book("DİYET","Aylık bütçe planı","Her kategoriye bir sınır koy, harcadıkça yaz; hangi kalemde sınırı aştığını ve ay sonuna kadar günde ne harcayabileceğini gör.")
    NP,NH=40,2000; P="Plan"; H="Hareketler"
    wp=B.sheet(P,"Plan","Her kategori için türünü ve aylık plan tutarını yaz.",{"A":28,"B":18,"C":16})
    B.head(wp,4,["Kategori","Tür","Aylık plan (₺)"]); B.body(wp,R0,NP,3,fmt={3:TLI},center=(2,))
    plan=[("Maaş","Gelir",75000),("Ek gelir","Gelir",5000),("Kira","Sabit gider",22000),("Faturalar","Sabit gider",3500),("Market","Değişken gider",14000),("Ulaşım","Değişken gider",4000),("Dışarıda yemek","Değişken gider",5000),("Giyim","Değişken gider",3000),("Eğlence","Değişken gider",2500),("Birikim hesabı","Birikim",10000)]
    B.put(wp,R0,plan); B.dv(wp,"uniq",f"A{R0}:A{R0+NP-1}",f"=COUNTIF($A${R0}:$A${R0+NP-1},A{R0})=1"); B.dv(wp,"list",f"B{R0}:B{R0+NP-1}","Gelir,Sabit gider,Değişken gider,Birikim"); B.dv(wp,"nonneg",f"C{R0}:C{R0+NP-1}")
    PA,PB,PC=[rg(P,c,NP) for c in "ABC"]
    wh=B.sheet(H,"Hareketler","Her gelir, harcama ya da birikim bir satır. Tutarı her zaman artı sayı yaz; türünü kategori belirler.",{"A":12,"B":26,"C":15,"D":34,"E":16})
    B.head(wh,4,["Tarih","Kategori","Tutar (₺)","Not","Tür"],calc=(5,)); B.body(wh,R0,NH,5,calc=(5,),fmt={1:DT,3:TLI},center=(1,5))
    har=[(date(2026,10,1),"Maaş",75000),(date(2026,10,1),"Kira",22000),(date(2026,10,2),"Market",3200),(date(2026,10,2),"Birikim hesabı",10000),(date(2026,10,3),"Ulaşım",900),(date(2026,10,3),"Dışarıda yemek",1800),(date(2026,10,4),"Market",2600),(date(2026,10,4),"Giyim",4200),
         (date(2026,9,1),"Maaş",75000),(date(2026,9,1),"Kira",22000),(date(2026,9,5),"Market",15500),(date(2026,9,10),"Faturalar",3300),(date(2026,9,12),"Ulaşım",3800),(date(2026,9,15),"Ek gelir",4000),(date(2026,9,20),"Eğlence",2900),(date(2026,9,25),"Birikim hesabı",8000)]
    B.put(wh,R0,[(a,b,c,"Örnek") for a,b,c in har])
    for i in range(NH): r=R0+i; wh.cell(r,5,f'=IF(B{r}="","",IFERROR(INDEX({PB},MATCH(B{r},{PA},0))&"","Kategori yok"))')
    wh.freeze_panes="A5"; B.dv(wh,"date",f"A{R0}:A{R0+NH-1}"); B.dv(wh,"list",f"B{R0}:B{R0+NH-1}","="+PA); B.dv(wh,"pos",f"C{R0}:C{R0+NH-1}"); B.cf_eq(wh,f"E{R0}:E{R0+NH-1}","Kategori yok",PINK,RED)
    HA,HB,HC,HE=[rg(H,c,NH) for c in "ABCE"]
    wa=B.sheet("Ay","Ay Görünümü","Yılı ve ayı yaz. Kategori kategori plan, gerçekleşen ve kalan.",{"A":3,"B":28,"C":16,"D":16,"E":16,"F":16,"G":12},grid=False)
    wa["B3"]="Yıl"; B.inp(wa,"C3",2026,"0"); wa["D3"]="Ay (1-12)"; wa["D3"].alignment=Alignment(horizontal="right"); B.inp(wa,"E3",10,"0"); B.dv(wa,"int","E3",(1,12))
    d1="DATE($C$3,$E$3,1)"; d2="DATE($C$3,$E$3+1,1)"
    def tur(t): return f'SUMIFS({HC},{HE},"{t}",{HA},">="&{d1},{HA},"<"&{d2})'
    B.cards(wa,5,[("Gelir","="+tur("Gelir"),TL),("Sabit gider","="+tur("Sabit gider"),TL),("Değişken gider","="+tur("Değişken gider"),TL),("Birikim","="+tur("Birikim"),TL),("Kalan","=B6-C6-D6-E6",TL)])
    wa["B8"]="Bugünden ay sonuna kadar günde harcayabileceğin (değişken gider planına göre):"; wa["B8"].font=font(10,True)
    wa["F8"]=f'=IF(AND(TODAY()>={d1},TODAY()<{d2}),MAX(0,SUMIFS({PC},{PB},"Değişken gider")-D6)/({d2}-TODAY()),"Seçili ay içinde değilsin")'; wa["F8"].number_format=TL; wa["F8"].font=font(11,True,GREEN)
    B.table(wa,10,2,["Kategori","Tür","Plan (₺)","Gerçekleşen (₺)","Fark (₺)","Kullanım"],NP,fmt={4:TL,5:TL,6:TL,7:"0%"},center=(3,7))
    for i in range(NP):
        r=11+i; s_=R0+i
        wa.cell(r,2,f"=IF('{P}'!A{s_}=\"\",\"\",'{P}'!A{s_})"); wa.cell(r,3,f"=IF(B{r}=\"\",\"\",'{P}'!B{s_}&\"\")"); wa.cell(r,4,f"=IF(B{r}=\"\",\"\",'{P}'!C{s_})")
        wa.cell(r,5,f'=IF(B{r}="","",SUMIFS({HC},{HB},B{r},{HA},">="&{d1},{HA},"<"&{d2}))')
        wa.cell(r,6,f'=IF(B{r}="","",IF(OR(C{r}="Gelir",C{r}="Birikim"),E{r}-D{r},D{r}-E{r}))'); wa.cell(r,7,f'=IF(OR(B{r}="",D{r}=0),"",E{r}/D{r})')
    B.cf_f(wa,f"G11:G{10+NP}",'AND(ISNUMBER($G11),$G11>1,OR($C11="Sabit gider",$C11="Değişken gider"))',PINK,RED); B.cf_f(wa,f"F11:F{10+NP}",'AND(ISNUMBER($F11),$F11<0)',None,RED)
    wy=B.sheet("Yıl","Yıl Görünümü","Yılı Ay sayfasından değiştir.",{"A":3,"B":14,"C":16,"D":16,"E":16,"F":16,"G":16},grid=False)
    B.table(wy,4,2,["Ay","Gelir (₺)","Sabit gider (₺)","Değişken gider (₺)","Birikim (₺)","Kalan (₺)"],13,fmt={c:TL for c in range(3,8)})
    for m,ay in enumerate(months(),1):
        r=4+m; wy.cell(r,2,ay)
        for col,t in ((3,"Gelir"),(4,"Sabit gider"),(5,"Değişken gider"),(6,"Birikim")): wy.cell(r,col,f"=SUMIFS({HC},{HE},\"{t}\",{HA},\">=\"&DATE('Ay'!$C$3,{m},1),{HA},\"<\"&DATE('Ay'!$C$3,{m+1},1))")
        wy.cell(r,7,f"=C{r}-D{r}-E{r}-F{r}")
    wy.cell(17,2,"TOPLAM")
    for col,L in zip(range(3,8),"CDEFG"): wy.cell(17,col,f"=SUM({L}5:{L}16)"); wy.cell(17,col).font=font(10,True)
    wy["B19"]="Aylık plan toplamları"; wy["B19"].font=font(11,True,GREEN)
    for k,t in enumerate(("Gelir","Sabit gider","Değişken gider","Birikim")):
        wy.cell(20+k,2,t).font=font(10); c=wy.cell(20+k,3,f'=SUMIFS({PC},{PB},"{t}")'); c.number_format=TL; c.border=box
    chart(wy,"I4","Aylık gelir, gider ve birikim (₺)",Reference(wy,min_col=2,min_row=5,max_row=16),[Reference(wy,min_col=c,min_row=4,max_row=16) for c in (3,4,5,6)],[GREEN,"475569","D92D20","F79009"],w=17)
    B.basla([("NASIL KULLANILIR",None),("1. Plan","Gelir ve gider kategorilerini, türünü ve aylık plan tutarını yaz. Tür dört seçenekten biridir: Gelir, Sabit gider, Değişken gider, Birikim."),("2. Hareketler","Para girdikçe ve çıktıkça bir satır yaz: tarih, kategori, tutar."),("3. Ay","Yılı ve ayı seç. Her kategoride plan, gerçekleşen, fark ve kullanım oranını gör. Sınırı aşan gider kırmızı olur."),("4. Yıl","On iki ayın gelir, gider, birikim ve kalanını yan yana gör."),
      (None,None),("BİLMEN GEREKENLER",None),("Günlük harcanabilir tutar","Değişken gider planından bu ay harcadığın düşülür, kalan tutar ay sonuna kadar kalan güne bölünür. Yalnızca içinde bulunduğun ay seçiliyken hesaplanır."),("Tek plan","Plan tutarı bütün aylar için aynıdır. Ay ay farklı plan tutmaz."),("Kapasite",f"{NP} kategori, {NH} hareket. Kategori adını değiştirirsen eski hareketlerdeki adı da değiştir.")])
    B.ai("Açıklamalardaki kişi adlarını ve hesap numaralarını sil.",[("Sınırı aşanlar","[Ay sayfasındaki tablo] Planı aşan gider kategorilerini ve aşım tutarını listele. Her biri için gelecek ay uygulanabilir bir öneri ver."),("Gerçekçi plan","[Yıl sayfası] Son üç ayın gerçekleşen harcamalarına bakarak her değişken gider için daha gerçekçi bir aylık plan öner."),("Birikim hedefi","[tutar] biriktirmek istiyorum. Mevcut kalan tutarımla bu kaç ay sürer? Süreyi kısaltmak için hangi kalemlerden ne kadar kısmalıyım?"),("Haftalık bütçe","[Plan sayfası] Değişken gider planımı haftalık sınırlara böl ve her hafta için kısa bir kontrol listesi yaz."),("Harcama alışkanlığı","[Hareketler sayfası] Harcamalarımı haftanın günlerine göre topla. En çok hangi gün ve hangi kategoride harcıyorum?"),("Ay sonu özeti","[Ay sayfası] Bu ayı beş cümleyle özetle: ne iyi gitti, ne kötü gitti, gelecek ay neye dikkat etmeliyim.")])
    fn=OUT+"ExcelKlinigi_DIYET_Aylik_Butce_Plani.xlsx"; B.save(["Başla","Hareketler","Ay","Yıl",P,"AI Komutları"],fn)
    wv=load_workbook(fn,data_only=True); a=wv["Ay"]; y=wv["Yıl"]
    assert [a.cell(6,c).value for c in range(2,7)]==[75000,22000,12700,10000,30300],[a.cell(6,c).value for c in range(2,7)]
    assert [y.cell(13,c).value for c in range(3,8)]==[79000,25300,22200,8000,23500],[y.cell(13,c).value for c in range(3,8)]
    assert a["E15"].value==5800 and a["F15"].value==8200 and a["E18"].value==4200 and a["F18"].value==-1200 and abs(a["G18"].value-1.4)<1e-9
    f8=a["F8"].value; assert isinstance(f8,(int,float)),f8; print("  DİYET doğrulandı; günlük harcanabilir:",round(f8,2))
if __name__=="__main__": recete(); doz(); diyet()
