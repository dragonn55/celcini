from lib2 import *
from openpyxl.chart import Reference
D=date; TD=date.today()
def kredi():
    p=P("KREDİ","Kredi taksit ve teklif karşılaştırma","Taksit ne kadar, toplamda ne ödeyeceksin, erken ödeme yaparsan ne kazanırsın, hangi banka teklifi gerçekten daha ucuz."); B=p.B; NT=120
    k=B.sheet("Kredi","Kredi Hesabı","Sarı hücreleri doldur. Ödeme planı aşağıda kendiliğinden çıkar.",{"A":3,"B":38,"C":18,"D":16,"E":16,"F":16,"G":16},grid=False)
    rows=[(3,"Kredi tutarı (₺)",200000,TLI),(4,"Aylık faiz oranı (vergiler hariç)",0.035,"0.00%"),(5,"Faiz üzerinden vergi ve fon oranı",0.30,"0%"),(6,"Vade (ay)",24,"0"),(7,"İlk taksit tarihi",D(2026,11,5),DT),(8,"Tek seferlik masraflar (₺)",2000,TLI),(9,"Erken ödeme yapılacak taksit no (istersen)",6,"0"),(10,"Erken ödeme tutarı (₺)",40000,TLI)]
    for r,lab,val,nf in rows: k.cell(r,2,lab).font=font(10,True); B.inp(k,f"C{r}",val,nf)
    k["D5"]="Kullanıcı girişi. Kredi türüne göre değişir; bankandan doğrula."; k["D5"].font=font(9,c="475467",i=True)
    k["B12"]="Vergili aylık oran"; k["C12"]="=C4*(1+C5)"; k["C12"].number_format="0.000%"
    k["B13"]="Aylık taksit (₺)"; k["C13"]="=IF(OR(C3=\"\",C6=\"\",C6=0),0,IF(C12=0,C3/C6,-PMT(C12,C6,C3)))"; k["C13"].number_format=TL; k["C13"].font=font(12,True,GREEN)
    for r in (12,13): k.cell(r,2).font=font(10,True)
    r0=22; B.head(k,r0,["Taksit no","Tarih","Ödeme (₺)","Faiz ve vergi (₺)","Anapara (₺)","Kalan borç (₺)"],calc=range(2,8),start=2,h=22); B.body(k,r0+1,NT,6,calc=range(2,8),fmt={3:DT,4:TL,5:TL,6:TL,7:TL},center=(2,3),start=2,stripe=True)
    for t in range(1,NT+1):
        r=r0+t; prev="$C$3" if t==1 else f"G{r-1}"
        k.cell(r,2,f'=IF(OR({t}>$C$6,{prev}<=0.005),"",{t})'); k.cell(r,3,f'=IF(B{r}="","",EDATE($C$7,{t-1}))'); k.cell(r,5,f'=IF(B{r}="","",ROUND({prev}*$C$12,2))')
        k.cell(r,4,f'=IF(B{r}="","",MIN({prev}+E{r},$C$13+IF({t}=$C$9,$C$10,0)))'); k.cell(r,6,f'=IF(B{r}="","",D{r}-E{r})'); k.cell(r,7,f'=IF(B{r}="",0,ROUND({prev}-F{r},2))')
    a,b=r0+1,r0+NT
    B.cards(k,15,[("Aylık taksit","=C13",TL),("Toplam geri ödeme",f"=SUM(D{a}:D{b})+C8",TL),("Toplam faiz ve vergi",f"=SUM(E{a}:E{b})",TL,1),("Son taksit",f"=IF(COUNT(B{a}:B{b})=0,\"\",INDEX(C{a}:C{b},COUNT(B{a}:B{b})))","MMMM YYYY"),("Taksit sayısı",f"=COUNT(B{a}:B{b})",NUM)])
    k["B18"]="Erken ödeme yapmasaydın toplam faiz ve vergi (₺)"; k["E18"]="=C13*C6-C3"; k["E18"].number_format=TL
    k["B19"]="Erken ödemenin kazandırdığı (₺)"; k["E19"]="=E18-D16"; k["E19"].number_format=TL; k["E19"].font=font(11,True,GREEN)
    k["B20"]="Kısalan vade (ay)"; k["E20"]="=C6-F16"
    for r in (18,19,20): k.cell(r,2).font=font(10,True)
    B.dv(k,"pos","C3"); B.dv(k,"pct","C4:C5"); B.dv(k,"int","C6",(1,NT)); B.dv(k,"date","C7"); B.dv(k,"nonneg","C8"); B.dv(k,"nonneg","C10")
    c=p.table("Teklifler","Teklif Karşılaştırma","Farklı bankaların tekliflerini yaz; en ucuzu toplam maliyete göre sıralanır.",[("Banka","t",None,22),("Tutar (₺)","m"),("Aylık faiz oranı","n","0.00%"),("Vade (ay)","n"),("Masraf (₺)","m"),
       ("Aylık taksit (₺)","c",("=IF(OR(A{r}=\"\",B{r}=\"\",D{r}=\"\",D{r}=0),\"\",IF(C{r}*(1+'Kredi'!$C$5)=0,B{r}/D{r},-PMT(C{r}*(1+'Kredi'!$C$5),D{r},B{r})))",TL),15),("Toplam geri ödeme (₺)","c",('=IF(F{r}="","",F{r}*D{r}+E{r})',TL),17),("Toplam maliyet (₺)","c",('=IF(G{r}="","",G{r}-B{r})',TL),16),("Sıra","c",(f'=IF(H{{r}}="","",RANK(H{{r}},$H${R0}:$H${R0+9},1))',NUM))],10,
       [("Banka Alfa",200000,0.035,24,2000),("Banka Beta",200000,0.033,24,6000),("Banka Gama",200000,0.036,18,0)],freeze=None)
    p.B.cf_f(c.ws,f"I{R0}:I{R0+9}",f"I{R0}=1",GREEN,WHITE); p.caps.append(f"{NT} aya kadar vade")
    def chk(w):
        x=w["Kredi"]; r=0.035*1.3; pmt=200000*r/(1-(1+r)**-24); bal=200000.0; fz=0; od=0; nn=0
        for t in range(1,25):
            if bal<=0.005: break
            f=round(bal*r,2); pay=min(bal+f,pmt+(40000 if t==6 else 0)); bal=round(bal-(pay-f),2); fz+=f; od+=pay; nn=t
        assert near(x["C13"].value,pmt,0.01) and near(x["D16"].value,fz,0.05) and near(x["C16"].value,od+2000,0.05) and x["F16"].value==nn and near(x["E18"].value,pmt*24-200000,0.05) and x["E20"].value==24-nn,(x["D16"].value,fz,x["F16"].value,nn)
        t=w["Teklifler"]; ex=[]
        for a_,tu,fo,vd,ms in [("A",200000,0.035,24,2000),("B",200000,0.033,24,6000),("C",200000,0.036,18,0)]:
            rr=fo*1.3; ex.append(tu*rr/(1-(1+rr)**-vd)*vd+ms-tu)
        assert all(near(t.cell(R0+i,8).value,ex[i],0.05) for i in range(3)) and [t.cell(R0+i,9).value for i in range(3)]==[sorted(ex).index(e)+1 for e in ex]; print(f"   taksit {pmt:,.2f} ₺; erken ödemeyle {nn} taksit, kazanç {pmt*24-200000-fz:,.0f} ₺")
    p.done([("1. Kredi","Tutarı, aylık faizi, vergi oranını, vadeyi ve ilk taksit tarihini yaz. Taksit ve ödeme planı çıkar."),("2. Erken ödeme","Bir taksitte fazladan ödeme yapmayı düşünüyorsan taksit numarasını ve tutarı yaz; kazancını ve kısalan vadeyi gör."),("3. Teklifler","Farklı bankaların tekliflerini yan yana yaz; toplam maliyete göre sıralanır.")],
      [("Vergili oran","Faiz oranı, yazdığın vergi ve fon oranıyla büyütülür. Varsayılan yüzde 30 bir kullanıcı girişidir; kredi türüne göre değişir."),("Erken ödeme","Fazladan ödenen tutar anaparadan düşer, taksit aynı kalır, vade kısalır. Bankalar erken ödeme ücreti alabilir; dosya bunu hesaba katmaz."),("Banka tablosuyla fark","Bankanın ödeme planı gün hesabı ve yuvarlama nedeniyle birkaç lira farklı olabilir. Bağlayıcı olan bankanın verdiği plandır."),("Finansal tavsiye değildir","Dosya hesap yapar; kredi kullanıp kullanmama kararını vermez.")],
      "Banka adı dışında kişisel bilgi yazma; kimlik ve hesap numarası paylaşma.",[("Teklif yorumu","[Teklifler sayfası] Tekliflerin farkını taksit, toplam maliyet ve vade açısından yorumla."),("Bütçe uyumu","Aylık gelirim [tutar], sabit giderim [tutar]. Bu taksit bütçeme ne kadar yük getirir?"),("Sorular","Bankaya kredi öncesinde sormam gereken soruları listele: masraflar, sigorta, erken ödeme."),("Erken ödeme zamanı","Elimde [tutar] var. Erken ödemeyi şimdi mi yapmalıyım, yoksa acil durum için mi tutmalıyım? Artı ve eksileri yaz.")],"ExcelKlinigi_KREDI_Taksit_Hesaplama.xlsx",order=["Başla","Kredi","Teklifler","AI Komutları"],check=chk)

def icerik():
    p=P("İÇERİK","Sosyal medya içerik performansı","Hangi içerik türü, hangi gün, hangi konu daha çok erişim ve etkileşim getiriyor."); n=2000
    L=p.lists({"Platform":["Instagram","YouTube","TikTok","LinkedIn"],"Tür":["Reels","Gönderi","Hikâye","Video","Shorts","Kaydırmalı gönderi"]})
    c=p.table("İçerikler","İçerikler","Her paylaşım bir satır. Rakamları paylaşımdan birkaç gün sonra, platformun istatistik ekranından yaz.",[("Tarih","d"),("Platform","l",L["Platform"],13),("Tür","l",L["Tür"],16),("Konu","t",None,30),("Erişim","n"),("Beğeni","n"),("Yorum","n"),("Kaydetme","n"),("Paylaşım","n"),("Yeni takipçi","n"),
       ("Etkileşim","c",('=IF(A{r}="","",F{r}+G{r}+H{r}+I{r})',NUM)),("Etkileşim oranı","c",('=IF(OR(K{r}="",E{r}=0,E{r}=""),"",K{r}/E{r})',"0.0%"),14),("y","h",('=IF(OR(A{r}="",E{r}=""),"",E{r}+ROW()/1000000)',None))],n,
       [(D(2026,9,28),"Instagram","Reels","Dağınık Excel önce ve sonra",12400,640,38,410,95,62),(D(2026,9,30),"Instagram","Gönderi","5 Excel kısayolu",3100,210,12,180,22,9),(D(2026,10,1),"YouTube","Shorts","Elle toplamı formüle çevir",8800,350,20,0,40,31),(D(2026,10,2),"Instagram","Reels","Birleştirilmiş hücre hatası",20500,1100,74,820,210,140),(D(2026,10,3),"Instagram","Hikâye","Soru cevap",900,40,6,0,2,1)],freeze="C5")
    o=p.ozet(); A_,E_,K_,J_,T_=c.r("Tarih"),c.r("Erişim"),c.r("Etkileşim"),c.r("Yeni takipçi"),c.r("Tür")
    tr=p.monthly(o,10,2,[("İçerik",mcnt(A_),NUM),("Erişim",msum(E_,A_),NUM),("Etkileşim",msum(K_,A_),NUM),("Yeni takipçi",msum(J_,A_),NUM)])
    p.B.cards(o,5,[("İçerik (yıl)",f"=C{tr}",NUM),("Toplam erişim",f"=D{tr}",NUM),("Etkileşim oranı",f"=IF(D{tr}=0,0,E{tr}/D{tr})","0.0%"),("Yeni takipçi",f"=F{tr}",NUM),("İçerik başına erişim",f"=IF(C{tr}=0,0,D{tr}/C{tr})",NUM)])
    p.bylist(o,10,9,"Türe göre (tüm kayıtlar)",L["Tür"],10,[("İçerik",f"COUNTIF({T_},{{k}})",NUM),("Ortalama erişim",f"IF(COUNTIF({T_},{{k}})=0,0,SUMIFS({E_},{T_},{{k}})/COUNTIF({T_},{{k}}))",NUM),("Etkileşim oranı",f"IF(SUMIFS({E_},{T_},{{k}})=0,0,SUMIFS({K_},{T_},{{k}})/SUMIFS({E_},{T_},{{k}}))","0.0%")],"Tür")
    p.B.label(o,"B25","Haftanın günlerine göre ortalama erişim"); p.B.table(o,26,2,["Gün","İçerik","Ortalama erişim"],7,fmt={3:NUM,4:NUM})
    for i,g in enumerate(["Pazartesi","Salı","Çarşamba","Perşembe","Cuma","Cumartesi","Pazar"],1):
        r=26+i; o.cell(r,2,g); o.cell(r,3,f"=SUMPRODUCT((WEEKDAY({A_},2)={i})*({A_}>0))"); o.cell(r,4,f"=IF(C{r}=0,0,SUMPRODUCT((WEEKDAY({A_},2)={i})*({A_}>0)*{E_})/C{r})")
    p.top(o,23,9,"En çok erişim alan 5 içerik",["Konu","Erişim","Etkileşim oranı"],c.r("y"),c.r("Konu"),[E_,c.r("Etkileşim oranı")],5,[NUM,"0.0%"])
    def chk(w):
        x=w["İçerikler"]; assert x.cell(R0,11).value==1183 and near(x.cell(R0,12).value,1183/12400,1e-9)
        o=w["Özet"]; assert o["C19"].value==2 and o["D20"].value==8800+20500+900 and o["B6"].value==5 and near(o["D6"].value,(1183+424+410+2204+48)/45700,1e-9) and o["I24"].value=="Birleştirilmiş hücre hatası" and o["J11"].value==2 and o["K11"].value==16450
        assert o["C27"].value==1 and o["D27"].value==12400 and o["C31"].value==1 and o["D31"].value==20500
    p.done([("1. Listeler","Kullandığın platformları ve içerik türlerini yaz."),("2. İçerikler","Her paylaşımı tarih, tür, konu ve istatistikleriyle gir."),("3. Özet","Aylık gidiş, türe ve güne göre ortalamalar, en iyi içerikler.")],
      [("Etkileşim oranı","Beğeni, yorum, kaydetme ve paylaşım toplamının erişime bölümüdür."),("Elle giriş","Dosya platformlardan veri çekmez; rakamları sen yazarsın."),("Karşılaştırma","Farklı platformların erişim tanımı aynı değildir; platformları kendi içinde karşılaştırmak daha sağlıklıdır.")],
      "Hesap adını paylaşmak zorunda değilsin.",[("Ne işe yarıyor","[İçerikler sayfası] En çok erişim ve en yüksek etkileşim oranı alan içeriklerin ortak özelliklerini bul."),("Takvim","[Özet sayfası] Tür ve gün ortalamalarına bakarak gelecek hafta için beş içerik fikri ve paylaşım günü öner."),("Başlık","[konu] için dikkat çekici beş açılış cümlesi yaz."),("Zayıf içerik","En düşük erişim alan içeriklerin neden tutmamış olabileceğini yorumla; kesin bilemeyeceğini belirt.")],"ExcelKlinigi_ICERIK_Performans_Takibi.xlsx",check=chk)
if __name__=="__main__":
    kredi(); icerik()
