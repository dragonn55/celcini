from lib import *
import lib
from openpyxl import load_workbook
from datetime import date, timedelta
lib.OUT=os.path.join(os.environ.get("KLINIK_CIKTI","out2"),""); OUT=lib.OUT
T="TODAY()"
class Tab:
    def __init__(s,name,n,lt,ws): s.name,s.n,s.lt,s.ws=name,n,lt,ws
    def r(s,h): return rg(s.name,s.lt[h],s.n)
    def c(s,h,r): return f"'{s.name}'!${s.lt[h]}{r}"
class P:
    def __init__(s,kod,ne,slogan): s.B=Book(kod,ne,slogan); s.kod=kod; s.ne=ne; s.caps=[]; s.hide=[]
    def lists(s,d,name="Listeler",n=30):
        B=s.B; ws=B.sheet(name,"Listeler","Açılır listelerde çıkan seçenekler. Sarı hücrelere kendi seçeneklerini yaz.",{CL(i+1):24 for i in range(len(d))}); out={}
        for j,(t,items) in enumerate(d.items(),1):
            B.head(ws,4,[t],start=j,h=22); B.body(ws,R0,n,1,start=j); B.put(ws,R0,[[x] for x in items],start=j); out[t]=f"'{name}'!${CL(j)}${R0}:${CL(j)}${R0+n-1}"
        return out
    def table(s,name,title,sub,cols,n,samples=(),freeze="A5",uniq=False):
        B=s.B; ws=B.sheet(name,title,sub,{CL(i+1):(c[3] if len(c)>3 else 14) for i,c in enumerate(cols)})
        calc=[i for i,c in enumerate(cols,1) if c[1] in "ch"]; B.head(ws,4,[c[0] for c in cols],calc=calc); fmt={}; cen=[]
        for i,c in enumerate(cols,1):
            k=c[1]; a=c[2] if len(c)>2 else None
            if k=="d": fmt[i]=DT; cen.append(i)
            elif k=="m": fmt[i]=TLI
            elif k=="n": fmt[i]=a or NUM; cen.append(i)
            elif k=="p": fmt[i]="0%"; cen.append(i)
            elif k=="l": cen.append(i)
            elif k in "ch":
                if a[1]: fmt[i]=a[1]
                if a[1] not in (TL,TLI): cen.append(i)
        B.body(ws,R0,n,len(cols),calc=calc,fmt=fmt,center=cen); B.put(ws,R0,samples)
        for i,c in enumerate(cols,1):
            L=CL(i); ref=f"{L}{R0}:{L}{R0+n-1}"; k=c[1]; a=c[2] if len(c)>2 else None
            if k=="d": B.dv(ws,"date",ref)
            elif k in "mn": B.dv(ws,"nonneg",ref)
            elif k=="p": B.dv(ws,"pct",ref)
            elif k=="l": B.dv(ws,"list",ref,("="+a) if "!" in a else a)
            elif k in "ch":
                for r in range(R0,R0+n): ws.cell(r,i,a[0].replace("{r}",str(r)))
                if k=="h": ws.column_dimensions[L].hidden=True
        if uniq: B.dv(ws,"uniq",f"A{R0}:A{R0+n-1}",f"=COUNTIF($A${R0}:$A${R0+n-1},A{R0})=1")
        if freeze: ws.freeze_panes=freeze
        s.caps.append(f"{name}: {n} satır"); return Tab(name,n,{c[0]:CL(i) for i,c in enumerate(cols,1)},ws)
    def ozet(s,sub="Bu sayfada yalnızca yılı değiştir.",widths=None,year=2026,name="Özet"):
        ws=s.B.sheet(name,name,sub,widths or {"A":3,"B":18,"C":16,"D":16,"E":16,"F":16,"G":16,"H":3,"I":26,"J":15,"K":15,"L":15},grid=False)
        if year: ws["B3"]="Yıl"; ws["B3"].font=font(10,True); s.B.inp(ws,"C3",year,"0")
        return ws
    def monthly(s,ws,row,col,series,y="$C$3",title="Aylık tablo"):
        B=s.B; B.label(ws,f"{CL(col)}{row-1}",title); B.table(ws,row,col,["Ay"]+[x[0] for x in series],13,fmt={col+1+j:x[2] for j,x in enumerate(series)})
        for m,ay in enumerate(months(),1):
            ws.cell(row+m,col,ay)
            for j,x in enumerate(series): ws.cell(row+m,col+1+j,x[1](m,row+m))
        ws.cell(row+13,col,"TOPLAM").font=font(10,True)
        for j,x in enumerate(series):
            L=CL(col+1+j); c=ws.cell(row+13,col+1+j,f"=SUM({L}{row+1}:{L}{row+12})" if (len(x)<4 or x[3]) else '=""'); c.font=font(10,True)
        return row+13
    def bylist(s,ws,row,col,title,lref,n,cols,keyhead="Kategori"):
        B=s.B; B.label(ws,f"{CL(col)}{row-1}",title); B.table(ws,row,col,[keyhead]+[c[0] for c in cols],n,fmt={col+1+j:c[2] for j,c in enumerate(cols)})
        sh,rest=lref.split("!"); cl=rest.split("$")[1]
        for k in range(n):
            r=row+1+k; kc=f"{CL(col)}{r}"; ws.cell(r,col,f'=IF({sh}!{cl}{R0+k}="","",{sh}!{cl}{R0+k})')
            for j,c in enumerate(cols): ws.cell(r,col+1+j,f'=IF({kc}="","",{c[1].replace("{k}",kc).replace("{r}",str(r))})')
    def top(s,ws,row,col,title,heads,key,name,vals,n,fmts,mode="LARGE"):
        B=s.B; B.label(ws,f"{CL(col)}{row-1}",title); B.table(ws,row,col,heads,n,fmt={col+1+j:f for j,f in enumerate(fmts)},center=[col+1+j for j,f in enumerate(fmts) if f in (DT,None)])
        for k in range(1,n+1):
            r=row+k; ix=f"MATCH({mode}({key},{k}),{key},0)"; ws.cell(r,col,f'=IFERROR(INDEX({name},{ix})&"","")')
            for j,v in enumerate(vals): ws.cell(r,col+1+j,f'=IF({CL(col)}{r}="","",INDEX({v},{ix}))')
    def done(s,steps,notes,warn,prompts,fname,order=None,check=None):
        B=s.B; B.basla([("NASIL KULLANILIR",None)]+steps+[(None,None),("BİLMEN GEREKENLER",None)]+notes+[("Kapasite","; ".join(s.caps)+".")]); B.ai(warn,prompts)
        names=[w.title for w in B.wb.worksheets]; order=order or (["Başla"]+[x for x in names if x not in ("Başla","Listeler","AI Komutları")]+[x for x in ("Listeler",) if x in names]+["AI Komutları"])
        out=B.save(order,OUT+fname); assert out.get("total_errors")==0,out
        if check: check(load_workbook(OUT+fname,data_only=True)); print("   doğrulandı")
def msum(val,dt,extra="",y="$C$3"): return lambda m,r: "="+mrange(val,dt,y,m,extra)
def mcnt(dt,extra="",y="$C$3"): return lambda m,r: "="+crange(dt,y,m,extra)
def mf(t): return lambda m,r: t.replace("{r}",str(r)).replace("{m}",str(m))
def near(a,b,eps=0.01): return abs(a-b)<eps
