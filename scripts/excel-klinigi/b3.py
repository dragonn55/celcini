from lib2 import *
from openpyxl.chart import Reference
D=date; TD=date.today()
def uyelik():
    p=P("ÜYELİK","Üyelik ve paket takibi","Spor salonu, kurs, stüdyo: kimin üyeliği sürüyor, kiminki bitmek üzere, kimin borcu var."); nP,nU=20,2000
    uP,uB,uO,uD=[rg("Üyelikler",c,nU) for c in "CDEJ"]
    pk=p.table("Paketler","Paketler","Sattığın paketleri, süresini ve ücretini yaz.",[("Paket","t",None,24),("Süre (ay)","n"),("Ücret (₺)","m"),("Süren üyelik","c",(f'=IF(A{{r}}="","",COUNTIFS({uP},A{{r}},{uD},"Aktif")+COUNTIFS({uP},A{{r}},{uD},"Bitiyor"))',NUM))],nP,[("Aylık",1,3000),("3 aylık",3,8000),("Yıllık",12,27000)],uniq=True)
    lk=lambda h: f'IFERROR(INDEX({pk.r(h)},MATCH(C{{r}},{pk.r("Paket")},0)),0)'
    u=p.table("Üyelikler","Üyelikler","Her satış bir satır. Üye paketini yenilerse yeni satır aç.",[("Üye","t",None,24),("Telefon","t",None,14),("Paket","l",pk.r("Paket"),14),("Başlangıç","d"),("Ödenen (₺)","m"),
       ("Bitiş","c",(f'=IF(OR(C{{r}}="",D{{r}}=""),"",EDATE(D{{r}},{lk("Süre (ay)")})-1)',DT)),("Ücret (₺)","c",(f'=IF(F{{r}}="","",{lk("Ücret (₺)")})',TL)),("Borç (₺)","c",('=IF(F{r}="","",G{r}-E{r})',TL)),("Kalan gün","c",(f'=IF(F{{r}}="","",F{{r}}-{T})',NUM)),
       ("Durum","c",(f'=IF(F{{r}}="","",IF({T}<D{{r}},"Başlamadı",IF({T}>F{{r}},"Bitti",IF(F{{r}}-{T}<=7,"Bitiyor","Aktif"))))',None)),("y","h",('=IF(OR(J{r}="Aktif",J{r}="Bitiyor"),F{r}+ROW()/1000000,"")',None))],nU,
       [("Üye Alfa",None,"Aylık",TD-timedelta(27),3000),("Üye Beta",None,"3 aylık",TD-timedelta(40),5000),("Üye Gama",None,"Yıllık",D(2026,2,1),27000),("Üye Delta",None,"Aylık",TD-timedelta(60),3000),("Üye Epsilon",None,"Aylık",TD+timedelta(3),0)])
    for t,fc,tc in (("Bitiyor",AMB,AMBT),("Bitti",GREY,"667085"),("Aktif",MINT,GREEN)): p.B.cf_eq(u.ws,f"J{R0}:J{R0+nU-1}",t,fc,tc)
    p.B.cf_f(u.ws,f"H{R0}:H{R0+nU-1}",f"AND(ISNUMBER(H{R0}),H{R0}>0)",PINK,RED)
    o=p.ozet("Yılı değiştir. Üyelik ve tahsilat, başlangıç ayına yazılır."); I_=u.r("Kalan gün")
    tr=p.monthly(o,10,2,[("Yeni üyelik",mcnt(uB),NUM),("Tahsilat (₺)",msum(uO,uB),TL)])
    p.B.cards(o,5,[("Süren üyelik",f'=COUNTIF({uD},"Aktif")+COUNTIF({uD},"Bitiyor")',NUM),("7 gün içinde biten",f'=COUNTIF({uD},"Bitiyor")',NUM,1),("30 gün içinde biten",f'=COUNTIFS({I_},">=0",{I_},"<=30",{uD},"<>Başlamadı")',NUM),("Toplam borç",f'=SUMIFS({u.r("Borç (₺)")},{u.r("Borç (₺)")},">0")',TL,1),("Yıl tahsilatı",f"=D{tr}",TL)])
    p.top(o,10,6,"Üyeliği en yakın bitecek 10 kişi",["Üye","Bitiş","Paket"],u.r("y"),u.r("Üye"),[u.r("Bitiş"),uP],10,[DT,None],"SMALL"); o.column_dimensions["F"].width=22
    def chk(w):
        x=w["Üyelikler"]; assert [x.cell(R0+i,10).value for i in range(5)]==["Bitiyor","Aktif","Aktif","Bitti","Başlamadı"] and x.cell(R0+1,8).value==3000 and x.cell(R0+2,6).value.date()==D(2027,1,31)
        o=w["Özet"]; assert [o.cell(6,c).value for c in range(2,6)]==[3,1,1,6000] and o["F11"].value=="Üye Alfa" and w["Paketler"].cell(R0,4).value==1
    p.done([("1. Paketler","Paketlerini, süresini ve ücretini yaz."),("2. Üyelikler","Her satışı üye, paket, başlangıç tarihi ve ödenen tutarla gir."),("3. Özet","Süren üyelikler, bitmek üzere olanlar, borçlular ve aylık tahsilat.")],
      [("Bitiş","Başlangıç tarihine paketin süresi eklenir ve bir gün çıkarılır."),("Dondurma","Dosya üyelik dondurmayı tutmaz; donduran üyenin başlangıç tarihini elle kaydır."),("Yenileme","Yenileyen üye için yeni satır aç; eski satır Bitti olarak kalır.")],
      "Üye adlarını ve telefonlarını paylaşma.",[("Yenileme mesajı","Üyeliği bu hafta biten üyeye gönderilecek kısa ve samimi bir yenileme mesajı yaz."),("Kayıp üye","Üyeliği bitmiş ve yenilememiş üyeleri geri çağırmak için bir kampanya metni öner."),("Paket analizi","[Paketler sayfası] Hangi paket daha çok tercih ediliyor? Fiyat ve süre açısından yorumla."),("Borç hatırlatma","Ödemesi eksik üyeye gönderilecek kibar bir hatırlatma yaz.")],"ExcelKlinigi_UYELIK_Paket_Takibi.xlsx",order=["Başla","Üyelikler","Paketler","Özet","AI Komutları"],check=chk)

def yoklama():
    p=P("YOKLAMA","Sınıf yoklama ve devam takibi","Her derste kim geldi, kim gelmedi; devamsızlık sınırına yaklaşan öğrenciyi geç olmadan gör."); B=p.B; nO,nD=40,60; c0=8
    y=B.sheet("Yoklama","Yoklama","Üst satıra ders tarihlerini yaz. Her öğrenci için X geldi, D devamsız, G geç, İ izinli.",{"A":3,"B":24,"C":10,"D":8,"E":8,"F":10,"G":16,**{CL(c):5 for c in range(c0,c0+nD)}})
    y["B3"]="Devamsızlık sınırı"; y["B3"].font=font(10,True); B.inp(y,"C3",0.30,"0%")
    B.head(y,5,["Öğrenci","Devamsız","Geç","İzinli","Devam","Uyarı"]+[""]*nD,calc=(3,4,5,6,7),start=2,h=40)
    for j in range(nD):
        c=y.cell(5,c0+j); c.fill=fill(INPUT); c.font=font(8,True); c.number_format="DD.MM"; c.alignment=Alignment(text_rotation=90,horizontal="center")
    B.body(y,6,nO,6+nD,calc=(3,4,5,6,7),fmt={6:"0%"},center=range(3,c0+nD),start=2)
    dates=[D(2026,9,14)+timedelta(7*k) for k in range(4)]
    for j,d in enumerate(dates): y.cell(5,c0+j,d)
    ogr=[("Öğrenci Alfa","XXXX"),("Öğrenci Beta","XDDX"),("Öğrenci Gama","GXİD"),("Öğrenci Delta","DDXD")]
    L1,L2=CL(c0),CL(c0+nD-1)
    for i in range(nO):
        r=6+i; rg_=f"{L1}{r}:{L2}{r}"
        if i<len(ogr):
            y.cell(r,2,ogr[i][0])
            for j,ch in enumerate(ogr[i][1]): y.cell(r,c0+j,ch)
        y.cell(r,3,f'=IF(B{r}="","",COUNTIF({rg_},"D"))'); y.cell(r,4,f'=IF(B{r}="","",COUNTIF({rg_},"G"))'); y.cell(r,5,f'=IF(B{r}="","",COUNTIF({rg_},"İ"))')
        y.cell(r,6,f'=IF(OR(B{r}="",COUNTA({rg_})=0),"",(COUNTIF({rg_},"X")+COUNTIF({rg_},"G"))/COUNTA({rg_}))')
        y.cell(r,7,f'=IF(OR(B{r}="",COUNTA({rg_})=0),"",IF(C{r}/COUNTA({rg_})>$C$3,"Sınırı aştı",IF(C{r}/COUNTA({rg_})>$C$3*0.7,"Sınıra yakın","")))')
        y.cell(r,1,f'=IF(B{r}="","",C{r}+ROW()/1000000)').font=font(1,c=WHITE)
    r=6+nO; y.cell(r,2,"GELEN").font=font(10,True)
    for j in range(nD): L_=CL(c0+j); y.cell(r,c0+j,f'=IF({L_}5="","",COUNTIF({L_}6:{L_}{5+nO},"X")+COUNTIF({L_}6:{L_}{5+nO},"G"))').font=font(8,True)
    B.dv(y,"list",f"{L1}6:{L2}{5+nO}","X,D,G,İ"); B.dv(y,"date",f"{L1}5:{L2}5"); y.freeze_panes=f"{L1}6"
    for t,fc,tc in (("Sınırı aştı",PINK,RED),("Sınıra yakın",AMB,AMBT)): B.cf_eq(y,f"G6:G{5+nO}",t,fc,tc)
    B.cf_eq(y,f"{L1}6:{L2}{5+nO}","D",PINK,RED); B.cf_eq(y,f"{L1}6:{L2}{5+nO}","G",AMB,AMBT)
    o=p.ozet("Bu sayfaya yazma.",year=None); Y="'Yoklama'!"
    B.cards(o,5,[("Öğrenci",f"=COUNTA({Y}B6:B{5+nO})",NUM),("İşlenen ders",f"=COUNT({Y}{L1}5:{L2}5)",NUM),("Ortalama devam",f"=IFERROR(AVERAGE({Y}F6:F{5+nO}),0)","0%"),("Sınırı aşan",f'=COUNTIF({Y}G6:G{5+nO},"Sınırı aştı")',NUM,1),("Sınıra yakın",f'=COUNTIF({Y}G6:G{5+nO},"Sınıra yakın")',NUM)])
    KEY=f"{Y}$A$6:$A${5+nO}"; p.top(o,9,2,"En çok devamsızlığı olan 5 öğrenci",["Öğrenci","Devamsız","Devam"],KEY,f"{Y}$B$6:$B${5+nO}",[f"{Y}$C$6:$C${5+nO}",f"{Y}$F$6:$F${5+nO}"],5,[NUM,"0%"])
    p.caps.append(f"{nO} öğrenci, {nD} ders")
    def chk(w):
        x=w["Yoklama"]; assert [x.cell(6+i,3).value for i in range(4)]==[0,2,1,3] and [x.cell(6+i,6).value for i in range(4)]==[1,0.5,0.5,0.25] and [x.cell(6+i,7).value for i in range(4)]==[None,"Sınırı aştı","Sınıra yakın","Sınırı aştı"] and [x.cell(46,c0+j).value for j in range(4)]==[3,2,2,2]
        o=w["Özet"]; assert [o.cell(6,c).value for c in range(2,7)]==[4,4,0.5625,2,1] and o["B10"].value=="Öğrenci Delta" and o["B11"].value=="Öğrenci Beta"
    p.done([("1. Yoklama","Öğrencileri alt alta, ders tarihlerini üst satıra yaz."),("2. İşaretle","Her ders için X (geldi), D (devamsız), G (geç), İ (izinli) yaz ya da listeden seç."),("3. Özet","Ortalama devam, sınırı aşanlar ve en çok devamsızlığı olanlar.")],
      [("Devam oranı","Geldi ve geç sayılan derslerin, o öğrenci için işaretlenmiş bütün derslere oranıdır."),("Sınır","Devamsız ders oranı sınırı aşarsa Sınırı aştı, sınırın yüzde 70'ini geçerse Sınıra yakın yazar. İzinli dersler devamsız sayılmaz."),("Tek sınıf","Bir dosya bir sınıfı ya da grubu tutar.")],
      "Öğrenci adlarını takma adla ya da numarayla değiştir.",[("Veli mesajı","Devamsızlığı sınıra yaklaşan öğrencinin velisine gönderilecek kısa ve yapıcı bir bilgilendirme yaz."),("Örüntü","[Yoklama sayfası] Devamsızlığın yoğunlaştığı dersleri ya da haftaları bul."),("Özet rapor","[Özet sayfası] Sınıfın devam durumunu yöneticiye sunmak için üç cümlelik özet yaz."),("Katılım","Derse katılımı artırmak için uygulanabilir beş öneri ver.")],"ExcelKlinigi_YOKLAMA_Devam_Takibi.xlsx",order=["Başla","Yoklama","Özet","AI Komutları"],check=chk)

def karar():
    p=P("KARAR","Karar matrisi","Araba, ev, tedarikçi, iş teklifi: seçenekleri kendi ölçütlerinle puanla, hangisi önde ve neden, gör."); B=p.B; nK,nS=12,8
    m=B.sheet("Matris","Karar Matrisi","Ölçütleri ve önemlerini yaz, seçenekleri üst satıra yaz, her seçeneği her ölçütte 1 ile 10 arasında puanla.",{"A":3,"B":28,"C":11,"D":9,**{CL(c):14 for c in range(5,5+nS)}},grid=False)
    B.head(m,5,["Ölçüt","Önem (1-10)","Pay"]+[""]*nS,calc=(4,),start=2,h=30)
    for j in range(nS): c=m.cell(5,5+j); c.fill=fill(INPUT); c.font=font(10,True)
    B.body(m,6,nK,3+nS,calc=(4,),fmt={4:"0%"},center=range(3,5+nS),start=2)
    olc=[("Fiyat",9),("Yakıt tüketimi",7),("Güvenlik",8),("Bagaj hacmi",4),("İkinci el değeri",6)]; sec=["Araç Alfa","Araç Beta","Araç Gama"]; pu=[[6,8,7],[7,9,5],[8,6,9],[9,5,7],[6,7,8]]
    for j,s in enumerate(sec): m.cell(5,5+j,s)
    for i,(a,w) in enumerate(olc):
        m.cell(6+i,2,a); m.cell(6+i,3,w)
        for j in range(len(sec)): m.cell(6+i,5+j,pu[i][j])
    for i in range(nK): r=6+i; m.cell(r,4,f'=IF(OR(B{r}="",C{r}=""),"",C{r}/SUM($C$6:$C${5+nK}))')
    B.dv(m,"int",f"C6:C{5+nK}",(1,10)); B.dv(m,"int",f"E6:{CL(4+nS)}{5+nK}",(1,10))
    r1=6+nK+1
    for k,lab in enumerate(("Ağırlıklı puan","Sıra","Eksik puan")): m.cell(r1+k,2,lab).font=font(10,True)
    for j in range(nS):
        L=CL(5+j); pr=f"$D$6:$D${5+nK}"; cr=f"{L}6:{L}{5+nK}"
        a=m.cell(r1,5+j,f'=IF({L}5="","",SUMPRODUCT(--($D$6:$D${5+nK}<>""),IF(ISNUMBER($D$6:$D${5+nK}),$D$6:$D${5+nK},0),{cr}))'); a.number_format="0.00"; a.font=font(11,True); a.fill=fill(MINT); a.border=box; a.alignment=Alignment(horizontal="center")
        a.value=f'=IF({L}5="","",SUMPRODUCT({pr},{cr}))'
        b=m.cell(r1+1,5+j,f'=IF({L}{r1}="","",RANK({L}{r1},$E${r1}:${CL(4+nS)}${r1}))'); b.font=font(11,True); b.border=box; b.alignment=Alignment(horizontal="center")
        c=m.cell(r1+2,5+j,f'=IF({L}5="","",COUNTA($B$6:$B${5+nK})-COUNT({cr}))'); c.border=box; c.alignment=Alignment(horizontal="center"); c.font=font(10)
    B.cf_f(m,f"E{r1+1}:{CL(4+nS)}{r1+1}",f"E{r1+1}=1",GREEN,WHITE); B.cf_f(m,f"E{r1+2}:{CL(4+nS)}{r1+2}",f"AND(ISNUMBER(E{r1+2}),E{r1+2}>0)",PINK,RED)
    rr=r1+4; m.cell(rr,2,"Sonuç").font=font(11,True,GREEN)
    PR=f"$E${r1}:${CL(4+nS)}${r1}"; NM=f"$E$5:${CL(4+nS)}$5"
    m.cell(rr+1,2,"Öndeki seçenek"); m.cell(rr+1,3,f'=IFERROR(INDEX({NM},MATCH(MAX({PR}),{PR},0))&"","")'); m.cell(rr+1,3).font=font(12,True,GREEN)
    m.cell(rr+2,2,"İkinciyle puan farkı"); m.cell(rr+2,3,f'=IFERROR(LARGE({PR},1)-LARGE({PR},2),"")'); m.cell(rr+2,3).number_format="0.00"
    m.cell(rr+3,2,"Yorum"); m.cell(rr+3,3,f'=IF(C{rr+2}="","",IF(C{rr+2}<0.3,"Fark çok küçük; önemleri ve puanları bir kez daha gözden geçir.","Fark belirgin."))')
    p.caps.append(f"{nK} ölçüt, {nS} seçenek")
    def chk(w):
        x=w["Matris"]; tw=sum(w_ for _,w_ in olc); exp=[sum(olc[i][1]/tw*pu[i][j] for i in range(5)) for j in range(3)]
        assert all(near(x.cell(19,5+j).value,exp[j],1e-9) for j in range(3)); rk=sorted(exp,reverse=True); assert [x.cell(20,5+j).value for j in range(3)]==[rk.index(e)+1 for e in exp] and x["C24"].value==sec[exp.index(max(exp))] and near(x["C25"].value,rk[0]-rk[1],1e-9) and x.cell(21,5).value==0
    p.done([("1. Ölçütler","Kararda senin için önemli olan ölçütleri ve her birinin önemini 1 ile 10 arasında yaz."),("2. Seçenekler","Seçenekleri üst satırdaki sarı hücrelere yaz."),("3. Puanla","Her seçeneği her ölçütte 1 ile 10 arasında puanla. Yüksek puan her zaman daha iyi demek olmalı: fiyatta ucuz olan yüksek puan alır."),("4. Sonuç","Ağırlıklı puan, sıra ve öndeki seçenek hesaplanır.")],
      [("Ağırlıklı puan","Her ölçütün puanı, o ölçütün toplam önem içindeki payıyla çarpılıp toplanır."),("Eksik puan","Bir seçenekte puanlanmamış ölçüt varsa o seçenek haksız yere düşük görünür; eksik puan satırı bunu gösterir."),("Karar senin","Dosya sayısal bir karşılaştırma verir; son kararı vermez.")],
      "Paylaşılacak kişisel bilgi yoksa doğrudan yapıştırabilirsin.",[("Ölçüt önerisi","[karar konusu] için düşünmem gereken ölçütleri öner; hangilerini atlamış olabilirim?"),("Şeytanın avukatı","[Matris sayfası] Öndeki seçeneğin zayıf yanlarını ve ikinciyi seçmek için geçerli nedenleri yaz."),("Duyarlılık","[Matris sayfası] Hangi ölçütün önemi değişirse sıralama değişir?"),("Puan tutarlılığı","[Matris sayfası] Puanlarımda tutarsız görünen yer var mı?")],"ExcelKlinigi_KARAR_Matrisi.xlsx",order=["Başla","Matris","AI Komutları"],check=chk)

def nakit():
    p=P("NAKİT","13 haftalık nakit akış tahmini","Önümüzdeki üç ayda kasaya ne girecek, ne çıkacak; hangi hafta paran yetmeyecek, şimdiden gör."); n=200; NW=13; B=p.B; S0,BN="'Tahmin'!$C$4","'Tahmin'!$C$3"
    k=p.table("Kalemler","Kalemler","Beklediğin her girişi ve çıkışı yaz. Tekrar eden ödemeler için ilk tarihi ve sıklığı seç.",[("Kalem","t",None,28),("Tür","l","Giriş,Çıkış"),("Tutar (₺)","m"),("İlk tarih","d"),("Tekrar","l","Tek sefer,Haftalık,Aylık"),("Kesinlik","l","Kesin,Olası")],n,
       [("Müşteri Alfa tahsilatı","Giriş",60000,D(2026,10,9),"Tek sefer","Kesin"),("Haftalık nakit satış","Giriş",25000,D(2026,10,5),"Haftalık","Kesin"),("Yeni sipariş (beklenen)","Giriş",80000,D(2026,11,10),"Tek sefer","Olası"),("Kira","Çıkış",30000,D(2026,10,5),"Aylık","Kesin"),("Maaşlar","Çıkış",85000,D(2026,10,31),"Aylık","Kesin"),("Tedarikçi ödemesi","Çıkış",70000,D(2026,10,20),"Tek sefer","Kesin"),("Haftalık malzeme","Çıkış",12000,D(2026,10,7),"Haftalık","Kesin")])
    d=B.sheet("Dağılım","Dağılım (dokunma)","Her kalemin hangi haftaya ne kadar düştüğünün ara hesabı. Değiştirme.",{"A":28})
    for w in range(NW):
        L=CL(2+w); ws=f"({S0}+{7*w})"; we=f"({S0}+{7*w+6})"
        for i in range(n):
            r=R0+i; K="'Kalemler'!"; dd=f"{K}$D{r}"
            d1=f"DATE(YEAR({ws}),MONTH({ws}),MIN(DAY({dd}),DAY(EOMONTH({ws},0))))"; d2=f"DATE(YEAR({we}),MONTH({we}),MIN(DAY({dd}),DAY(EOMONTH({we},0))))"
            if w==0: d.cell(r,1,f'=IF({K}A{r}="","",{K}A{r})')
            d.cell(r,2+w,f'=IF(OR({K}$A{r}="",{K}$C{r}="",{dd}=""),0,{K}$C{r}*IF({K}$E{r}="Haftalık",IF({dd}<={we},1,0),IF({K}$E{r}="Aylık",IF(OR(AND({d1}>={ws},{d1}<={we},{d1}>={dd}),AND({d2}>={ws},{d2}<={we},{d2}<>{d1},{d2}>={dd})),1,0),IF(AND({dd}>={ws},{dd}<={we}),1,0))))')
    t=B.sheet("Tahmin","13 Haftalık Tahmin","Başlangıç nakdini ve ilk haftanın pazartesi tarihini yaz.",{"A":3,"B":9,"C":14,"D":16,"E":16,"F":16,"G":18,"H":3,"I":18,"J":20,"K":22},grid=False)
    t["B3"]="Bugün kasada ve bankada olan (₺)"; t["B4"]="İlk haftanın pazartesi tarihi"
    for rr in (3,4): t.cell(rr,2).font=font(10,True)
    t.merge_cells("C3:D3"); B.inp(t,"C3",40000,TLI); t.merge_cells("C4:D4"); B.inp(t,"C4",D(2026,10,5),DT); t.column_dimensions["B"].width=34
    B.table(t,9,2,["Hafta başı","Giriş (₺)","Çıkış (₺)","Net (₺)","Hafta sonu nakit (₺)"],NW,fmt={2:DT,3:TL,4:TL,5:TL,6:TL},center=(2,)); B.label(t,"B8","Bütün kalemlerle")
    B.table(t,9,9,["Kesin giriş (₺)","Hafta sonu nakit (₺)","Uyarı"],NW,fmt={9:TL,10:TL}); B.label(t,"I8","Yalnız kesin girişlerle (çıkışların tümü)")
    Tr,Ks=rg("Kalemler","B",n),rg("Kalemler","F",n)
    for w in range(NW):
        r=10+w; L=CL(2+w); col=f"'Dağılım'!${L}${R0}:${L}${R0+n-1}"
        t.cell(r,2,f"=$C$4+{7*w}"); t.cell(r,3,f'=SUMIFS({col},{Tr},"Giriş")'); t.cell(r,4,f'=SUMIFS({col},{Tr},"Çıkış")'); t.cell(r,5,f"=C{r}-D{r}"); t.cell(r,6,f"=$C$3+E{r}" if w==0 else f"=F{r-1}+E{r}")
        t.cell(r,9,f'=SUMIFS({col},{Tr},"Giriş",{Ks},"Kesin")'); t.cell(r,10,(f"=$C$3" if w==0 else f"=J{r-1}")+f"+I{r}-D{r}"); t.cell(r,11,f'=IF(J{r}<0,"Nakit yetmiyor",IF(F{r}<0,"Nakit yetmiyor",""))')
    B.cf_eq(t,"K10:K22","Nakit yetmiyor",PINK,RED); B.cf_f(t,"F10:F22","F10<0",None,RED); B.cf_f(t,"J10:J22","J10<0",None,RED)
    B.cards(t,6,[("13 hafta girişi","=SUM(C10:C22)",TL),("13 hafta çıkışı","=SUM(D10:D22)",TL),("En düşük nakit","=MIN(F10:F22)",TL),("En düşük olduğu hafta","=INDEX(B10:B22,MATCH(MIN(F10:F22),F10:F22,0))",DT),("Kesinlerle en düşük nakit","=MIN(J10:J22)",TL,1)],col0=2)
    chart(t,"B25","Hafta sonu nakit (₺)",Reference(t,min_col=2,min_row=10,max_row=22),[Reference(t,min_col=6,min_row=9,max_row=22),Reference(t,min_col=10,min_row=9,max_row=22)],[GREEN,"D92D20"],w=18,kind="line")
    p.caps.append("13 hafta")
    def chk(w):
        x=w["Tahmin"]; st=D(2026,10,5); items=[(a[1],a[2],a[3],a[4],a[5]) for a in k.ws.iter_rows(min_row=R0,max_row=R0+6,max_col=6,values_only=True)]
        import calendar
        def occ(dt,rep,ws,we):
            dt=dt.date() if hasattr(dt,"date") else dt
            if rep=="Haftalık": return 1 if dt<=we else 0
            if rep=="Aylık":
                c=set()
                for ref in (ws,we):
                    dd=D(ref.year,ref.month,min(dt.day,calendar.monthrange(ref.year,ref.month)[1]))
                    if ws<=dd<=we and dd>=dt: c.add(dd)
                return 1 if c else 0
            return 1 if ws<=dt<=we else 0
        bal=40000; bk=40000
        for i in range(13):
            ws=st+timedelta(7*i); we=ws+timedelta(6); gi=sum(t_*occ(dt,rep,ws,we) for tur,t_,dt,rep,ks in items if tur=="Giriş"); ck=sum(t_*occ(dt,rep,ws,we) for tur,t_,dt,rep,ks in items if tur=="Çıkış"); gk=sum(t_*occ(dt,rep,ws,we) for tur,t_,dt,rep,ks in items if tur=="Giriş" and ks=="Kesin")
            bal+=gi-ck; bk+=gk-ck; r=10+i; assert (x.cell(r,3).value,x.cell(r,4).value,x.cell(r,6).value,x.cell(r,10).value)==(gi,ck,bal,bk),(i,x.cell(r,3).value,gi,x.cell(r,4).value,ck)
        print("   haftalar:",[x.cell(10+i,6).value for i in range(13)])
    p.done([("1. Tahmin","Bugün elindeki nakdi ve ilk haftanın pazartesi tarihini yaz."),("2. Kalemler","Beklediğin tahsilatları ve yapacağın ödemeleri tutarı, tarihi, tekrarı ve kesinliğiyle gir."),("3. Tahmin sayfasına dön","13 haftanın giriş, çıkış ve hafta sonu nakdini gör. Nakdin eksiye düştüğü hafta kırmızı uyarı verir.")],
      [("İki senaryo","Soldaki tablo bütün kalemleri sayar. Sağdaki tablo olası girişleri saymaz ama bütün çıkışları sayar; kötü günü gösterir."),("Aylık tekrar","Aylık kalem, ilk tarihin ayın kaçı olduğuna bakılarak her ay o güne yazılır. Kısa aylarda ayın son gününe kayar."),("Tahmin","Sonuç, yazdığın tutar ve tarihlerin doğruluğu kadar doğrudur. Her hafta güncelle.")],
      "Müşteri ve tedarikçi adlarını takma adla değiştir.",[("Açık kapatma","[Tahmin sayfası] Nakdin eksiye düştüğü haftalar için hangi ödemeyi erteleyebilir ya da hangi tahsilatı öne çekebilirim?"),("Senaryo","En büyük tahsilatım iki hafta gecikirse hangi hafta sıkışırım?"),("Görüşme hazırlığı","Tedarikçiden ödeme vadesini uzatmasını istemek için kısa bir konuşma metni hazırla."),("Kontrol","[Kalemler sayfası] Unutmuş olabileceğim düzenli giderleri hatırlat: vergi, sigorta, kredi taksidi gibi.")],"ExcelKlinigi_NAKIT_13_Hafta_Tahmini.xlsx",order=["Başla","Tahmin","Kalemler","Dağılım","AI Komutları"],check=chk)
if __name__=="__main__":
    for f in (uyelik,yoklama,karar,nakit): f()
