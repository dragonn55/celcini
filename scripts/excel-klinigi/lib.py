from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.utils import get_column_letter as CL
import subprocess, json, os, sys
RECALC=os.environ.get("RECALC","/mnt/skills/public/xlsx/scripts/recalc.py")  # LibreOffice ile yeniden hesap (uret.py)
GREEN="1F6F43"; MINT="E8F5EE"; INPUT="FFF9DB"; GREY="F3F4F6"; RED="B42318"; DARK="1B1F23"; WHITE="FFFFFF"; AMB="FEF0C7"; AMBT="93370D"; PINK="FDE2E0"
F="Arial"
TL='#,##0.00 "₺";[Red]-#,##0.00 "₺";"-"'; TLI='#,##0.00 "₺"'; NUM='#,##0;[Red]-#,##0;"-"'; DT="DD.MM.YYYY"; PCT='0%;[Red]-0%;"-"'; PCT1='0.0%'
def font(sz=10,b=False,c=DARK,i=False): return Font(name=F,size=sz,bold=b,color=c,italic=i)
def fill(c): return PatternFill("solid",start_color=c,end_color=c)
thin=Side(style="thin",color="D0D5DD"); box=Border(left=thin,right=thin,top=thin,bottom=thin)
R0=5
class Book:
    def __init__(s,kod,ne,slogan):
        s.kod,s.ne=kod,ne; s.wb=Workbook(); ws=s.wb.active; ws.title="Başla"; s.b=ws; ws.sheet_view.showGridLines=False
        ws.column_dimensions["A"].width=3; ws.column_dimensions["B"].width=24; ws.column_dimensions["C"].width=88
        ws.merge_cells("B2:C2"); c=ws["B2"]; c.value="EXCEL KLİNİĞİ"; c.font=font(11,True,WHITE); c.fill=fill(GREEN); c.alignment=Alignment(vertical="center",indent=1); ws.row_dimensions[2].height=24
        ws.merge_cells("B3:C3"); ws["B3"]=f"{kod}  |  {ne}"; ws["B3"].font=font(20,True,GREEN); ws.row_dimensions[3].height=36
        ws.merge_cells("B4:C4"); ws["B4"]=slogan; ws["B4"].font=font(10,c="475467")
    def basla(s,rows):
        ws=s.b; r=6
        rows=list(rows)+[(None,None),("RENKLER",None),("Sarı hücre","Senin dolduracağın hücre. İçindeki örnekleri silip kendi bilgini yaz."),("Beyaz hücre, açık başlık","Formül. Değiştirme, kendiliğinden hesaplanır.")]
        for a,b in rows:
            if a and b is None:
                ws.merge_cells(start_row=r,start_column=2,end_row=r,end_column=3)
                c=ws.cell(r,2,a); c.font=font(10,True,GREEN); c.fill=fill(MINT); c.alignment=Alignment(indent=1,vertical="center"); ws.row_dimensions[r].height=20
            elif a:
                ws.cell(r,2,a).font=font(10,True); ws.cell(r,3,b).font=font(10)
                ws.cell(r,2).alignment=Alignment(vertical="top",indent=1,wrap_text=True); ws.cell(r,3).alignment=Alignment(wrap_text=True,vertical="top")
                ws.row_dimensions[r].height=44 if len(b)>190 else (30 if len(b)>92 else 18)
                if a=="Sarı hücre": ws.cell(r,2).fill=fill(INPUT)
            r+=1
    def sheet(s,name,title,sub,widths,grid=True):
        ws=s.wb.create_sheet(name); ws.sheet_view.showGridLines=grid
        ws["A1"]=title; ws["A1"].font=font(16,True,GREEN); ws["A2"]=sub; ws["A2"].font=font(9,c="475467",i=True)
        for k,v in widths.items(): ws.column_dimensions[k].width=v
        return ws
    def head(s,ws,row,heads,calc=(),start=1,h=28):
        for j,t in enumerate(heads):
            c=ws.cell(row,start+j,t); c.font=font(10,True,WHITE); c.fill=fill("3B8A5E" if (start+j) in calc else GREEN); c.border=box
            c.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
        ws.row_dimensions[row].height=h
    def body(s,ws,r0,n,ncols,calc=(),fmt=None,center=(),bold=(),start=1,stripe=False):
        fmt=fmt or {}
        for i in range(n):
            for j in range(start,start+ncols):
                c=ws.cell(r0+i,j); c.font=font(10,j in bold); c.border=box
                c.fill=fill(WHITE if j in calc else INPUT)
                if stripe and j in calc and i%2==1: c.fill=fill(GREY)
                if fmt.get(j): c.number_format=fmt[j]
                if j in center: c.alignment=Alignment(horizontal="center")
    def put(s,ws,r0,rows,start=1):
        for i,row in enumerate(rows):
            for j,v in enumerate(row):
                if v is not None: ws.cell(r0+i,start+j,v)
    def dv(s,ws,kind,ref,arg=None,msg=None):
        if kind=="list": v=DataValidation(type="list",formula1=arg if arg.startswith("=") else f'"{arg}"',allow_blank=True); m="Listeden seç."
        elif kind=="date": v=DataValidation(type="date",operator="greaterThan",formula1="36526",allow_blank=True); m="Geçerli bir tarih gir (örnek: 05.10.2026)."
        elif kind=="pos": v=DataValidation(type="decimal",operator="greaterThan",formula1="0",allow_blank=True); m="Sıfırdan büyük bir sayı gir."
        elif kind=="nonneg": v=DataValidation(type="decimal",operator="greaterThanOrEqual",formula1="0",allow_blank=True); m="Sıfır ya da daha büyük bir sayı gir."
        elif kind=="pct": v=DataValidation(type="decimal",operator="between",formula1="0",formula2="1",allow_blank=True); m="0 ile 1 arasında bir oran gir (örnek: yüzde 20 için 0,2)."
        elif kind=="int": v=DataValidation(type="whole",operator="between",formula1=str(arg[0]),formula2=str(arg[1]),allow_blank=True); m=f"{arg[0]} ile {arg[1]} arasında bir tam sayı gir."
        elif kind=="uniq": v=DataValidation(type="custom",formula1=arg,allow_blank=True); m="Bu ad zaten var. Her ad bir kez girilebilir."
        elif kind=="custom": v=DataValidation(type="custom",formula1=arg,allow_blank=True); m="Geçersiz değer."
        v.error=msg or m; v.errorTitle="Kontrol"; v.showErrorMessage=True; ws.add_data_validation(v); v.add(ref)
    def cards(s,ws,row,items,col0=2):
        for j,(lab,f,nf,*red) in enumerate(items):
            a=ws.cell(row,col0+j,lab); a.font=font(9,True,WHITE); a.fill=fill(GREEN); a.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True); a.border=box
            b=ws.cell(row+1,col0+j,f); b.font=font(12,True,RED if (red and red[0]) else DARK); b.fill=fill(MINT); b.alignment=Alignment(horizontal="center",vertical="center"); b.number_format=nf; b.border=box
        ws.row_dimensions[row].height=30; ws.row_dimensions[row+1].height=28
    def label(s,ws,ref,text,sz=11): ws[ref]=text; ws[ref].font=font(sz,True,GREEN)
    def inp(s,ws,ref,val,nf=None,b=True):
        c=ws[ref]; c.value=val; c.font=font(11,b); c.fill=fill(INPUT); c.border=box; c.alignment=Alignment(horizontal="center")
        if nf: c.number_format=nf
    def cf_eq(s,ws,ref,text,fc,tc): ws.conditional_formatting.add(ref,CellIsRule(operator="equal",formula=[f'"{text}"'],fill=fill(fc),font=Font(name=F,bold=True,color=tc)))
    def cf_f(s,ws,ref,formula,fc=None,tc=None,bold=True):
        kw={}
        if fc: kw["fill"]=fill(fc)
        if tc: kw["font"]=Font(name=F,bold=bold,color=tc)
        ws.conditional_formatting.add(ref,FormulaRule(formula=[formula],**kw))
    def table(s,ws,row,col,heads,n,fmt=None,center=()):
        s.head(ws,row,heads,start=col,h=22); s.body(ws,row+1,n,len(heads),calc=range(col,col+len(heads)),fmt=fmt or {},center=center,start=col,stripe=True)
    def ai(s,warn,prompts):
        ws=s.wb.create_sheet("AI Komutları"); ws.sheet_view.showGridLines=False
        ws.column_dimensions["A"].width=3; ws.column_dimensions["B"].width=26; ws.column_dimensions["C"].width=96
        ws["B1"]="Yapay Zeka Komutları"; ws["B1"].font=font(16,True,GREEN)
        ws.merge_cells("B2:C2"); ws["B2"]="İlgili sayfada başlık satırıyla birlikte kayıtlarını seç, kopyala, kullandığın yapay zeka asistanına yapıştır ve altına komutu ekle."
        ws["B2"].font=font(10,c="475467"); ws["B2"].alignment=Alignment(wrap_text=True,vertical="top"); ws.row_dimensions[2].height=30
        ws.merge_cells("B3:C3"); ws["B3"]="Dikkat: "+warn+" Yapay zekanın verdiği rakamları dosyadaki özet sayfasıyla karşılaştır."
        ws["B3"].font=font(10,True,RED); ws["B3"].alignment=Alignment(wrap_text=True,vertical="top"); ws.row_dimensions[3].height=32
        for j,h in enumerate(["Ne için","Komut (kopyala, yapıştır)"]):
            c=ws.cell(5,2+j,h); c.font=font(10,True,WHITE); c.fill=fill(GREEN); c.border=box; c.alignment=Alignment(indent=1)
        for i,(a,b) in enumerate(prompts):
            r=6+i; ws.cell(r,2,a).font=font(10,True); ws.cell(r,3,b).font=font(10)
            for col in (2,3):
                c=ws.cell(r,col); c.border=box; c.alignment=Alignment(wrap_text=True,vertical="center",indent=1)
                if i%2==1: c.fill=fill(GREY)
            ws.row_dimensions[r].height=32
    def save(s,order,fname):
        s.wb._sheets=[s.wb[n] for n in order]
        for w in s.wb.worksheets:
            w.sheet_properties.tabColor=GREEN; w.page_setup.orientation="landscape"; w.page_setup.fitToWidth=1; w.page_setup.fitToHeight=0; w.sheet_properties.pageSetUpPr.fitToPage=True
        s.wb.active=0; s.wb.save(fname)
        out=json.loads(subprocess.run([sys.executable,RECALC,fname,"600"],capture_output=True,text=True).stdout)
        print(f"{s.kod:9} formül={out.get('total_formulas')} hata={out.get('total_errors')} {out.get('error_summary') or ''}")
        return out
def rg(sheet,col,n,r0=R0): return f"'{sheet}'!${col}${r0}:${col}${r0+n-1}"
def months(): return ["Ocak","Şubat","Mart","Nisan","Mayıs","Haziran","Temmuz","Ağustos","Eylül","Ekim","Kasım","Aralık"]
def mrange(val,dt,ycell,m,extra=""): return f'SUMIFS({val},{dt},">="&DATE({ycell},{m},1),{dt},"<"&DATE({ycell},{m+1},1){extra})'
def crange(dt,ycell,m,extra=""): return f'COUNTIFS({dt},">="&DATE({ycell},{m},1),{dt},"<"&DATE({ycell},{m+1},1){extra})'
def chart(ws,anchor,title,cats,series,colors,h=7.5,w=17,kind="bar"):
    ch=BarChart() if kind=="bar" else LineChart()
    if kind=="bar": ch.type="col"
    ch.title=title; ch.height=h; ch.width=w
    for ref in series: ch.add_data(ref,titles_from_data=True)
    ch.set_categories(cats)
    for sr,c in zip(ch.series,colors):
        if kind=="bar": sr.graphicalProperties.solidFill=c
        else: sr.graphicalProperties.line.solidFill=c; sr.graphicalProperties.line.width=22000; sr.smooth=False
    ch.legend.position="b"; ch.y_axis.numFmt='#,##0'; ch.y_axis.majorGridlines=None; ch.x_axis.delete=False; ch.y_axis.delete=False
    ws.add_chart(ch,anchor)
OUT=os.path.join(os.environ.get("KLINIK_CIKTI","out"),"")
